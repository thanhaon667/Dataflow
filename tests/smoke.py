"""
Merge-gate smoke check for the erp_support project: fast, read-only, one PASS/FAIL table.

The autonomous improvement loop (.claude/AUTONOMY.md, step 6) merges a branch only when this exits 0.
It never writes to the database, never starts the Script Center and never triggers a real ClickUp /
DeepSeek / e-mail call. Checks (all required):

  1  Imports         every project .py compiles (in memory, no .pyc written); a module of the erp /
                     desktop / dashboard packages is imported ONLY when every statement that runs at
                     import time is on an ALLOWLIST (imports, defs, classes, constant assignments, a
                     guarded __main__, simple logging / os.environ / Path calls) and it imports no
                     compile-only project module; everything else is compile-only and listed with --verbose
  2  Database        PostgreSQL reachable in a READ ONLY session; the two views and the core tables SELECT
  3  Data Flow map   desktop.flow_data.check_map: no unmapped scripts, no stale nodes / IGNORE entries,
                     no inconsistent edges (unknown nodes, missing armed_by ...)
  4  ERP Desk server the FastAPI app is started in-process on a spare 127.0.0.1 port, its pages, assets and
                     JSON feeds (Reporting, Data Flow ...) must answer HTTP 200, then it is stopped
  5  Secrets & stray  everything since `git merge-base <base> HEAD` - the working-tree diff PLUS every commit
                     and commit message in merge-base..HEAD, so a secret added and removed again on the branch
                     is still caught - and untracked files are scanned for secret patterns (ClickUp pk_,
                     DeepSeek sk-, webhook URLs with tokens, password literals, values copied from .env)
                     and for stray files (*.log, scratch, csv/json outputs, empty files)
  6  Agent system    Law.md, .claude/** playbooks / lessons / workflow exist; lessons.md and the
                     workflow script contain no CR or control bytes; typography: no tracked file names a
                     retired font (only history / quotation LINES are exempt) and shell.css, shell.html and
                     erp/typography.py agree about the font
  7  Left nothing   no new files in the work tree, no logging handlers / threads / listening port left over
  8  Branch discipline  FAIL when the checked-out branch IS main / master and tracked source files are
                     modified (Law.md rule 1 / lesson L-053); the owner's own dirty files are excluded
  9  Gate self-test  tests/smoke_selftest.py: each protection above (run_smoke.bat + -B, merge-base and the
                     history scan, the import allowlist, the strict placeholder rules, git / OSError handling,
                     branch discipline, the typography line rules and the font cross-check) has a scenario
                     that must pass - a protection broken by a later edit fails HERE

Run (from the project root; -B is required so tests/__pycache__ is never created, lesson L-077):
    venv\\Scripts\\python.exe -B -m tests.smoke            (or run_smoke.bat)
    venv\\Scripts\\python.exe -B -m tests.smoke --verbose  (details for every check)
    venv\\Scripts\\python.exe -B -m tests.smoke --only 3,5 (subset; exits 3 even when it passes: not a gate)

Exit code: 0 = every check passed, 1 = at least one FAIL, 2 = the run itself timed out / crashed,
3 = a subset (--only) passed. A DB, server or git problem is a clean FAIL line, never a traceback.
"""
from __future__ import annotations

import argparse
import ast
import contextlib
import fnmatch
import hashlib
import importlib
import io
import json
import logging
import os
import re
import shutil
import socket
import subprocess
import sys
import threading
import time
import traceback
import urllib.error
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path

sys.dont_write_bytecode = True                 # the smoke run must not create .pyc files
os.environ.setdefault("MPLBACKEND", "Agg")     # importing matplotlib-using modules must never open a window

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

log = logging.getLogger("smoke")

# ----------------------------------------------------------------------------- configuration
PACKAGE_DIRS_SKIP = {"tests", "venv", ".venv", "env", "node_modules", "__pycache__", "site-packages", "build", "dist"}
# Modules that must not be imported even though they live in a package dir: "relative/path.py": "why".
# (Streamlit pages, __main__.py and modules that call an entry point at top level are detected automatically.)
IMPORT_SKIP: dict[str, str] = {}

# Import safety is an ALLOWLIST (lesson L-075): check 1 imports a module only when every statement that runs at import time is
# known to be harmless. A brand-new entry-point name nobody thought of is therefore NOT safe just because no denylist names it.
# Calls allowed at import time (dotted name as written in the source). They only read the environment or build an in-memory
# value; the importer snapshots and restores the root logger's handlers (L-010) and the working directory.
SAFE_CALLS = frozenset({
    "logging.getLogger", "logging.basicConfig", "logging.Formatter", "os.environ.get", "os.environ.setdefault", "os.getenv",
    "os.chdir", "os.path.join", "os.path.dirname", "os.path.abspath", "os.path.realpath", "load_dotenv", "dotenv.load_dotenv",
    "sys.path.insert", "sys.path.append", "sys.stdout.reconfigure", "sys.stderr.reconfigure",
    "Path", "pathlib.Path", "re.compile", "str", "int", "float", "bool", "bytes", "list", "dict", "set", "tuple", "frozenset",
    "len", "max", "min", "sorted", "range", "sum", "any", "all", "abs", "round", "getattr", "hasattr", "isinstance", "callable",
    "repr", "zip", "enumerate", "reversed", "time", "date", "datetime", "timedelta", "timezone", "field", "dataclasses.field",
    "TypeVar", "namedtuple", "collections.namedtuple", "defaultdict", "collections.defaultdict", "threading.Lock",
    "threading.RLock", "threading.Event", "urllib.request.build_opener", "urllib.request.ProxyHandler", "Enum", "enum.Enum",
})
# Methods allowed on a VALUE (a literal, an f-string, the result of a safe call ...), never on a bare name that could be a
# module (os.replace, shutil.copy ...): pure string / path / mapping helpers.
SAFE_METHODS = frozenset({
    "resolve", "absolute", "expanduser", "joinpath", "with_suffix", "with_name", "relative_to", "as_posix", "strip", "lstrip",
    "rstrip", "lower", "upper", "title", "casefold", "split", "rsplit", "splitlines", "join", "format", "startswith", "endswith",
    "replace", "encode", "decode", "removeprefix", "removesuffix", "keys", "values", "items", "isoformat", "strftime", "copy",
})
# Decorators allowed on a def / class (dotted name, optionally called: @dataclass(frozen=True)).
SAFE_DECORATORS = frozenset({
    "dataclass", "dataclasses.dataclass", "staticmethod", "classmethod", "property", "abstractmethod", "abc.abstractmethod",
    "lru_cache", "functools.lru_cache", "cache", "functools.cache", "wraps", "functools.wraps", "contextmanager",
    "contextlib.contextmanager", "overload", "typing.overload", "total_ordering", "functools.total_ordering",
})
# Reviewed exceptions, one module each: extra top-level calls / decorators that are safe IN THAT MODULE, with the reason.
# A new top-level call in any other module (or a new one here) makes that module compile-only until somebody reads it and
# adds it here. Keep the "why" honest: it is the only evidence that a human looked at the call.
IMPORT_SAFE_EXTRA: dict[str, dict] = {
    "erp/config.py": {"calls": {"_int_env", "_bool_env", "_float_env"}, "why": "the _*_env helpers only read an environment variable"},
    "erp/db.py": {"calls": {"make_engine", "make_read_engine", "sessionmaker", "declarative_base",
                            "threading.Semaphore", "threading.Lock"},
                  "why": "make_engine/make_read_engine build a lazy SQLAlchemy engine (no connection opened until "
                         "the first query - make_read_engine just calls make_engine with extra connect_args); "
                         "threading.Semaphore()/threading.Lock() only construct in-memory sync objects at import "
                         "time - no I/O, no connection, no thread started"},
    "erp/webhook_app.py": {"calls": {"FastAPI"}, "decorators": {"app.post", "app.get", "model_validator", "field_validator"},
                           "why": "builds the in-memory FastAPI app and registers its routes; nothing is served"},
    "desktop/flow_data.py": {"calls": {"_file_mtime", "_db_guard"},
                             "why": "_file_mtime is one os.stat(); _db_guard only wraps a function"},
    "desktop/today_sql.py": {"calls": {"lead_past_sla_sql", "lead_awaiting_sql"},
                            "why": "pure string builders of SQL fragments (no query is run at import)"},
    "desktop/leads_data.py": {"calls": {"td.lead_past_sla_sql", "td.lead_awaiting_sql"},
                              "why": "the same pure SQL-fragment builders of desktop/today_data.py (no query is run at import)"},
    "desktop/launcher.py": {"calls": {"typography.streamlit_theme_args"},
                            "why": "builds the Streamlit --theme.* flag list from erp/typography.py: pure strings, nothing is started"},
    "desktop/winutil.py": {"calls": {"ctypes.WinDLL", "ctypes.POINTER", "ctypes.WINFUNCTYPE", "ctypes.windll.kernel32.GetConsoleWindow"},
                           "why": "loads kernel32 / user32 and declares call signatures; nothing is done to a window"},
}

DB_VIEWS = ["v_leads_summary", "v_tickets_summary"]
DB_TABLES = ["departments", "users", "customers", "tickets", "ticket_comments", "clickup_sync",
             "leads", "lead_assignments", "lead_ai_analysis", "lead_clickup_sync", "lead_updates"]
DB_CONNECT_TIMEOUT = 5

RESERVED_PORTS = {47650, 47651, 8000, 8502}    # ERP Desk shell + Script Center, lead webhook, standalone Script Center
AGENT_FILES = ["Law.md", "README.md", "PROJECT_NOTES.md", ".gitattributes", ".claude/AUTONOMY.md", ".claude/MODEL_POLICY.md",
               ".claude/backlog.md", ".claude/journal.md", ".claude/lessons.md",
               ".claude/agents/erp-builder.md", ".claude/agents/erp-reviewer.md", ".claude/agents/erp-auditor.md",
               ".claude/workflows/build-feature.js"]
NO_CR_FILES = [".claude/lessons.md", ".claude/workflows/build-feature.js"]   # must be byte-clean (L-066 / L-068)

# The owner's own uncommitted files on main (Law.md rule 9): tolerated as uncommitted edits, never as part of a change.
OWNER_FILES = {"processed_orders.csv", "summary_report.json"}
# Real runtime logs of the apps (git-ignored). Any OTHER log file lying around is a scratch file somebody forgot.
KNOWN_LOGS = {"erp_desktop.log", "erp_desktop_streamlit.log", "daily_digest.log", "script_center.log"}
# Runtime outputs of the project: never part of a change.
KNOWN_OUTPUTS = {"daily_digest_latest.json", "daily_digest_history.json",
                 "processed_orders.csv", "summary_report.json"}
STRAY_GLOBS = ["*.log", "*.tmp", "*.temp", "*.bak", "*.orig", "*.rej", "*.swp", "*.swo", "*~", "*.pyc", "*.pyo",
               "*.pid", "*.dmp", "nohup.out", "nul"]
STRAY_NAME_RE = re.compile(r"(?i)^(scratch|tmp|temp)([._\-0-9]|$)")
DATA_OUTPUT_EXT = {".csv", ".json", ".jsonl", ".ndjson", ".xlsx", ".xls", ".parquet", ".pkl", ".sqlite", ".db"}
DATA_OUTPUT_OK_PREFIX = (".claude/", "tests/fixtures/", "desktop/static/", ".vscode/")
SECRET_FILE_GLOBS = [".env", ".env.*", "*.pem", "*.key", "*.pfx", "*.p12", "id_rsa*", "id_ed25519*"]
SECRET_FILE_OK = {".env.example"}
BINARY_EXT = {".png", ".jpg", ".jpeg", ".gif", ".ico", ".woff", ".woff2", ".ttf", ".pdf", ".zip", ".gz", ".pyc"}

# ----------------------------------------------------------------------------- results
PASS, FAIL = "PASS", "FAIL"


@dataclass
class Result:
    n: int
    name: str
    status: str = FAIL
    reason: str = "did not run"
    details: list[str] = field(default_factory=list)
    seconds: float = 0.0


RESULTS: list[Result] = []
_SECRETS: list[str] | None = None


def _env_secret_values() -> list[str]:
    """Secret-looking values from .env, kept in memory only so they can be (a) searched for and (b) scrubbed from output."""
    global _SECRETS
    if _SECRETS is None:
        vals: list[str] = []
        try:
            from dotenv import dotenv_values
            for k, v in dotenv_values(ROOT / ".env").items():
                if v and len(v) >= 6 and re.search(r"(?i)pass|secret|token|api[_-]?key|webhook|hook", k or "") \
                        and v.strip().lower() not in ("true", "false", "localhost", "127.0.0.1"):
                    vals.append(v)
        except Exception:  # noqa: BLE001 - no .env / no dotenv: nothing to protect or search for
            pass
        _SECRETS = sorted(set(vals), key=len, reverse=True)
    return _SECRETS


def scrub(text: str) -> str:
    """Never let a secret value reach the console (exception texts can echo connection details)."""
    for v in _env_secret_values():
        text = text.replace(v, "***")
    return text


def one_line(text: str, limit: int = 200) -> str:
    first = next((ln.strip() for ln in scrub(str(text)).splitlines() if ln.strip()), "")
    return first if len(first) <= limit else first[: limit - 3] + "..."


# ----------------------------------------------------------------------------- shared helpers
class GateError(RuntimeError):
    """A prerequisite of the gate (git, the repository, the base branch ...) is unusable. It always becomes a clean FAIL row
    carrying this message - never a traceback on the console (lesson L-076)."""


_RUN = subprocess.run          # one seam so tests/smoke_selftest.py can simulate "git is missing" / "git hangs"


def git(*args: str, timeout: int = 30, cwd: Path | None = None) -> subprocess.CompletedProcess:
    """Read-only git call (no optional index refresh / lock).

    Raises GateError when git cannot be run AT ALL (not on PATH, cannot start, timed out). A non-zero exit code of git
    itself is returned unchanged, because only the caller knows whether that is a failure (`rev-parse --verify` uses it
    as a question) or not.
    """
    env = dict(os.environ, GIT_OPTIONAL_LOCKS="0", GIT_TERMINAL_PROMPT="0", LC_ALL="C")
    what = "git " + (args[0] if args else "")
    try:
        return _RUN(["git", "-c", "core.quotepath=false", "--no-pager", *args], cwd=cwd or ROOT, env=env, capture_output=True,
                    text=True, encoding="utf-8", errors="replace", timeout=timeout, check=False)
    except subprocess.TimeoutExpired:
        raise GateError(f"{what} did not finish within {timeout} s") from None
    except FileNotFoundError:
        if shutil.which("git") is None:
            raise GateError("git is not installed / not on PATH - the gate cannot inspect the repository") from None
        raise GateError(f"{what} could not be started (is the folder {cwd or ROOT} missing?)") from None
    except OSError as exc:
        raise GateError(f"{what} could not be started: {type(exc).__name__}: {one_line(str(exc), 100)}") from None


def current_branch(root: Path = ROOT) -> str:
    """Name of the checked-out branch, for the table header and check 8. Never raises."""
    try:
        r = git("branch", "--show-current", cwd=root, timeout=10)
    except GateError as exc:
        return f"(unknown: {one_line(str(exc), 60)})"
    if r.returncode != 0:
        return "(not a git repository)"
    return r.stdout.strip() or "(detached HEAD)"


def list_py_files() -> list[Path]:
    """Every project .py file (hidden folders and virtualenvs pruned)."""
    out: list[Path] = []
    for dirpath, dirnames, filenames in os.walk(ROOT):
        dirnames[:] = sorted(d for d in dirnames if not d.startswith(".") and d not in ("__pycache__", "node_modules", "site-packages")
                             and not (Path(dirpath, d) / "pyvenv.cfg").exists())
        out.extend(Path(dirpath, f) for f in sorted(filenames) if f.endswith(".py"))
    return out


def rel(p: Path) -> str:
    return p.relative_to(ROOT).as_posix()


def _dotted(node: ast.AST) -> str:
    if isinstance(node, ast.Name):
        return node.id
    if isinstance(node, ast.Attribute):
        base = _dotted(node.value)
        return f"{base}.{node.attr}" if base else node.attr
    return ""


def _is_main_guard(test: ast.AST) -> bool:
    return isinstance(test, ast.Compare) and isinstance(test.left, ast.Name) and test.left.id == "__name__"


_SAFE_EXPR_NODES = (ast.Constant, ast.Name, ast.Attribute, ast.Subscript, ast.Slice, ast.BinOp, ast.UnaryOp, ast.BoolOp,
                    ast.Compare, ast.IfExp, ast.Tuple, ast.List, ast.Set, ast.Dict, ast.JoinedStr, ast.FormattedValue,
                    ast.Starred, ast.ListComp, ast.SetComp, ast.DictComp, ast.GeneratorExp, ast.comprehension, ast.keyword)


def _expr_problem(node: ast.AST | None, calls: frozenset[str] = frozenset()) -> str | None:
    """None when evaluating this expression at import time only computes a value; else the first reason why it does not."""
    if node is None:
        return None
    if isinstance(node, ast.Lambda):                      # the body runs only when called; the default values run now
        for d in [*node.args.defaults, *node.args.kw_defaults]:
            problem = _expr_problem(d, calls)
            if problem:
                return problem
        return None
    if isinstance(node, ast.Call):
        name = _dotted(node.func)
        ok = bool(name) and (name in SAFE_CALLS or name in calls)
        if not ok and isinstance(node.func, ast.Attribute) and node.func.attr in SAFE_METHODS \
                and not isinstance(node.func.value, (ast.Name, ast.Attribute)):     # a method of a VALUE, not of a module
            ok = _expr_problem(node.func.value, calls) is None
        if not ok:
            return f"line {node.lineno}: calls {name or ast.unparse(node.func)[:40]}() at import time (not on the allowlist)"
        parts: list[ast.AST] = [*node.args, *[k.value for k in node.keywords]]
    elif isinstance(node, _SAFE_EXPR_NODES):
        parts = list(ast.iter_child_nodes(node))
    else:
        return f"line {getattr(node, 'lineno', '?')}: a {type(node).__name__} expression at import time (not on the allowlist)"
    for child in parts:
        if isinstance(child, (ast.expr_context, ast.operator, ast.unaryop, ast.cmpop, ast.boolop)):
            continue                                       # an operator token is not code that runs
        problem = _expr_problem(child, calls)
        if problem:
            return problem
    return None


def _decorator_problem(dec: ast.AST, extra: dict) -> str | None:
    target = dec.func if isinstance(dec, ast.Call) else dec
    name = _dotted(target)
    if isinstance(target, ast.Attribute) and target.attr in ("setter", "getter", "deleter"):
        name, ok = "property." + target.attr, True         # @x.setter: rebinding a property
    else:
        ok = bool(name) and (name in SAFE_DECORATORS or name in extra.get("decorators", ()))
    if not ok:
        return f"line {dec.lineno}: decorator @{name or '?'} runs at import time (not on the allowlist)"
    if isinstance(dec, ast.Call):
        for a in [*dec.args, *[k.value for k in dec.keywords]]:
            problem = _expr_problem(a, frozenset(extra.get("calls", ())))
            if problem:
                return problem
    return None


def _stmt_problem(stmt: ast.stmt, extra: dict) -> str | None:
    """None when this statement is on the import-time allowlist; else why the module must stay compile-only."""
    calls = frozenset(extra.get("calls", ()))

    def exprs(*nodes) -> str | None:
        for n in nodes:
            problem = _expr_problem(n, calls)
            if problem:
                return problem
        return None

    def block(body: list[ast.stmt]) -> str | None:
        for s in body:
            problem = _stmt_problem(s, extra)
            if problem:
                return problem
        return None

    if isinstance(stmt, (ast.Import, ast.ImportFrom)):
        mods = [a.name for a in stmt.names] if isinstance(stmt, ast.Import) else [stmt.module or ""]
        if any(m.split(".")[0] == "streamlit" for m in mods):
            return "a Streamlit page: importing it renders the UI"
        return None
    if isinstance(stmt, (ast.FunctionDef, ast.AsyncFunctionDef)):
        for d in stmt.decorator_list:
            problem = _decorator_problem(d, extra)
            if problem:
                return problem
        return exprs(*stmt.args.defaults, *stmt.args.kw_defaults)      # default values run at def time; the body does not
    if isinstance(stmt, ast.ClassDef):
        for d in stmt.decorator_list:
            problem = _decorator_problem(d, extra)
            if problem:
                return problem
        return exprs(*stmt.bases, *[k.value for k in stmt.keywords]) or block(stmt.body)   # a class body RUNS at import
    if isinstance(stmt, ast.Assign):
        return exprs(*stmt.targets, stmt.value)
    if isinstance(stmt, ast.AnnAssign):
        return exprs(stmt.target, stmt.annotation, stmt.value)
    if isinstance(stmt, ast.AugAssign):
        return exprs(stmt.target, stmt.value)
    if isinstance(stmt, ast.Expr):
        if isinstance(stmt.value, ast.Constant):
            return None                                    # a docstring
        return exprs(stmt.value)                           # only a call on the allowlist gets through
    if isinstance(stmt, ast.Pass):
        return None
    if isinstance(stmt, ast.If):
        if _is_main_guard(stmt.test):
            return block(stmt.orelse)                      # the guarded body never runs on import
        return exprs(stmt.test) or block(stmt.body) or block(stmt.orelse)
    if isinstance(stmt, (ast.Try, getattr(ast, "TryStar", ast.Try))):
        for h in stmt.handlers:
            problem = exprs(h.type) or block(h.body)
            if problem:
                return problem
        return block(stmt.body) or block(stmt.orelse) or block(stmt.finalbody)
    return f"line {stmt.lineno}: a top-level {type(stmt).__name__} statement (not on the allowlist)"


def _rel_or_name(path: Path) -> str:
    try:
        return rel(path)
    except ValueError:
        return path.name


def import_hazard(tree: ast.Module, path: Path, extra: dict | None = None) -> str | None:
    """Why the gate must NOT import this module (importing it would run a page, a program, an unknown call), else None.

    An ALLOWLIST (L-075): every statement that runs at import time must be known-safe. `extra` is the module's reviewed
    IMPORT_SAFE_EXTRA entry (looked up by relative path when not given).
    """
    if path.name == "__main__.py":
        return "a __main__ module runs its program"
    if extra is None:
        extra = IMPORT_SAFE_EXTRA.get(_rel_or_name(path), {})
    for stmt in tree.body:
        problem = _stmt_problem(stmt, extra)
        if problem:
            return problem
    return None


def _import_scope(body: list[ast.stmt]):
    """Statements that run when the module is imported (nested through if / try / class bodies; def bodies and the
    __main__ guard excluded)."""
    for s in body:
        yield s
        if isinstance(s, ast.If):
            yield from _import_scope(s.orelse if _is_main_guard(s.test) else s.body + s.orelse)
        elif isinstance(s, ast.ClassDef):
            yield from _import_scope(s.body)
        elif isinstance(s, (ast.Try, getattr(ast, "TryStar", ast.Try))):
            yield from _import_scope(s.body + [x for h in s.handlers for x in h.body] + s.orelse + s.finalbody)


def project_imports(tree: ast.Module, modname: str, is_pkg: bool, known: set[str]) -> set[str]:
    """Project modules (names in `known`) that importing this one imports too, including their parent packages."""
    found: set[str] = set()

    def add(name: str) -> None:
        parts = name.split(".")
        for i in range(1, len(parts) + 1):
            if ".".join(parts[:i]) in known:
                found.add(".".join(parts[:i]))

    pkg = modname if is_pkg else modname.rpartition(".")[0]
    for s in _import_scope(tree.body):
        if isinstance(s, ast.Import):
            for a in s.names:
                add(a.name)
        elif isinstance(s, ast.ImportFrom):
            if s.level:
                anchor = pkg.split(".")
                anchor = anchor[:len(anchor) - (s.level - 1)] if s.level > 1 else anchor
                base = ".".join([*anchor, *(s.module.split(".") if s.module else [])])
            else:
                base = s.module or ""
            if base:
                add(base)
                for a in s.names:
                    add(f"{base}.{a.name}")
    found.discard(modname)
    return found


def classify_imports(mods: dict[str, tuple[Path, ast.Module]]) -> dict[str, str | None]:
    """{module name: reason it stays compile-only, or None when the gate may import it}.

    A module that imports a compile-only project module runs that module's top level too, so it is compile-only as
    well; the verdict is propagated until nothing changes.
    """
    verdict: dict[str, str | None] = {m: import_hazard(t, p) for m, (p, t) in mods.items()}
    known = set(mods)
    deps = {m: project_imports(t, m, p.name == "__init__.py", known) for m, (p, t) in mods.items()}
    changed = True
    while changed:
        changed = False
        for m, ds in deps.items():
            if verdict[m] is None:
                bad = sorted(d for d in ds if verdict[d] is not None)
                if bad:
                    verdict[m] = f"imports {bad[0]}, which is compile-only"
                    changed = True
    return verdict


# ============================================================================ check 1: imports
def check_imports() -> tuple[str, str, list[str]]:
    files = list_py_files()
    details: list[str] = []
    bad: list[str] = []
    trees: dict[Path, ast.Module] = {}
    for p in files:
        try:
            src = p.read_bytes()
            trees[p] = ast.parse(src, filename=str(p))
            compile(src, str(p), "exec", dont_inherit=True)      # what py_compile does, without writing a .pyc
        except SyntaxError as exc:
            bad.append(f"{rel(p)}:{exc.lineno}: {exc.msg}")
        except Exception as exc:  # noqa: BLE001 - unreadable file, bad encoding declaration, ...
            bad.append(f"{rel(p)}: {type(exc).__name__}: {one_line(str(exc), 100)}")
    if bad:
        return FAIL, f"{len(bad)} file(s) do not compile: {bad[0]}" + (" (+more)" if len(bad) > 1 else ""), bad

    pkg_roots = sorted({p.relative_to(ROOT).parts[0] for p in files if len(p.relative_to(ROOT).parts) > 1
                        and p.relative_to(ROOT).parts[0] not in PACKAGE_DIRS_SKIP})
    mods: dict[str, tuple[Path, ast.Module]] = {}
    compile_only: list[str] = []
    for p in files:
        parts = p.relative_to(ROOT).parts
        if len(parts) < 2 or parts[0] not in pkg_roots:
            compile_only.append(f"{rel(p)} (root script: never imported)" if len(parts) == 1 else f"{rel(p)} (not a project package)")
            continue
        mod = ".".join(parts[:-1]) if parts[-1] == "__init__.py" else ".".join(parts)[:-3]
        mods[mod] = (p, trees[p])
    verdict = classify_imports(mods)      # an ALLOWLIST: anything not known to be harmless at import time stays compile-only (L-075)
    to_import: list[tuple[str, Path]] = []
    for mod, (p, _t) in mods.items():
        why = IMPORT_SKIP.get(rel(p)) or verdict[mod]
        if why:
            compile_only.append(f"{rel(p)} ({why})")
        else:
            to_import.append((mod, p))
    compile_only.sort()

    root_logger = logging.getLogger()
    failures: list[str] = []
    stray_handlers: list[str] = []
    printed: list[str] = []
    cwd0 = os.getcwd()
    for mod, p in to_import:
        before, level = list(root_logger.handlers), root_logger.level
        out = io.StringIO()
        try:
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(out):
                importlib.import_module(mod)
        except BaseException as exc:  # noqa: BLE001 - SystemExit from a module that exits at import counts as a failure too
            if isinstance(exc, KeyboardInterrupt):
                raise
            failures.append(f"{mod}: {type(exc).__name__}: {one_line(str(exc), 120)}")
        finally:
            for h in list(root_logger.handlers):
                if h not in before:            # a script-style module called logging.basicConfig at import time (L-010)
                    root_logger.removeHandler(h)
                    if isinstance(h, logging.FileHandler):
                        stray_handlers.append(f"{mod} opened a log FILE at import time ({getattr(h, 'baseFilename', '?')})")
                        h.close()
            root_logger.setLevel(level)
            if os.getcwd() != cwd0:                 # an allowlisted os.chdir at import time must not move the gate
                os.chdir(cwd0)
        if out.getvalue().strip():
            printed.append(mod)
    if failures:
        return FAIL, f"{len(failures)} module(s) fail to import: {failures[0]}" + (" (+more)" if len(failures) > 1 else ""), failures
    if stray_handlers:
        return FAIL, stray_handlers[0], stray_handlers
    details.append(f"compile-only ({len(compile_only)}; only what the import allowlist covers is imported, L-075):")
    details.extend(f"  compile-only  {c}" for c in compile_only)
    if printed:
        details.append("printed something at import (harmless, output discarded): " + ", ".join(printed))
    return (PASS, f"{len(files)} files compile; {len(to_import)} modules imported cleanly; {len(compile_only)} compile-only "
                  "(pages / entry points)", details)


# ============================================================================ check 2: database
def check_database() -> tuple[str, str, list[str]]:
    from sqlalchemy import create_engine, text
    from erp.config import DB_HOST, DB_NAME, DB_PORT, database_url

    for ident in DB_VIEWS + DB_TABLES:
        assert re.fullmatch(r"[a-z_]+", ident), ident            # identifiers are constants, never user input
    eng = create_engine(database_url(), connect_args={"connect_timeout": DB_CONNECT_TIMEOUT})
    details: list[str] = []
    try:
        try:
            conn = eng.connect().execution_options(postgresql_readonly=True)   # the whole session is READ ONLY
        except Exception as exc:  # noqa: BLE001
            msg = one_line(getattr(exc, "orig", None) or exc, 150)
            return FAIL, f"cannot connect to PostgreSQL at {DB_HOST}:{DB_PORT}/{DB_NAME}: {msg}", [scrub(str(exc))[:600]]
        with conn:
            ro = conn.execute(text("SHOW transaction_read_only")).scalar()
            if str(ro).lower() != "on":
                return FAIL, "could not put the session in READ ONLY mode - refusing to run any query", []
            version = conn.execute(text("SHOW server_version")).scalar()
            missing: list[str] = []
            for kind, names in (("view", DB_VIEWS), ("table", DB_TABLES)):
                for name in names:
                    try:
                        conn.execute(text(f'SELECT * FROM "{name}" LIMIT 1')).fetchall()
                        details.append(f"{kind} {name}: queryable")
                    except Exception as exc:  # noqa: BLE001
                        conn.rollback()
                        missing.append(f"{kind} {name}")
                        details.append(f"{kind} {name}: {one_line(getattr(exc, 'orig', None) or exc, 140)}")
            conn.rollback()
        if missing:
            return FAIL, f"not queryable: {', '.join(missing)}", details
        return PASS, (f"PostgreSQL {version} at {DB_HOST}:{DB_PORT}/{DB_NAME}, read-only session; "
                      f"{len(DB_VIEWS)} views + {len(DB_TABLES)} tables SELECT fine"), details
    finally:
        eng.dispose()


# ============================================================================ check 3: Data Flow map
def check_flow_map() -> tuple[str, str, list[str]]:
    from desktop import flow_definition as fd
    from desktop.flow_data import check_map

    res = check_map(ROOT)
    if not res.get("available"):
        return FAIL, f"the map check itself is unavailable: {one_line(res.get('error') or '?')}", []
    details: list[str] = []
    for u in res["unmapped"]:
        details.append(f"unmapped script: {u['path']}" + (f" - {u['doc']}" if u.get("doc") else ""))
    for s in res["stale_nodes"]:
        details.append(f"stale node {s['node']}: {s['path']} ({s['reason']})")
    for s in res["stale_ignores"]:
        details.append(f"IGNORE entry {s['path']}: {s['reason']}")
    for i in res["integrity"]:
        details.append(f"inconsistent [{i['kind']}] {i['where']}: {i['message']}")
    if not res["in_sync"]:
        parts = [(len(res["unmapped"]), "unmapped script(s)"), (len(res["stale_nodes"]), "stale node file(s)"),
                 (len(res["stale_ignores"]), "stale IGNORE entr(ies)"), (len(res["integrity"]), "inconsistent edge/node(s)")]
        head = ", ".join(f"{n} {t}" for n, t in parts if n)
        first = details[0] if details else ""
        return FAIL, f"map out of date: {head} - {first}", details + [f"fix: {res.get('hint', '')}"]
    unknown = [e for e in fd.EDGES if e.get("trigger") in ("event", "scheduled", "background") and e.get("armed_by") is None]
    if unknown:   # belt and braces: map_integrity reports the same thing
        return FAIL, f"{len(unknown)} automated edge(s) without armed_by (they would show as 'unknown')", [e["id"] for e in unknown]
    return PASS, (f"map in sync: {res['mapped']} scripts mapped, {len(res['ignored'])} ignored on purpose, "
                  f"{len(fd.NODES)} nodes, {len(fd.EDGES)} edges, no unknown edges"), []


# ============================================================================ check 4: ERP Desk server
_OPENER = urllib.request.build_opener(urllib.request.ProxyHandler({}))      # local requests never go through a proxy


def _http(method: str, url: str, timeout: float = 25) -> tuple[int, dict, bytes]:
    req = urllib.request.Request(url, method=method)
    try:
        with _OPENER.open(req, timeout=timeout) as r:
            return r.status, dict(r.headers), r.read()
    except urllib.error.HTTPError as exc:
        return exc.code, dict(exc.headers or {}), exc.read() or b""


def free_port() -> int:
    for _ in range(20):
        with socket.socket() as s:
            s.bind(("127.0.0.1", 0))
            port = s.getsockname()[1]
        if port not in RESERVED_PORTS:
            return port
    raise RuntimeError("no spare port found")


def port_open(port: int) -> bool:
    with socket.socket() as s:
        s.settimeout(0.5)
        return s.connect_ex(("127.0.0.1", port)) == 0


def _json(body: bytes):
    return json.loads(body.decode("utf-8"))


def check_server(state: dict) -> tuple[str, str, list[str]]:
    from desktop import flow_definition as fd
    from desktop.digest_service import DigestService
    from desktop.flow_data import FlowStore
    from desktop.report_data import ReportStore
    from desktop.server import STATIC_DIR, AppContext, ServerThread, create_app

    # The server's own loggers are noisy on purpose (DB down -> warnings): keep them out of the table.
    captured: list[str] = []

    class _Cap(logging.Handler):
        def emit(self, record):      # noqa: D401
            captured.append(f"{record.levelname} {record.name}: {one_line(record.getMessage(), 160)}")

    desk_log = logging.getLogger("erp_desk")
    cap, old_prop = _Cap(level=logging.WARNING), desk_log.propagate
    desk_log.addHandler(cap)
    desk_log.propagate = False

    port = free_port()
    state["server_port"] = port
    report = ReportStore(3)
    digest = DigestService(ROOT, sys.executable, auto_generate=False)      # never generates, never mails
    flow = FlowStore(ROOT, report=report, digest=digest, interval_seconds=3)
    ctx = AppContext(shell_port=port, streamlit_port=port + 1, report=report, digest=digest, flow=flow)
    srv = ServerThread(create_app(ctx), port)
    details: list[str] = []
    failures: list[str] = []
    base = f"http://127.0.0.1:{port}"
    try:
        report.refresh()          # one synchronous, read-only query round each; the refresher threads are NOT started
        flow.refresh()
        srv.start()
        up = False
        for _ in range(100):
            try:
                st, _h, body = _http("GET", base + "/api/ping", timeout=2)
                if st == 200:
                    ping = _json(body)
                    up = ping.get("app") == "erp-desk" and ping.get("pid") == os.getpid()
                    break
            except (OSError, ValueError):
                pass
            time.sleep(0.1)
        if not up:
            return FAIL, f"ERP Desk did not answer on 127.0.0.1:{port} within 10 s", captured[-3:]

        def expect(path: str, ctype: str, validate=None, method: str = "GET", want: int = 200, label: str | None = None) -> bytes | None:
            t0 = time.monotonic()
            try:
                status, headers, body = _http(method, base + path)
            except Exception as exc:  # noqa: BLE001
                failures.append(f"{method} {path}: {type(exc).__name__}: {one_line(str(exc), 100)}")
                return None
            ok, why = status == want, ""
            if not ok:
                why = f"HTTP {status}, expected {want}"
            elif ctype and ctype not in headers.get("content-type", headers.get("Content-Type", "")):
                ok, why = False, f"content-type {headers.get('content-type', headers.get('Content-Type', '?'))}, expected {ctype}"
            elif want == 200 and not body:
                ok, why = False, "empty body"
            elif validate:
                try:
                    problem = validate(body)
                except Exception as exc:  # noqa: BLE001
                    problem = f"{type(exc).__name__}: {one_line(str(exc), 100)}"
                if problem:
                    ok, why = False, problem
            details.append(f"{'ok  ' if ok else 'FAIL'} {method:4} {path:34} {status} {(time.monotonic() - t0) * 1000:5.0f} ms  {label or ''}{('  <- ' + why) if why else ''}")
            if not ok:
                failures.append(f"{method} {path}: {why}")
            return body if ok else None

        def v_shell(b: bytes):
            if b"__TOKEN__" in b or b"__STREAMLIT_URL__" in b:
                return "placeholders were not substituted"
            if b"ERP Desk" not in b:
                return "not the ERP Desk shell"
            return None

        def v_status(b: bytes):
            d = _json(b)
            return None if d.get("app") == "erp-desk" and isinstance(d.get("ports"), dict) else "unexpected /api/status payload"

        db_ok = bool(state.get("db_ok"))

        def v_report(b: bytes):
            d = _json(b)
            if not isinstance(d, dict) or "has_data" not in d:
                return "unexpected payload (no has_data)"
            if db_ok and not d["has_data"]:
                return f"the database is reachable but the report has no data ({one_line(d.get('error') or '?', 90)})"
            return None

        def v_flow(b: bytes):
            d = _json(b)
            data = d.get("data")
            if not isinstance(data, dict):
                return f"no data ({one_line(d.get('error') or 'feed not ready', 90)})"
            if len(data.get("nodes", [])) != len(fd.NODES) or len(data.get("edges", [])) != len(fd.EDGES):
                return "the page shows a different number of nodes / edges than flow_definition.py declares"
            undeclared = [e["id"] for e in data["edges"] if str(e.get("mode_why", "")).startswith("No armed_by declared")]
            if undeclared:
                return f"edges without armed_by (would show as unknown): {', '.join(undeclared[:4])}"
            mc = data.get("map_check") or {}
            if mc.get("available") and mc.get("in_sync") is False:
                return "the Data Flow page itself reports 'The map is out of date'"
            if db_ok and not (data.get("db") or {}).get("ok"):
                return "the Data Flow feed cannot read the database although the smoke check could"
            return None

        TODAY_KPIS = ["new_leads", "overdue_leads", "open_tickets", "tickets_today", "awaiting", "last_job"]
        TODAY_HEALTH = ["database", "webhook", "clickup", "deepseek", "email", "jobs"]

        def v_today(b: bytes):
            """/api/today: the Today landing page's feed. Shape first, then (only when the DB is reachable) that nothing degraded."""
            d = _json(b)
            if not isinstance(d, dict):
                return "not an object"
            miss = [k for k in ("ok", "generated_at", "day_label", "status", "kpis", "attention", "health", "sources", "rules") if k not in d]
            if miss:
                return f"missing keys: {', '.join(miss)}"
            st = d["status"]
            if st.get("level") not in ("ok", "attention", "watch", "unknown") or not str(st.get("headline", "")).strip():
                return "status has no level / headline sentence"
            for k in ("caveat", "blind", "sep"):
                if k not in st:
                    return f"status has no '{k}' field (the honesty sub-line, lesson L-084)"
            if st["caveat"] is not None and not str(st["caveat"]).strip():
                return "status.caveat is an empty string (must be null or a sentence)"
            if st["level"] == "ok" and (st["caveat"] or st["blind"]):
                return "level 'ok' ('All good') together with a blind input"
            if st["level"] != "unknown" and d["health"]:
                cells = {c.get("id"): c for c in d["health"]}
                amber = [n for n, bad in (("webhook", cells.get("webhook", {}).get("state") in ("warn", "unknown")),
                                          ("pull job", cells.get("jobs", {}).get("pull", "unreadable") != "scheduled"),
                                          ("ClickUp", cells.get("clickup", {}).get("state") == "off")) if bad]
                if amber and not st["caveat"]:
                    return f"an input is blind ({', '.join(amber)}) but status.caveat is empty: the headline would reassure about data that cannot arrive"
                if amber and "All good" in st["headline"]:
                    return "headline says 'All good' although an input is blind"
            if [t.get("id") for t in d["kpis"]] != TODAY_KPIS and not (st["level"] == "unknown" and not d["kpis"]):
                return f"KPI tiles are {[t.get('id') for t in d['kpis']]}, expected {TODAY_KPIS}"
            for t in d["kpis"]:
                if t.get("state") not in ("ok", "unavailable") or not t.get("label"):
                    return f"tile {t.get('id')}: bad state / label"
                if t["state"] == "ok" and t.get("value") is None and t.get("text") is None:
                    return f"tile {t.get('id')} is 'ok' but has no value"
            att = d["attention"]
            if not isinstance(att.get("items"), list) or "available" not in att or "total" not in att:
                return "attention list has the wrong shape"
            if "all_same_what" not in att:
                return "the attention list has no 'all_same_what' field (the hoisted shared phrase, lesson L-145)"
            whats = [i.get("what") for i in att["items"]]
            want_same = whats[0] if len(whats) > 1 and len(set(whats)) == 1 else None
            if att["all_same_what"] != want_same:
                return f"attention.all_same_what is {att['all_same_what']!r} but the rows say {want_same!r} (L-145)"
            if [c.get("id") for c in d["health"]] != TODAY_HEALTH and d["health"]:
                return f"health cells are {[c.get('id') for c in d['health']]}, expected {TODAY_HEALTH}"
            if any(c.get("state") not in ("ok", "warn", "off", "bad", "unknown") for c in d["health"]):
                return "a health cell has an unknown state"
            if db_ok:
                if st["level"] == "unknown" or d["sources"].get("leads") != "ok" or d["sources"].get("tickets") != "ok":
                    return "the database is reachable but the Today feed reports its numbers as unavailable"
                if not att["available"]:
                    return "the database is reachable but the attention list is unavailable"
            return None

        today_seen: dict = {}
        _v_today_plain = v_today

        def v_today_keep(b: bytes):
            problem = _v_today_plain(b)
            try:
                today_seen["d"] = _json(b)
            except ValueError:
                pass
            return problem

        HEALTH_CARDS = ["database", "webhook", "clickup", "deepseek", "email",
                        "job_clickup_pull", "job_digest_daily", "job_weekly_report",
                        "desk_report", "desk_flow", "desk_digest", "script_center"]
        # Settings whose VALUE may never appear in a payload. Checked against the machine's real .env, so this is a
        # live proof, not a re-run of the scenario table's sentinel check (tests/health_scenarios.py).
        HEALTH_SECRET_SETTINGS = ["CLICKUP_API_TOKEN", "CLICKUP_LIST_ID", "DEEPSEEK_API_KEY", "SMTP_HOST", "SMTP_USER",
                                  "SMTP_PASSWORD", "ALERT_EMAIL_FROM", "ALERT_EMAIL_TO", "DB_PASSWORD"]

        def v_health(b: bytes):
            """/api/health: the Health page's feed. Shape, vocabulary, fix hints, no secrets, and agreement with Today."""
            from desktop import health_data as hd
            d = _json(b)
            if not isinstance(d, dict):
                return "not an object"
            miss = [k for k in ("ok", "generated_at", "headline", "today", "groups", "cards", "numbers", "config",
                                "never_runs", "rules", "sources") if k not in d]
            if miss:
                return f"missing keys: {', '.join(miss)}"
            h = d["headline"]
            if h.get("level") not in ("ok", "watch", "attention", "unknown") or not str(h.get("text", "")).strip():
                return "the headline has no level / sentence"
            cards = d["cards"]
            ids = [c.get("id") for c in cards]
            if len(ids) != len(set(ids)):
                return "two cards share an id"
            gone = [i for i in HEALTH_CARDS if i not in ids]
            if gone:
                return f"cards missing: {', '.join(gone)}"
            groups = {g.get("id") for g in d["groups"]}
            listed = [i for g in d["groups"] for i in g.get("cards", [])]
            if sorted(listed) != sorted(ids):
                return "a card is not rendered in exactly one group"
            for c in cards:
                if c.get("state") not in hd.STATES or c.get("word") not in hd.WORDS:
                    return f"card {c.get('id')}: state {c.get('state')!r} / word {c.get('word')!r} is outside the vocabulary"
                if c.get("group") not in groups:
                    return f"card {c.get('id')} is in group {c.get('group')!r}, which is not rendered"
                if not str(c.get("why", "")).strip() or not str(c.get("impact", "")).strip():
                    return f"card {c.get('id')} does not say why, or what stops working without it"
                if not c.get("fix") or not any(f.get("copy") for f in c["fix"]):
                    return f"card {c.get('id')} has no copyable fix hint"
                if any(f.get("kind") not in ("command", "click", "config", "owner") or not str(f.get("text", "")).strip()
                       for f in c["fix"]):
                    return f"card {c.get('id')} has a fix hint with an unknown kind or no text"
                if not c.get("checked_text") or not c.get("ok_text"):
                    return f"card {c.get('id')} does not say when it was last checked / last succeeded"
                if c.get("impact_title") != hd._impact_title(c["state"]):
                    return f"card {c.get('id')} heads its consequence block {c.get('impact_title')!r}, which does not match its state"
                if c.get("tag") not in (None, hd._TAG.get(c["state"])):
                    return f"card {c.get('id')} carries the severity badge {c.get('tag')!r}, which is outside the vocabulary"
            n = d["numbers"]
            if "complete" not in n or not isinstance(n.get("items"), list) or not str(n.get("text", "")).strip():
                return "the 'what this means for the numbers' block has the wrong shape"
            if n["items"] and any(not it.get("clause") for it in n["items"]):
                return "a blind input is listed without saying what it is"
            cfg = d["config"]
            if cfg.get("available"):
                for r in cfg["rows"]:
                    if r.get("state") not in ("set", "not set"):
                        return f"configuration row {r.get('key')} has state {r.get('state')!r}, expected 'set' / 'not set'"
                    if set(r) - {"key", "label", "integration", "why", "state"}:
                        return f"configuration row {r.get('key')} carries a field that could hold a value"
                for i in cfg["integrations"]:
                    if i.get("mode") not in ("configured", "partly configured", "mocked", "not set up"):
                        return f"integration {i.get('id')} has mode {i.get('mode')!r}"
            if "never runs anything" not in str(d["never_runs"]).lower():
                return "the payload does not state that the page runs nothing"
            # no configuration VALUE may appear anywhere in the response
            try:
                from erp import config as _cfg
                low = b.lower()
                leaked = [k for k in HEALTH_SECRET_SETTINGS
                          if len(str(getattr(_cfg, k, "") or "").strip()) >= 4
                          and str(getattr(_cfg, k)).strip().lower().encode() in low]
                if leaked:
                    return f"{len(leaked)} setting value(s) from .env appear in the payload"   # never name or print them
            except Exception as exc:  # noqa: BLE001
                return f"the no-secrets check could not read erp.config: {type(exc).__name__}"
            # the Today strip is the summary of these very cards: they may not disagree
            t = today_seen.get("d")
            if isinstance(t, dict) and t.get("health"):
                cells = {c.get("id"): c.get("state") for c in t["health"]}
                by = {c["id"]: c for c in cards}
                for spec in hd.SUMMARY_CARDS:
                    want, got = cells.get(spec["cell"]), by[spec["id"]]["state"]
                    if want is not None and want != got:
                        return (f"the Health card {spec['id']} says {got!r} but the Today health strip says {want!r} "
                                "about the same dependency")
                # Only ONE direction is an invariant, and it is the one health_data.mark_attention() really enforces:
                # a green Health page may never sit above a Today page that admits a blind input. The other direction
                # is NOT a rule and must not be asserted: the Health page counts every recommended job that is not in
                # Task Scheduler (the daily digest, the weekly report), while today_data.blind_inputs() only knows the
                # webhook, ClickUp and the comment pull - so a correct machine with the pull scheduled and the webhook
                # running legitimately reads "All good" on Today and "2 things need attention" here. The two pages
                # answer different questions (machinery vs. workload); what they may not do is contradict each other
                # about the SAME dependency, which the per-cell loop above already checks.
                if h["level"] == "ok" and t.get("status", {}).get("blind"):
                    return "the Health page says all clear although the Today page reports a blind input"
            return None

        LEADS_PANELS = ["funnel", "sla", "ttfr", "trend", "by_rep", "by_source", "table"]
        LEADS_KPIS = ["leads", "rate", "past", "waiting", "median"]

        def v_leads(b: bytes):
            """/api/leads/analysis: the Leads page's feed. Shape, internal consistency, agreement with the Today page."""
            from desktop import today_data as td
            d = _json(b)
            if not isinstance(d, dict):
                return "not an object"
            miss = [k for k in ("ok", "generated_at", "filters", "options", "panels", "kpis", "definitions", "export_url", "config",
                                "caveat", "warnings", "page") if k not in d]
            if miss:
                return f"missing keys: {', '.join(miss)}"
            if sorted(d["panels"]) != sorted(LEADS_PANELS):
                return f"panels are {sorted(d['panels'])}, expected {sorted(LEADS_PANELS)}"
            for k, pnl in d["panels"].items():
                if pnl.get("state") not in ("ok", "empty", "unavailable"):
                    return f"panel {k} has state {pnl.get('state')!r}"
                if pnl["state"] == "unavailable" and not pnl.get("error"):
                    return f"panel {k} is unavailable without saying why"
            if [t.get("id") for t in d["kpis"]] != LEADS_KPIS:
                return f"KPI tiles are {[t.get('id') for t in d['kpis']]}, expected {LEADS_KPIS}"
            for t in d["kpis"]:
                if t.get("state") not in ("ok", "unavailable") or not t.get("label") or not t.get("hint"):
                    return f"tile {t.get('id')}: bad state / label / hint"
                if t["state"] == "ok" and t.get("value") is None and t.get("text") is None:
                    return f"tile {t.get('id')} is 'ok' but has no value"
            if not str(d["export_url"]).startswith("/api/leads/export.csv"):
                return "export_url does not point at the CSV endpoint"
            defs = d["definitions"]
            if not defs or any(not x.get("term") or not x.get("text") for x in defs):
                return "the Definitions drawer text is empty"
            joined = " ".join(x["text"] for x in defs)
            if f"{td._sla_hours():g} business hours" not in joined:
                return "the Definitions text does not quote the configured SLA hours (LEAD_SLA_HOURS)"
            opt = d["options"]
            if not isinstance(opt.get("reps"), list) or not isinstance(opt.get("sources"), list) or len(opt.get("statuses", [])) != 5:
                return "options (reps / sources / statuses) have the wrong shape"
            sla, fun, tbl = d["panels"]["sla"], d["panels"]["funnel"], d["panels"]["table"]
            if sla["state"] == "ok":
                c = sla["counts"]
                if sum(c.values()) != sla["total"]:
                    return "SLA outcome counts do not add up to the number of leads in view"
                for r in [sla["on_time_rate"], sla["breached_rate"]] + [s["rate"] for s in sla["shares"]]:
                    if r["pct"] is not None and (r["of"] < d["config"]["low_n"] or not 0 <= r["pct"] <= 100):
                        return f"a percentage is shown for {r['of']} lead(s) (low-n rule) or is out of range: {r}"
            if fun["state"] == "ok":
                counts = [s["count"] for s in fun["stages"]]
                if counts != sorted(counts, reverse=True) or counts[0] != fun["total"]:
                    return f"funnel counts {counts} are not a non-increasing funnel from the total {fun['total']}"
                if sla["state"] == "ok" and fun["total"] != sla["total"]:
                    return "funnel total and SLA total disagree"
            if tbl["state"] == "ok" and sla["state"] == "ok" and tbl["total"] != sla["total"]:
                return "table total and SLA total disagree (the panels do not share one filter)"
            if db_ok:
                if not d["ok"] or any(v["state"] == "unavailable" for v in d["panels"].values()):
                    return "the database is reachable but a Leads panel is unavailable: " + ", ".join(k for k, v in d["panels"].items() if v["state"] == "unavailable")
                if d["filters"]["active"] != 0:
                    return "an unfiltered request came back with active filters"
                if sla["total"] != opt["total"]:
                    return f"no filter set but {sla['total']} leads in view of {opt['total']}"
            t = today_seen.get("d")                     # the two pages must never disagree about who is late (they share one SQL helper)
            if isinstance(t, dict) and db_ok and sla["state"] == "ok":
                tile = next((k for k in t.get("kpis", []) if k.get("id") == "overdue_leads"), None)
                if tile and tile.get("state") == "ok" and tile.get("value") != sla["past_sla_now"]:
                    return f"'Past SLA now' is {sla['past_sla_now']} here but the Today page says {tile.get('value')}"
                if t.get("status", {}).get("level") != "unknown" and (t["status"].get("caveat") is None) != (d["caveat"] is None):
                    return "the Leads page and the Today page disagree about whether the numbers may be incomplete (caveat)"
            return None

        def v_leads_400(b: bytes):
            d = _json(b)
            if d.get("ok") is not False or not d.get("problems") or "panels" in d:
                return "a hostile filter must come back as an error with a 'problems' list and no data"
            return None

        def v_health_400(b: bytes):
            d = _json(b)
            if d.get("ok") is not False or not d.get("problems") or "cards" in d:
                return "an unknown parameter must come back as an error with a 'problems' list and no cards"
            if d["problems"][0].get("param") != "bogus":
                return f"the problem does not name the offending parameter: {d['problems']}"
            return None

        def v_leads_csv(b: bytes):
            if b[:3] != b"\xef\xbb\xbf":
                return "the CSV does not start with a UTF-8 byte-order mark"
            text_ = b[3:].decode("utf-8")
            lines = text_.split("\r\n")
            head = lines[0].split(",")
            if head[:3] != ["lead_id", "created_at_utc", "lead_name"] or len(head) != 16:
                return f"unexpected CSV header: {lines[0][:80]}"
            low = lines[0].lower()
            if any(w in low for w in ("email", "phone", "payload", "token", "password", "secret")):
                return "the CSV header exposes a sensitive column"
            import csv as _csv
            import io as _io
            rows = list(_csv.reader(_io.StringIO(text_, newline="")))
            if any(len(r) != len(head) for r in rows if r):
                return "a CSV row has a different number of cells than the header"
            for r in rows[1:]:
                for cell in r:
                    if cell[:1] in ("=", "+", "@", "\t", "\r", ";"):
                        return f"a CSV cell starts with a formula character: {cell[:20]!r}"
            if db_ok and len(rows) < 1:
                return "empty CSV"
            return None

        CHANNEL_STATES = ("not_installed", "empty", "no_match", "ready")
        channels_seen: dict = {}

        def v_channels(b: bytes):
            """/api/channels: the Channels page's feed (marketing channel performance). Shape and honesty in WHATEVER state the
            real database is in - today the rollup tables are not installed, so the normal answer is a setup card with the exact
            fix and not one number; once installed and loaded it must add up and never show a percentage without a real n."""
            d = _json(b)
            if not isinstance(d, dict):
                return "not an object"
            miss = [k for k in ("ok", "state", "generated_at", "time_zone", "schema", "headline", "kpis", "panels", "definitions", "filters", "sort",
                                "sorts", "options", "read", "freshness", "config", "export", "install", "warnings", "money") if k not in d]
            if miss:
                return f"missing keys: {', '.join(miss)}"
            channels_seen["d"] = d
            if d["state"] not in CHANNEL_STATES:
                return f"unknown page state {d['state']!r}"
            if d["read"].get("raw_fact_read") is not False:
                return "the page does not state that it never reads the raw fact table"
            head = d["headline"]
            if not str(head.get("text") or "").strip() or head.get("tone") not in ("grey", "blue", "amber", "red"):
                return f"the headline sentence is missing or has no tone: {head}"
            if not d["definitions"] or any(not x.get("term") or not x.get("text") for x in d["definitions"]):
                return "the Definitions drawer text is empty"
            joined = " ".join(x["text"] for x in d["definitions"])
            if "not scheduled unless the owner schedules it" not in joined or "python -m erp.marketing.rollup" not in joined:
                return "the Definitions text does not say how the rollup is refreshed (and that it is not scheduled)"
            if d["time_zone"] and d["time_zone"] not in joined:
                return "the Definitions text does not name the database time zone (L-097)"
            if {s["value"] for s in d["sorts"]} != {"sessions", "channel", "clicks", "impressions", "ctr", "conversions", "cvr", "spend", "revenue", "cpa"}:
                return f"the sort vocabulary changed: {[s['value'] for s in d['sorts']]}"
            if d["state"] in ("not_installed", "empty"):
                if d["kpis"] or d["panels"] or d["export"] or d["range"] is not None:
                    return f"a {d['state']} answer must carry no tile, no panel, no export and no range - it has data-shaped fields"
                if re.search(r"\d", head["text"] + head.get("sub", "")):
                    return "a setup headline contains a number: nothing may be shown that does not exist"
                steps = d["install"]["steps"] if d["install"] else []
                cmds = [s.get("command", "") for s in steps]
                if d["state"] == "not_installed" and not (d["install"]["missing"] and any("09_marketing_rollup.sql" in c or "07_marketing_schema.sql" in c for c in cmds)):
                    return f"not installed, but the exact psql command for the missing SQL file is not offered: {cmds}"
                if d["state"] == "empty" and not any("ingest" in c for c in cmds):
                    return "empty, but the load command is not offered"
                if d["state"] == "not_installed" and d["install"]["missing"] and any("rollup" in m for m in d["install"]["missing"]):
                    p10 = next((i for i, c in enumerate(cmds) if "10_marketing_currency.sql" in c), None)
                    p09 = next((i for i, c in enumerate(cmds) if "09_marketing_rollup.sql" in c), None)
                    if p10 is None or p09 is None or p10 > p09:
                        return f"the currency migration (db/sql/10) must be offered BEFORE the rollup tables (db/sql/09): {cmds}"
                if d["state"] == "not_installed" and d["install"]["missing"] and any("rollup" in m for m in d["install"]["missing"]) \
                        and not any("psql.exe" in c and "09_marketing_rollup.sql" in c for c in cmds):
                    return "the rollup tables are missing but the owner-run psql command for db/sql/09 is not shown"
                if not d["ok"]:
                    return "a setup state is not a failure: ok must be true"
                return None
            if db_ok and not d["ok"]:
                return "the database is reachable but a Channels panel or tile is unavailable: " + ", ".join(k for k, v in d["panels"].items() if v["state"] == "unavailable")
            if d["state"] == "ready":
                low_n = d["config"]["low_n"]
                if [t["id"] for t in d["kpis"]] != ["sessions", "clicks", "conversions", "cvr", "spend", "revenue"]:
                    return f"unexpected tiles: {[t['id'] for t in d['kpis']]}"
                tb = d["panels"]["table"]
                if tb["state"] == "ok":
                    if tb["total"]["sessions"] != sum(r["sessions"] for r in tb["rows"]) and not tb["more"]:
                        return "the table total is not the sum of its rows"
                    tile = next(t for t in d["kpis"] if t["id"] == "sessions")
                    if tile["state"] == "ok" and tile["value"] != tb["total"]["sessions"] and not tb["more"]:
                        return f"the sessions tile ({tile['value']}) and the table total ({tb['total']['sessions']}) disagree"
                    for r in tb["rows"] + [tb["total"]]:
                        for k in ("ctr", "cvr", "share"):
                            rt = r[k]
                            if rt["pct"] is not None and (rt["of"] < low_n or not 0 <= rt["pct"] <= 100):
                                return f"{r['name']}: a {k} percentage is shown for a denominator of {rt['of']} (low-n {low_n})"
                        if r["spend"]["state"] == "none" and r["spend"]["value"] is not None:
                            return f"{r['name']}: spend is 'none' but carries a value"
                        for k in ("spend", "revenue", "cpa"):        # money is never added across currencies: a multi-currency cell has parts and NO value
                            cell = r[k]
                            if cell["state"] == "multi" and (cell["value"] is not None or len(cell["parts"]) < 2):
                                return f"{r['name']}: a {k} cell that spans currencies must have no single value and one part per currency: {cell}"
                            if cell["state"] == "ok" and not cell.get("currency"):
                                return f"{r['name']}: a {k} amount is shown without its currency"
                    for t in d["kpis"]:
                        if t["state"] == "multi" and t["value"] is not None:
                            return f"the {t['id']} tile spans currencies but carries a single value"
                        if t["id"] in ("spend", "revenue") and t["state"] == "ok" and not t.get("unit"):
                            return f"the {t['id']} tile shows an amount without its currency"
                if not d["range"] or not d["range"]["from"] or not d["time_zone"]:
                    return "a ready payload must carry its date range and the time zone (L-097)"
            return None

        def v_channels_400(b: bytes):
            d = _json(b)
            if d.get("ok") is not False or not d.get("problems") or "panels" in d or "kpis" in d:
                return "a bad parameter must come back as an error with a 'problems' list and no data"
            return None

        def v_channels_409(b: bytes):
            d = _json(b)
            if d.get("ok") is not False or d.get("state") not in ("not_installed", "empty"):
                return f"an export with nothing to export must say why: {d}"
            return None

        def v_channels_csv(b: bytes):
            if b[:3] != b"\xef\xbb\xbf":
                return "the CSV does not start with a UTF-8 byte-order mark"
            text_ = b[3:].decode("utf-8")
            if "\r\n" not in text_:
                return "the CSV does not use CRLF line ends"
            import csv as _csv
            import io as _io
            rows = list(_csv.reader(_io.StringIO(text_, newline="")))
            head = rows[0]
            if any(re.search(r"e-?mail|phone|identity|user|lead|payload|token|password|secret", h, re.I) for h in head) or "time_zone" not in head:
                return f"the CSV header exposes an identity column or omits the time zone: {head}"
            if any(len(r) != len(head) for r in rows if r):
                return "a CSV row has a different number of cells than the header"
            for r in rows[1:]:
                for cell in r:
                    if cell[:1] in ("=", "+", "@", "\t", "\r", ";"):
                        return f"a CSV cell starts with a formula character: {cell[:20]!r}"
            return None

        leads_seen: dict = {}
        _v_leads_plain = v_leads

        def v_leads_keep(b: bytes):
            problem = _v_leads_plain(b)
            try:
                leads_seen["d"] = _json(b)
            except ValueError:
                pass
            return problem

        SOURCES_PANELS = ["cohorts", "sources"]
        SOURCES_KPIS = ["sources", "top", "cohorts", "replied", "past"]

        def v_sources(b: bytes):
            """/api/leads/sources: the Leads page's "Sources & cohorts" tab. Shape, the low-n and cohort rules that make it
            honest, and agreement with the other two pages about who is past SLA."""
            from desktop import sources_data as sx
            from desktop import today_data as td
            d = _json(b)
            if not isinstance(d, dict):
                return "not an object"
            miss = [k for k in ("ok", "generated_at", "time_zone", "filters", "view", "options", "panels", "kpis",
                                "definitions", "export_url", "export_cohort_url", "config", "caveat", "warnings", "banner") if k not in d]
            if miss:
                return f"missing keys: {', '.join(miss)}"
            if sorted(d["panels"]) != SOURCES_PANELS:
                return f"panels are {sorted(d['panels'])}, expected {SOURCES_PANELS}"
            for k, pnl in d["panels"].items():
                if pnl.get("state") not in ("ok", "empty", "unavailable", "too_big"):
                    return f"panel {k} has state {pnl.get('state')!r}"
                if pnl["state"] in ("unavailable", "too_big") and not pnl.get("error"):
                    return f"panel {k} is {pnl['state']} without saying why"
            if [t.get("id") for t in d["kpis"]] != SOURCES_KPIS:
                return f"KPI tiles are {[t.get('id') for t in d['kpis']]}, expected {SOURCES_KPIS}"
            for t in d["kpis"]:
                if t.get("state") not in ("ok", "unavailable") or not t.get("label") or not t.get("hint"):
                    return f"tile {t.get('id')}: bad state / label / hint"
                if t["state"] == "ok" and t.get("value") is None and t.get("text") is None:
                    return f"tile {t.get('id')} is 'ok' but has no value"
            if not str(d["export_url"]).startswith("/api/leads/sources.csv") or "part=sources" not in d["export_url"]:
                return "export_url does not point at this tab's CSV endpoint"
            if "part=cohorts" not in str(d["export_cohort_url"]):
                return "export_cohort_url does not point at the cohort CSV"
            view = d["view"]
            if view.get("bucket") not in sx.BUCKET_KEYS or view.get("metric") not in sx.METRIC_KEYS:
                return f"view bucket / metric are not whitelisted values: {view}"
            if not str(view.get("bucket_why") or "").strip():
                return "the page does not say WHY it grouped the leads the way it did"
            defs = d["definitions"]
            joined = " ".join(x.get("text", "") for x in defs)
            if not defs or any(not x.get("term") or not x.get("text") for x in defs):
                return "the Definitions drawer text is empty"
            if f"{td._sla_hours():g} business hours" not in joined:
                return "the Definitions text does not quote the configured SLA hours (LEAD_SLA_HOURS)"
            if d["time_zone"] and d["time_zone"] not in joined:
                return "the Definitions text does not name the time zone the day buckets are built in (L-097)"
            low_n = d["config"]["low_n"]
            coh = d["panels"]["cohorts"]
            if coh["state"] == "ok":
                if coh["y_mode"] == "pct" and any(c["n"] < low_n for c in coh["cohorts"]):
                    return "the chart shows percentages although a cohort has fewer than the low-n threshold"
                for c in coh["cohorts"]:
                    cums = [p["cum"] for p in c["points"]]
                    if any(b_ < a for a, b_ in zip(cums, cums[1:])):
                        return f"cohort {c['key']}: a cumulative curve goes down"
                    if cums and cums[-1] > c["n"]:
                        return f"cohort {c['key']}: more leads reached the metric than the cohort holds"
                    if c["single_point"] != (len(c["points"]) < 2):
                        return f"cohort {c['key']}: single_point does not match the number of points (a line through one point)"
                    if any(p["rate"]["pct"] is not None for p in c["points"]) and c["n"] < low_n:
                        return f"cohort {c['key']}: a percentage is shown for {c['n']} lead(s)"
                    if any(p["n"] != c["n"] for p in c["points"]):
                        return f"cohort {c['key']}: the denominator changes along the curve"
            src = d["panels"]["sources"]
            if src["state"] == "ok":
                for r in src["rows"]:
                    stages = [r["stages"][k]["n"] for k in ("assigned", "analyzed", "task", "update")]
                    if stages != sorted(stages, reverse=True):
                        return f"source {r['name']}: the strict stage counts {stages} are not a funnel"
                    if sum(r["counts"].values()) != r["n"]:
                        return f"source {r['name']}: the SLA outcomes do not add up to its leads"
                    for rt in [r["share"], r["on_time_rate"]] + [r["stages"][k] for k in r["stages"]]:
                        if rt["pct"] is not None and (rt["of"] < low_n or not 0 <= rt["pct"] <= 100):
                            return f"source {r['name']}: a percentage is shown for {rt['of']} lead(s) or is out of range"
            if db_ok:
                if not d["ok"] or any(v["state"] in ("unavailable", "too_big") for v in d["panels"].values()):
                    return "the database is reachable but a Sources panel is unavailable: " + ", ".join(
                        k for k, v in d["panels"].items() if v["state"] in ("unavailable", "too_big"))
                if d["filters"]["active"] != 0:
                    return "an unfiltered request came back with active filters"
                if not d["time_zone"]:
                    return "the database is reachable but the session time zone was not read (L-097)"
            past = sum(r["past_sla"] for r in src.get("rows", [])) if src["state"] == "ok" else None
            l = leads_seen.get("d")     # the three pages share one "past SLA" definition: they may never disagree (L-102)
            if past is not None and isinstance(l, dict) and db_ok and l.get("panels", {}).get("sla", {}).get("state") == "ok":
                if past != l["panels"]["sla"]["past_sla_now"]:
                    return f"'Past SLA now' is {past} on the Sources tab but {l['panels']['sla']['past_sla_now']} on the Funnel & SLA tab"
                if src["total_leads"] != l["panels"]["sla"]["total"]:
                    return f"the two Leads tabs disagree about how many leads are in view ({src['total_leads']} vs {l['panels']['sla']['total']})"
            t = today_seen.get("d")
            if past is not None and isinstance(t, dict) and db_ok:
                tile = next((k for k in t.get("kpis", []) if k.get("id") == "overdue_leads"), None)
                if tile and tile.get("state") == "ok" and tile.get("value") != past:
                    return f"'Past SLA now' is {past} here but the Today page says {tile.get('value')}"
            return None

        def v_sources_csv(b: bytes):
            if b[:3] != b"\xef\xbb\xbf":
                return "the CSV does not start with a UTF-8 byte-order mark"
            text_ = b[3:].decode("utf-8")
            lines = text_.split("\r\n")
            head = lines[0].split(",")
            low = lines[0].lower()
            if any(w in low for w in ("email", "phone", "payload", "token", "password", "secret", "lead_id", "lead_name")):
                return "the CSV header exposes a per-lead or sensitive column"
            if head[0] not in ("source", "cohort_start"):
                return f"unexpected CSV header: {lines[0][:80]}"
            import csv as _csv
            import io as _io
            rows = list(_csv.reader(_io.StringIO(text_, newline="")))
            if any(len(r) != len(head) for r in rows if r):
                return "a CSV row has a different number of cells than the header"
            for r in rows[1:]:
                for cell in r:
                    if cell[:1] in ("=", "+", "@", "\t", "\r", ";"):
                        return f"a CSV cell starts with a formula character: {cell[:20]!r}"
            return None

        # pages (the shell is one document: Today, Reporting, Management, Data Flow and Leads are its views)
        shell = expect("/", "text/html", v_shell, label="app shell (Reporting / Management / Data Flow views)")
        assets = ["/favicon.ico"]
        if shell:
            for m in re.finditer(rb'(?:src|href)="(/(?:static|vendor)/[^"?#]+)"', shell):
                p = m.group(1).decode()
                if p not in assets:
                    assets.append(p)
        else:
            assets += ["/static/shell.js", "/static/report.js", "/static/flow.js", "/static/today.js", "/static/leads.js",
                       "/static/sources.js", "/static/health.js", "/static/channels.js", "/vendor/plotly.min.js"]
        for a in assets:
            expect(a, "", label="asset")
        for js in ("report.js", "flow.js", "today.js", "leads.js", "sources.js", "health.js", "channels.js", "insights.js", "placements.js"):
            if f"/static/{js}" not in assets and (STATIC_DIR / js).is_file():
                failures.append(f"the shell page does not load /static/{js}")
        # JSON feeds
        expect("/api/ping", "application/json")
        expect("/api/status", "application/json", v_status, label="app status")
        expect("/api/report/data", "application/json", v_report, label="Reporting page feed")
        expect("/api/digest", "application/json", lambda b: None if isinstance(_json(b), dict) else "not an object", label="today's briefing feed")
        expect("/api/flow", "application/json", v_flow, label="Data Flow page feed")
        expect("/api/today", "application/json", v_today_keep, label="Today page feed (landing page)")
        expect("/api/health", "application/json", v_health, label="Health page feed (one card per dependency)")
        expect("/api/health?fresh=1", "application/json", v_health, label="Health page feed with ?fresh=1")
        expect("/api/health?bogus=1", "application/json", v_health_400, want=400,
               label="guard: an unknown parameter NAME is refused (L-098)")
        expect("/api/leads/analysis", "application/json", v_leads_keep, label="Leads page feed (funnel + SLA explorer)")
        expect("/api/leads/analysis?rep=x%27%20OR%20%271%27%3D%271&sort=%3BDROP", "application/json", v_leads_400, want=400,
               label="guard: a hostile filter is refused (400)")
        expect("/api/leads/export.csv", "text/csv", v_leads_csv, label="Leads CSV export (BOM, formula-safe)")
        expect("/api/leads/export.csv?status=past_sla&sort=hours&dir=asc", "text/csv", v_leads_csv, label="Leads CSV export with filters")
        expect("/api/leads/export.csv?status=nope", "application/json", v_leads_400, want=400, label="guard: a bad export filter is refused, not ignored")
        expect("/api/leads/analysis?statuss=past_sla", "application/json", v_leads_400, want=400,
               label="guard: an unknown parameter NAME is refused (L-098)")
        expect("/api/leads/sources", "application/json", v_sources, label="Leads page feed (sources + arrival cohorts)")
        expect("/api/leads/sources?bucket=week&metric=past_sla", "application/json", v_sources, label="... with another bucket and metric")
        expect("/api/leads/sources?bucket=hour", "application/json", v_leads_400, want=400, label="guard: an unknown bucket is refused (400)")
        expect("/api/leads/sources?cohort=1999-01-01", "application/json", v_leads_400, want=400, label="guard: a cohort that does not exist is refused")
        expect("/api/leads/sources?sourc=web", "application/json", v_leads_400, want=400, label="guard: an unknown parameter NAME is refused (L-098)")
        expect("/api/leads/sources.csv", "text/csv", v_sources_csv, label="source table CSV (BOM, formula-safe, no per-lead column)")
        expect("/api/leads/sources.csv?part=cohorts", "text/csv", v_sources_csv, label="cohort curves CSV")
        expect("/api/leads/sources.csv?part=everything", "application/json", v_leads_400, want=400, label="guard: a bad export part is refused")
        expect("/api/channels", "application/json", v_channels, label="Channels page feed (rollup report: not installed / empty / populated)")
        expect("/api/channels?fresh=1", "application/json", v_channels, label="Channels page feed with ?fresh=1")
        expect("/api/channels?chanel=facebook", "application/json", v_channels_400, want=400, label="guard: an unknown parameter NAME is refused (L-098)")
        expect("/api/channels?from=2026-02-30&sort=%3BDROP", "application/json", v_channels_400, want=400, label="guard: a hostile value is refused (400)")
        expect("/api/channels?part=daily", "application/json", v_channels_400, want=400, label="guard: `part` belongs to the CSV endpoint only")
        expect("/api/channels.csv?chanel=facebook", "application/json", v_channels_400, want=400, label="guard: an unknown parameter NAME is refused on the CSV too")
        expect("/api/channels.csv?part=everything", "application/json", v_channels_400, want=400, label="guard: a bad CSV part is refused")
        def v_insights(b: bytes):
            d = _json(b)
            if not isinstance(d, dict) or d.get("state") not in ("not_installed", "empty", "no_data_in_window", "ready", "unavailable"):
                return f"unexpected insights answer: {str(d)[:120]}"
            if d["read"].get("raw_fact_read") is not False or not d.get("how", {}).get("thresholds"):
                return "the insights payload must state it never reads the raw facts and carry its thresholds"
            if d["state"] in ("not_installed", "empty") and (d["findings"] or d["automation"]):
                return "a setup state must carry no finding"
            for f in d["findings"]:
                if f["severity"] not in ("info", "warn", "crit") or f["confidence"] not in ("enough data", "thin data") or not f["evidence"] or not f["action"]:
                    return f"a finding lacks severity, confidence, evidence or action: {f.get('id')}"
                if f["confidence"] == "thin data" and f["severity"] != "info":
                    return f"a thin-data finding is louder than info: {f.get('id')}"
            return None

        expect("/api/channels/insights", "application/json", v_insights, label="Channels Insights feed (rules over the rollup)")
        expect("/api/channels/insights?chanel=facebook", "application/json", v_channels_400, want=400, label="guard: an unknown parameter NAME is refused on the insights feed")
        def v_placements(b: bytes):
            d = _json(b)
            if not isinstance(d, dict) or d.get("state") not in ("not_installed", "placements_not_installed", "empty", "no_match", "outside_filter", "ready", "unavailable"):
                return f"unexpected placements answer: {str(d)[:120]}"
            if d["read"].get("raw_fact_read") is not False:
                return "the placements payload must state it never reads the raw facts"
            if d["state"] in ("not_installed", "placements_not_installed", "empty") and (d["table"]["rows"] or d["totals"]):
                return "a setup state must carry no number"
            return None

        placements_body = expect("/api/channels/placements", "application/json", v_placements, label="Channels Placements feed (placement rollup: not installed / empty / populated)")
        expect("/api/channels/placements?chanel=facebook", "application/json", v_channels_400, want=400, label="guard: an unknown parameter NAME is refused on the placements feed")
        expect("/api/channels/placements?psort=nope", "application/json", v_channels_400, want=400, label="guard: a bad placements sort is refused")
        expect("/api/channels/placements.csv?chanel=facebook", "application/json", v_channels_400, want=400, label="guard: an unknown parameter NAME is refused on the placements CSV")
        expect("/api/channels/report.xlsx?chanel=facebook", "application/json", v_channels_400, want=400, label="guard: an unknown parameter NAME is refused on the Excel report")
        ch_state = (channels_seen.get("d") or {}).get("state")
        if ch_state in ("not_installed", "empty"):
            expect("/api/channels.csv", "application/json", v_channels_409, want=409, label=f"Channels CSV with nothing to export ({ch_state}): 409 with the reason")
        elif ch_state:
            expect("/api/channels.csv", "text/csv", v_channels_csv, label="Channels CSV export (BOM, CRLF, formula-safe, no identity column)")
        def v_channels_xlsx(b: bytes):
            if not b.startswith(b"PK"):
                return "the body does not start with the zip magic 'PK'"
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True)
            names = wb.sheetnames
            want_sheets = ["Summary", "Channels", "Campaigns", "Insights", "Definitions"]
            placed = (_json(placements_body) or {}).get("state") == "ready" if placements_body else False
            if placed and "Placements" not in names:
                return f"the placements feed is ready but the workbook has no Placements sheet: {names}"
            if [n for n in names if n != "Placements"] != want_sheets:
                return f"unexpected sheets: {names}"
            wb.close()
            return None

        if ch_state in ("not_installed", "empty"):
            expect("/api/channels/report.xlsx", "application/json", v_channels_409, want=409, label=f"Excel report with nothing to export ({ch_state}): 409 with the reason")
        elif ch_state:
            expect("/api/channels/report.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", v_channels_xlsx, label="Excel report (200, xlsx type, zip magic, opens with openpyxl, the five sheets)")
        # the token guard must still refuse a state-changing call without the token (nothing is changed either way: throw-away context)
        expect("/api/heartbeat", "", method="POST", want=403, label="guard: POST without token is refused")
        # the pure Today builders (headline / caveat / rules text) and the hang-safety scenarios: tests/today_scenarios.py (L-089)
        try:
            from tests import today_scenarios
            rows = today_scenarios.run(db_ok)
            bad = [r for r in rows if not r[1]]
            details.append(f"{'ok  ' if not bad else 'FAIL'} unit today_scenarios: {len(rows) - len(bad)}/{len(rows)} scenarios (headline, caveat, rules text, timeouts)")
            for name, _ok, why in bad:
                details.append(f"FAIL   scenario: {name} <- {one_line(str(why), 120)}")
            if bad:
                failures.append(f"today scenarios: {bad[0][0]} ({one_line(str(bad[0][2]), 90)})" + (f" (+{len(bad) - 1} more)" if len(bad) > 1 else ""))
        except Exception as exc:  # noqa: BLE001
            failures.append(f"today scenarios crashed: {type(exc).__name__}: {one_line(str(exc), 100)}")
        # the Health page's wording / severity / fix-hint mapping and its no-secrets property: tests/health_scenarios.py
        try:
            from tests import health_scenarios
            rows = health_scenarios.run(db_ok)
            bad = [r for r in rows if not r[1]]
            details.append(f"{'ok  ' if not bad else 'FAIL'} unit health_scenarios: {len(rows) - len(bad)}/{len(rows)} scenarios (states, wording, fix hints, no secrets)")
            for name, _ok, why in bad:
                details.append(f"FAIL   scenario: {name} <- {one_line(str(why), 120)}")
            if bad:
                failures.append(f"health scenarios: {bad[0][0]} ({one_line(str(bad[0][2]), 90)})" + (f" (+{len(bad) - 1} more)" if len(bad) > 1 else ""))
        except Exception as exc:  # noqa: BLE001
            failures.append(f"health scenarios crashed: {type(exc).__name__}: {one_line(str(exc), 100)}")
        # the pure Leads builders (filters / whitelists / funnel math / CSV escaping) and the SQL scenarios on a private synthetic
        # session (temporary tables only, nothing written to a real table): tests/leads_scenarios.py
        try:
            from tests import leads_scenarios
            rows = leads_scenarios.run(db_ok)
            bad = [r for r in rows if not r[1]]
            details.append(f"{'ok  ' if not bad else 'FAIL'} unit leads_scenarios: {len(rows) - len(bad)}/{len(rows)} scenarios (filters, whitelists, funnel, SLA, CSV, failures)")
            for name, _ok, why in bad:
                details.append(f"FAIL   scenario: {name} <- {one_line(str(why), 120)}")
            if bad:
                failures.append(f"leads scenarios: {bad[0][0]} ({one_line(str(bad[0][2]), 90)})" + (f" (+{len(bad) - 1} more)" if len(bad) > 1 else ""))
        except Exception as exc:  # noqa: BLE001
            failures.append(f"leads scenarios crashed: {type(exc).__name__}: {one_line(str(exc), 100)}")
        # the pure Sources & cohorts builders (cohort maths, low-n rules, whitelists, CSV) and the SQL scenarios on a
        # private synthetic session with FOUR arrival cohorts: tests/sources_scenarios.py
        try:
            from tests import sources_scenarios
            rows = sources_scenarios.run(db_ok)
            bad = [r for r in rows if not r[1]]
            details.append(f"{'ok  ' if not bad else 'FAIL'} unit sources_scenarios: {len(rows) - len(bad)}/{len(rows)} scenarios (cohort maths, low-n, buckets, whitelists, CSV)")
            for name, _ok, why in bad:
                details.append(f"FAIL   scenario: {name} <- {one_line(str(why), 120)}")
            if bad:
                failures.append(f"sources scenarios: {bad[0][0]} ({one_line(str(bad[0][2]), 90)})" + (f" (+{len(bad) - 1} more)" if len(bad) > 1 else ""))
        except Exception as exc:  # noqa: BLE001
            failures.append(f"sources scenarios crashed: {type(exc).__name__}: {one_line(str(exc), 100)}")
        # the Channels page (marketing rollup report): filter whitelist incl. the CSV, rate maths, the four page states, tiles /
        # panels / CSV, and the SQL scenarios against the real 07 + 10 + 09 DDL in the session's TEMPORARY schema: tests/channels_scenarios.py
        try:
            from tests import channels_scenarios
            rows = channels_scenarios.run(db_ok)
            bad = [r for r in rows if not r[1]]
            details.append(f"{'ok  ' if not bad else 'FAIL'} unit channels_scenarios: {len(rows) - len(bad)}/{len(rows)} scenarios (whitelist, rates, low-n, states, tiles, CSV, temp-schema SQL, failures)")
            for name, _ok, why in bad:
                details.append(f"FAIL   scenario: {name} <- {one_line(str(why), 120)}")
            if bad:
                failures.append(f"channels scenarios: {bad[0][0]} ({one_line(str(bad[0][2]), 90)})" + (f" (+{len(bad) - 1} more)" if len(bad) > 1 else ""))
        except Exception as exc:  # noqa: BLE001
            failures.append(f"channels scenarios crashed: {type(exc).__name__}: {one_line(str(exc), 100)}")
        # the Channels Insights engine + Excel report: rules (fire / quiet / thin), window maths, per-currency isolation, xlsx contents,
        # formula guard, states, all against the temporary schema: tests/insights_scenarios.py
        try:
            from tests import insights_scenarios
            rows = insights_scenarios.run(db_ok)
            bad = [r for r in rows if not r[1]]
            details.append(f"{'ok  ' if not bad else 'FAIL'} unit insights_scenarios: {len(rows) - len(bad)}/{len(rows)} scenarios (six rules, thresholds, windows, currencies, endpoints, xlsx)")
            for name, _ok, why in bad:
                details.append(f"FAIL   scenario: {name} <- {one_line(str(why), 120)}")
            if bad:
                failures.append(f"insights scenarios: {bad[0][0]} ({one_line(str(bad[0][2]), 90)})" + (f" (+{len(bad) - 1} more)" if len(bad) > 1 else ""))
        except Exception as exc:  # noqa: BLE001
            failures.append(f"insights scenarios crashed: {type(exc).__name__}: {one_line(str(exc), 100)}")
        try:
            from tests import placement_scenarios
            rows = placement_scenarios.run(db_ok)
            bad = [r for r in rows if not r[1]]
            details.append(f"{'ok  ' if not bad else 'FAIL'} unit placement_scenarios: {len(rows) - len(bad)}/{len(rows)} scenarios (connector, rules, rollup, page, csv, xlsx, autorun end to end)")
            for name, _ok, why in bad:
                details.append(f"FAIL   scenario: {name} <- {one_line(str(why), 120)}")
            if bad:
                failures.append(f"placement scenarios: {bad[0][0]} ({one_line(str(bad[0][2]), 90)})" + (f" (+{len(bad) - 1} more)" if len(bad) > 1 else ""))
        except Exception as exc:  # noqa: BLE001
            failures.append(f"placement scenarios crashed: {type(exc).__name__}: {one_line(str(exc), 100)}")
        # the marketing ingestion pipeline's pure logic (typing, validation, dedupe, malformed-row handling,
        # the scoped GIN allowlist, the flat-file connector, DDL/code agreement): tests/marketing_scenarios.py.
        # Needs no database and no server (db_ok is accepted for a uniform signature only) - the scale
        # verification itself is a one-off run of erp/marketing/perf_check.py, not a gate check (see its docstring).
        try:
            from tests import marketing_scenarios
            rows = marketing_scenarios.run(db_ok)
            bad = [r for r in rows if not r[1]]
            details.append(f"{'ok  ' if not bad else 'FAIL'} unit marketing_scenarios: {len(rows) - len(bad)}/{len(rows)} scenarios (typing, validation, dedupe, malformed rows, GIN allowlist, connector, DDL agreement)")
            for name, _ok, why in bad:
                details.append(f"FAIL   scenario: {name} <- {one_line(str(why), 120)}")
            if bad:
                failures.append(f"marketing scenarios: {bad[0][0]} ({one_line(str(bad[0][2]), 90)})" + (f" (+{len(bad) - 1} more)" if len(bad) > 1 else ""))
        except Exception as exc:  # noqa: BLE001
            failures.append(f"marketing scenarios crashed: {type(exc).__name__}: {one_line(str(exc), 100)}")
        if failures:
            return FAIL, f"{len(failures)} request(s) failed: {failures[0]}" + (" (+more)" if len(failures) > 1 else ""), details + captured[-3:]
        n_req = len([d for d in details if d.startswith("ok")])
        return PASS, f"{n_req} requests OK on 127.0.0.1:{port} (shell, {len(assets)} assets, Today + Reporting + Data Flow + Leads + Sources + Health + Channels feeds); server stopped", details
    finally:
        srv.stop()
        flow.stop()
        report.stop()
        try:
            from erp.db import engine
            engine.dispose()              # close the pooled connections the stores opened
        except Exception:  # noqa: BLE001
            pass
        desk_log.removeHandler(cap)
        desk_log.propagate = old_prop


# ============================================================================ check 5: secrets & stray files
SECRET_PATTERNS: list[tuple[str, re.Pattern]] = [
    ("ClickUp token (pk_...)", re.compile(r"\bpk_\d{3,}_[A-Za-z0-9]{16,}|\bpk_[A-Za-z0-9]{24,}")),
    ("DeepSeek / OpenAI style key (sk-...)", re.compile(r"\bsk-[A-Za-z0-9_\-]{20,}")),
    ("webhook URL with a token", re.compile(
        r"(?i)https?://[^\s'\"<>]*(?:hooks?|webhooks?)[^\s'\"<>]*[?&/](?:token|key|secret|auth|sig|signature)[=/][A-Za-z0-9._\-%]{6,}")),
    ("webhook URL that is itself the secret", re.compile(
        r"(?i)https?://hooks\.zapier\.com/hooks/catch/\d+/[A-Za-z0-9]+|https?://hook\.[a-z0-9.]*make\.com/[a-z0-9]{16,}|"
        r"https://discord(?:app)?\.com/api/webhooks/\d+/[\w\-]{20,}|https://hooks\.slack\.com/services/[A-Z0-9/]{20,}")),
    ("credentials inside a URL", re.compile(r"(?i)\b[a-z][a-z0-9+.\-]*://[^/\s:@'\"<>]+:(?P<pw>[^/\s@'\"<>]{3,})@[^\s'\"<>]+")),
    ("token / key in a URL query", re.compile(r"(?i)[?&](?:token|access_token|api_key|apikey|secret)=[A-Za-z0-9._\-%]{16,}")),
    ("private key block", re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----")),
    ("GitHub / AWS token", re.compile(r"\bghp_[A-Za-z0-9]{30,}|\bgithub_pat_[A-Za-z0-9_]{30,}|\bAKIA[0-9A-Z]{16}\b")),
    ("bearer token literal", re.compile(r"\bBearer\s+[A-Za-z0-9._\-]{24,}")),
]
# "strong" names (password / secret / api key ...): any non-placeholder literal is a finding.
# "weak" names (a bare `token`, `webhook_token` ...): too common in normal code (max_tokens, token_header ...), so the value
# must also LOOK like a random secret (see _looks_random). Both kinds are matched by _NAME.
_STRONG_WORDS = r"password|passwd|pwd|passphrase|secret|api[_-]?key|api[_-]?token|access[_-]?token|auth[_-]?token"
_WEAK_WORDS = r"token|webhook"
_NAME = r"[A-Za-z0-9_.\-]*(?:" + _STRONG_WORDS + "|" + _WEAK_WORDS + r")[A-Za-z0-9_\-]*"
_STRONG_NAME_RE = re.compile(r"(?i)" + _STRONG_WORDS)
# NAME = "literal"   /   NAME: 'literal'   /   "X-Webhook-Token": "literal"
QUOTED_ASSIGN_RE = re.compile(r"(?i)\b(?P<name>" + _NAME + r")\b[\"']?\s*[:=]\s*(?:f|r|b)?([\"'])(?P<val>[^\"'\n]{4,})\2")
# KEY=value (.env / .bat / .ps1 / shell / config style, unquoted)
BARE_ASSIGN_RE = re.compile(r"(?i)^\s*(?:export\s+|set\s+|\$env:)?(?P<name>" + _NAME + r")\s*=\s*(?P<val>[^\s\"'#;]{4,})\s*$")
# HTTP header / YAML style, unquoted, value runs to the end of the line:  X-Webhook-Token: <value>
HEADER_RE = re.compile(r"(?i)\b(?P<name>" + _NAME + r")[\"']?\s*:\s*[\"']?(?P<val>[A-Za-z0-9_\-+=]{16,})[\"']?[,;)\]}]?\s*$")
# os.getenv("CLICKUP_API_TOKEN", "<literal default>") - a secret hidden in the default argument
GETENV_DEFAULT_RE = re.compile(
    r"(?i)\b(?:getenv|environ\.get|env\.get|get_env|config\.get|settings\.get)\(\s*[\"'](?P<name>" + _NAME + r")[\"']\s*,\s*(?:f|r|b)?([\"'])(?P<val>[^\"'\n]{4,})\2")
SQL_PASSWORD_RE = re.compile(r"(?i)\bpassword\s+'(?P<val>[^']{4,})'")
# a URL (scanned for opaque secret path segments and for known-leaked values)
URL_RE = re.compile(r"(?i)\bhttps?://[^\s'\"<>()\[\]{}\\`]+")
_LONG_ALNUM_RE = re.compile(r"[A-Za-z0-9]{16,}")
# path segments that legitimately carry a long hex / alnum id (commit URLs ...): not secrets
_PUBLIC_ID_PARENTS = {"commit", "commits", "tree", "blob", "raw", "compare", "releases", "tag", "runs", "pull", "issues", "sha", "digest"}

# sha256 (of the raw value) -> (length, why) of secrets that were leaked once. The value itself must NEVER be stored here:
# a future agent that pastes it into a new file / diff / note fails the scan wherever it appears (URL, header, comment ...).
KNOWN_LEAKED: dict[str, tuple[int, str]] = {
    "db2c00e48c6c102ecf9fa5599dfc88855fe6f13c05a4eb3d3aa63312a197c00a": (32, "the webhook token that was committed once (lessons L-052)"),
}
BARE_ASSIGN_SUFFIX = {".env", ".bat", ".cmd", ".ps1", ".sh", ".ini", ".cfg", ".toml", ".yml", ".yaml", ".txt", ".example", ""}
# What a STRONG secret name (password / secret / api_key / api_token ...) may be assigned WITHOUT being a finding: only a
# value that is CLEARLY fake. Nothing that merely smells like a placeholder is exempt any more - an ALL_CAPS word, a value
# starting with "my" / "the" / "a", a word with a few digits - because a real value such as hunter22xyz, ADMIN2026 or
# my-real-pass1 fits all of those and must always be reported (lessons L-072, L-124). Matched against the WHOLE value.
_TEMPLATE_RE = re.compile(
    r"<[^<>\n]+>|\$\{[^{}\n]+\}|\{\{[^{}\n]+\}\}|\{[^{}\n]*\}|\$[A-Za-z_][A-Za-z0-9_]*|%[A-Za-z_][A-Za-z0-9_]*%|"
    r"%\([A-Za-z_]\w*\)[sd]|%[sd]|\*{3,}|\.{3,}")                    # <x>  ${X}  {{x}}  {x}  $X  %X%  %(x)s  %s  ***  ...
_FAKE_VALUE_RE = re.compile(
    r"(?i:your(?:[_\- ][a-z]+){1,5}\d{0,2}"                          # YOUR_PASSWORD, your-api-key-here
    r"|change[_\- ]?me(?:[_\- ]?(?:now|please))?\d{0,2}"             # changeme, CHANGE_ME
    r"|x{3,}(?:[_\-. ]x{3,})*"                                       # xxx, xxxxxxxx, xxxx-xxxx
    r"|example"
    r"|(?:password|passwd|secret|token|key|value|pass|api[_\-]?key)[_\- ]example"
    r"|example[_\- ](?:password|passwd|secret|token|key|value|pass|api[_\-]?key))")


# Kept for _looks_random() only: a WEAK name (a bare `token` ...) still needs the looser "is this even a word?" test.
_PLACEHOLDER_WORD_RE = re.compile(r"(?i)(?:your|my|the|a|an|example|sample|dummy|fake|change_?me|placeholder|replace|todo|none|null|redacted|hidden)(?:[^A-Za-z0-9]|$)")


def _is_placeholder(val: str) -> bool:
    """Is this value of a STRONG-named secret clearly a template / fake (exempt), rather than possibly the real thing?"""
    v = val.strip().strip("\"'`")
    return not v or bool(_TEMPLATE_RE.fullmatch(v) or _FAKE_VALUE_RE.fullmatch(v))


def _looks_random(val: str) -> bool:
    """Does a value look like a machine-generated token (not a word, a file name, an env var name or a number)?"""
    v = val.strip().strip("\"'`")
    if len(v) < 16 or re.search(r"[\s/:.<>{}$%*]", v) or _PLACEHOLDER_WORD_RE.match(v) or re.search(r"(?i)x{4,}|\.{3,}", v):
        return False
    if v.isupper() and "_" in v:
        return False                                   # SOME_ENV_VARIABLE_NAME
    letters, digits = sum(c.isalpha() for c in v), sum(c.isdigit() for c in v)
    return letters >= 6 and digits >= 2


def _is_secret_literal(name: str, val: str) -> bool:
    """Is `NAME = "val"` a hard-coded secret? Strong names (password ...): any non-placeholder value; weak names (token ...): random-looking."""
    if _STRONG_NAME_RE.search(name):
        return not _is_placeholder(val)
    return _looks_random(val)


def _opaque_url_token(url: str) -> bool:
    """A URL whose PATH holds a long letters+digits segment (24+ chars): the shape of a secret webhook / callback URL, e.g.
    https://host/api/<name>/webhooks/<32-char token> - the exact class of secret this project once leaked."""
    m = re.match(r"(?i)https?://([^/?#]*)([^?#]*)", url)
    if not m:
        return False
    segs = [x for x in m.group(2).split("/")]
    for i, seg in enumerate(segs):
        if len(seg) < 24 or not seg.isalnum() or not seg.isascii():
            continue
        if not (any(c.isdigit() for c in seg) and any(c.isalpha() for c in seg)):
            continue                                   # a long word / a plain number is not a token
        if i > 0 and segs[i - 1].lower() in _PUBLIC_ID_PARENTS:
            continue                                   # .../commit/<sha>: public identifier
        return True
    return False


def _known_leaked(text: str, deny: dict[str, tuple[int, str]] | None = None) -> str | None:
    """Reason if `text` contains a value whose sha256 is in the denylist (the value itself is never stored)."""
    deny = KNOWN_LEAKED if deny is None else deny
    lengths = sorted({n for n, _ in deny.values()})
    for run in _LONG_ALNUM_RE.findall(text):
        for n in lengths:
            for i in range(len(run) - n + 1):
                hit = deny.get(hashlib.sha256(run[i:i + n].encode("utf-8")).hexdigest())
                if hit:
                    return hit[1]
    return None


def scan_line(path: str, line: str) -> list[str]:
    """Names of the secret patterns found in one line (never the matched text)."""
    found: list[str] = []

    def add(name: str) -> None:
        if name not in found:
            found.append(name)

    step = 4000                                                     # bound the regex work on huge minified lines
    for start in range(0, max(len(line), 1), step - 200):
        chunk = line[start:start + step]
        for name, rx in SECRET_PATTERNS:
            if name in found:
                continue
            for m in rx.finditer(chunk):
                if "pw" in rx.groupindex and _is_placeholder(m.group("pw")):
                    continue                       # postgresql://{USER}:{PASSWORD}@host style templates
                found.append(name)
                break
        if any(_opaque_url_token(u) for u in URL_RE.findall(chunk)):
            add("opaque token in a URL path (secret webhook URL?)")
        why = _known_leaked(chunk)
        if why:
            add("known leaked secret: " + why)
        for rx in (QUOTED_ASSIGN_RE, GETENV_DEFAULT_RE, HEADER_RE):
            for m in rx.finditer(chunk):
                if _is_secret_literal(m.group("name"), m.group("val")):
                    add("password / secret literal")
        m = SQL_PASSWORD_RE.search(chunk)
        if m and not _is_placeholder(m.group("val")):
            add("password literal in SQL")
        if Path(path).suffix.lower() in BARE_ASSIGN_SUFFIX:
            m = BARE_ASSIGN_RE.match(chunk)
            if m and _is_secret_literal(m.group("name"), m.group("val")):
                add("password / secret literal")
        if start + step >= len(line):
            break
    return found


def _selftest_scanner() -> str | None:
    """The detectors must catch known-bad lines and let clean ones through (samples are assembled at run time)."""
    rnd = "a1b2c3d4" * 4                                   # 32 chars, letters + digits: shaped like a real opaque token
    short = "Qw8Er5Ty2Ui9Op3A"                             # 16 chars
    bad = {
        "clickup": "CLICKUP_" + "TOKEN = '" + "pk_" + "12345678" + "_" + "A1" * 16 + "'",
        "deepseek": "key = 'sk" + "-" + "a1B2c3D4" * 4 + "'",
        "webhook": "url = 'https://hooks.example.com/hooks/" + "abc" + "?tok" + "en=" + "Zz9" * 6 + "'",
        "pw literal": "DB_PASS" + "WORD = 'hunter" + "22xyz'",
        # a real value that only LOOKS like a placeholder must never be exempt (L-072): ALL_CAPS, a "my..." word, a short word
        "ALL_CAPS pw": "DB_PASS" + "WORD = 'ADMIN" + "2026'",
        "my-prefixed pw": "SMTP_PASS" + "WORD = 'my-real" + "-pass1'",
        "short pw": "api_" + "key = 'Qw8Er" + "5Ty2'",
        "env ALL_CAPS pw": "PGPASS" + "WORD=ERP_APP" + "_PASSWORD",
        "url creds": "postgresql://erp:" + "s3cr3tpw" + "@127.0.0.1:5432/db",
        "env line": "SMTP_PASS" + "WORD=" + "Abcd1234efgh",
        # the class of secret this project was burned by (L-052): the token is an opaque PATH segment on an arbitrary host
        "opaque-token url (py)": "hook_u" + "rl = \"https://crm.example.test/api/integration-builder/webh" + "ooks/" + rnd + "\"",
        "opaque-token url (bare)": "curl -X POST https://crm.example.test/api/in/" + rnd + " -d @x.json",
        "opaque-token url (short path)": "https://crm.example.test/" + rnd[:26] + "/",
        "WEBHOOK_" + "TOKEN literal": "WEBHOOK_" + "TOKEN = '" + rnd + "'",
        "bare token literal": "tok" + "en = '" + short + "'",
        "token dict entry": "{\"tok" + "en\": \"" + rnd + "\"}",
        "X-Webhook-Token header (line)": "X-Webhook-" + "Token: " + rnd,
        "X-Webhook-Token header (dict)": "headers = {\"X-Webhook-" + "Token\": \"" + rnd + "\"}",
        "getenv default": "os.getenv('CLICKUP_API_" + "TOKEN', '" + short + "')",
        "token in .env": "WEBHOOK_" + "TOKEN=" + rnd,
    }
    for label, line in bad.items():
        if not scan_line("x.env" if label in ("env line", "token in .env", "env ALL_CAPS pw") else "x.py", line):
            return f"scanner missed a known-bad line ({label})"
    good = ["DB_PASSWORD = os.getenv('DB_PASSWORD', '')", "password=''", "PASSWORD=", "DB_PASSWORD=your_password",
            # the ONLY exemptions left for a strong name: clearly fake values (L-072 / L-124)
            "DB_PASS" + "WORD = '<ERP_APP_PASSWORD>'", "SMTP_PASS" + "WORD = 'changeme'", "API_" + "SECRET = '${API_SECRET}'",
            "api_" + "key = 'YOUR_API_KEY_HERE'", "DB_PASS" + "WORD = 'example'", "SMTP_PASS" + "WORD = 'xxxxxxxx'",
            "-v pw=\"'ERP_APP_PASSWORD'\"", "TOKEN_HEADER = 'x-erp-desk-token'", "the ClickUp token starts with pk_ and DeepSeek with sk-",
            "https://fonts.googleapis.com/css2?family=Montserrat:wght@400;500;600;700;800&display=swap", "smtp_password = cfg.smtp_password",
            "password = st.text_input('Password')",
            "max_tokens = 4000", "self.token = token", "token: str", "webhook_url = os.getenv('WEBHOOK_URL', '')",
            "WEBHOOK_TOKEN = os.getenv('WEBHOOK_TOKEN', '')", "WEBHOOK_TOKEN=your_webhook_token", "token = st.session_state['token']",
            "X-Webhook-Token: <your token here>", "WEBHOOK_URL = 'https://your-host.example/api/webhooks/<TOKEN>'",
            "https://cdnjs.cloudflare.com/ajax/libs/three.js/r128/three.min.js",
            "https://github.com/acme/erp/commit/" + "0123456789abcdef" * 2 + "01234567",
            "https://fonts.gstatic.com/s/montserrat/v31/" + "6NUh8FyLNQOQZAnv9bYEvDiIdE9Ea92uemAk" + ".woff2",
            "https://example.com/a-very-long-readable-article-title-about-erp-systems-2026",
            "https://example.com/Pneumonoultramicroscopicsilicovolcanoconiosis/page"]
    for line in good:
        if scan_line("x.py", line):
            return f"scanner flagged a clean line ({line[:40]}...)"
    # the known-leaked denylist: hashes only, must work wherever the value appears (comment, header, prose, glued to text)
    for h, (n, why) in KNOWN_LEAKED.items():
        if not (re.fullmatch(r"[0-9a-f]{64}", h) and n >= 16 and why):
            return "KNOWN_LEAKED entry is malformed (needs a sha256 hex digest, the value length and a reason)"
    fake = {hashlib.sha256(rnd.encode()).hexdigest(): (len(rnd), "unit-test value")}
    for line in ("# see " + rnd, "auth=" + rnd + ";", "prefix" + rnd + "suffix", "x-key: " + rnd):
        if not _known_leaked(line, fake):
            return "the known-leaked denylist missed a value"
    if _known_leaked("nothing to see: " + rnd[:-1] + "z", fake):
        return "the known-leaked denylist flagged a different value"
    return None


def parse_added_lines(diff: str) -> list[tuple[str, int, str]]:
    """(path, line number, text) of every line a unified diff adds."""
    out: list[tuple[str, int, str]] = []
    path, lineno, in_hunk = None, 0, False
    for raw in diff.splitlines():
        if raw.startswith("diff --git"):
            in_hunk, path = False, None
        elif not in_hunk and raw.startswith("+++ "):
            p = raw[4:].strip()
            path = None if p == "/dev/null" else (p[2:] if p.startswith("b/") else p)
        elif raw.startswith("@@"):
            in_hunk = True
            m = re.match(r"@@ -\S+ \+(\d+)", raw)
            lineno = int(m.group(1)) if m else 0
        elif in_hunk and raw.startswith("+") and path:
            out.append((path, lineno, raw[1:]))
            lineno += 1
        elif in_hunk and not raw.startswith("-"):
            lineno += 1
    return out


def _stray_reason(path: str, size: int | None) -> str | None:
    base = path.rsplit("/", 1)[-1]
    low = base.lower()
    if any(fnmatch.fnmatch(low, g) for g in SECRET_FILE_GLOBS) and low not in SECRET_FILE_OK:
        return "a secrets / credentials file"
    if any(fnmatch.fnmatch(low, g) for g in STRAY_GLOBS):
        return "a log / temporary / editor file"
    if STRAY_NAME_RE.match(base):
        return "a scratch file"
    if base in KNOWN_OUTPUTS:
        return "a generated output of the project"
    if Path(base).suffix.lower() in DATA_OUTPUT_EXT and not path.startswith(DATA_OUTPUT_OK_PREFIX):
        return "a data output (csv / json / ...): put test data under tests/fixtures/ if it is really needed"
    if size == 0 and base not in ("__init__.py", ".gitkeep"):
        return "an empty file (accidental shell redirect?)"
    return None


def worktree_files(root: Path = ROOT) -> set[str]:
    """Untracked + git-ignored files and folders (ignored folders collapsed): what a run could have left behind."""
    out: list[str] = []
    for args in (("ls-files", "--others", "--exclude-standard"),
                 ("ls-files", "--others", "--ignored", "--exclude-standard", "--directory")):
        r = git(*args, cwd=root)
        if r.returncode != 0:      # never let "git failed" look like "no files": check 7 would pass for the wrong reason
            raise GateError(f"git ls-files failed in {root}: {one_line(r.stderr) or 'not a git repository?'}")
        out.append(r.stdout)
    return {ln.strip() for ln in "\n".join(out).splitlines() if ln.strip()}


_REF_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._/@\-]*")


def merge_base(base_ref: str, root: Path = ROOT) -> str:
    """The commit where HEAD left `base_ref`.

    The gate diffs against THIS, never against the tip of the base branch: a main that moved on after the branch was cut
    would otherwise present its own newer commits as reversed changes of this branch (lesson L-074).
    """
    r = git("merge-base", base_ref, "HEAD", cwd=root)
    sha = r.stdout.strip()
    if r.returncode != 0 or not re.fullmatch(r"[0-9a-f]{40,64}", sha):
        raise GateError(f"HEAD and '{base_ref}' have no common ancestor (unrelated histories?): "
                        f"{one_line(r.stderr) or 'git merge-base found none'}")
    return sha


def branch_history(mb: str, root: Path = ROOT) -> list[tuple[str, str, str]]:
    """[(sha, commit message, its unified diff)] for every commit in merge-base..HEAD (a merge has no diff of its own)."""
    r = git("log", "-p", "--no-color", "--no-ext-diff", "--no-renames", "-U0", "--format=%x01%H%x1f%B%x1e",
            f"{mb}..HEAD", timeout=60, cwd=root)
    if r.returncode != 0:
        raise GateError(f"git log {mb[:8]}..HEAD failed: {one_line(r.stderr)}")
    commits: list[tuple[str, str, str]] = []
    for chunk in r.stdout.split("\x01")[1:]:
        sha, _, rest = chunk.partition("\x1f")
        message, _, diff = rest.partition("\x1e")
        commits.append((sha.strip(), message, diff))
    return commits


def check_secrets_and_strays(base_ref: str, root: Path = ROOT) -> tuple[str, str, list[str]]:
    problem = _selftest_scanner()
    if problem:
        return FAIL, f"secret scanner self-test failed: {problem}", []
    if not _REF_RE.fullmatch(base_ref or ""):
        return FAIL, f"the --base value {base_ref!r} is not a plain branch / commit name", []
    if git("rev-parse", "--is-inside-work-tree", cwd=root).returncode != 0:
        return FAIL, f"not inside a git work tree ({root})", []
    if git("rev-parse", "--verify", "--quiet", base_ref + "^{commit}", cwd=root).returncode != 0:
        return FAIL, f"base branch '{base_ref}' does not exist (use --base <branch>)", []
    if git("rev-parse", "--verify", "--quiet", "HEAD^{commit}", cwd=root).returncode != 0:
        return FAIL, "the current branch has no commit yet (unborn branch): there is nothing to compare with the base", []
    mb = merge_base(base_ref, root)
    branch = current_branch(root)

    findings: list[str] = []
    notes: list[str] = []
    scanned_files = 0

    # -- what changed since the MERGE BASE: committed on this branch + staged + unstaged (tracked) ----------
    changed = git("diff", "--name-status", "--no-renames", mb, cwd=root)
    if changed.returncode != 0:
        return FAIL, f"git diff {mb[:8]} failed: {one_line(changed.stderr)}", []
    paths = []
    for ln in changed.stdout.splitlines():
        st, _, p = ln.partition("\t")
        if st and not st.startswith("D") and p:
            paths.append(p)
    committed = set(git("diff", "--name-only", "--no-renames", mb, "HEAD", cwd=root).stdout.splitlines())
    staged = set(git("diff", "--cached", "--name-only", "--no-renames", cwd=root).stdout.splitlines())
    untracked = [p for p in git("ls-files", "--others", "--exclude-standard", cwd=root).stdout.splitlines() if p.strip()]
    tracked_env = [p for p in git("ls-files", cwd=root).stdout.splitlines()
                   if any(fnmatch.fnmatch(p.rsplit("/", 1)[-1].lower(), g) for g in (".env", ".env.*")) and p.lower().rsplit("/", 1)[-1] not in SECRET_FILE_OK]
    for p in tracked_env:
        findings.append(f"{p}: a .env file is tracked by git")

    owner_left = []
    for p in sorted(set(paths) | set(untracked)):
        size = None
        fp = root / p
        try:
            size = fp.stat().st_size if fp.is_file() else None
        except OSError:
            pass
        if p in OWNER_FILES:
            if p in committed or p in staged:
                findings.append(f"{p}: the owner's uncommitted file was {'committed on this branch' if p in committed else 'staged'} (Law.md rule 9)")
            else:
                owner_left.append(p)
            continue
        why = _stray_reason(p, size)
        if why and p not in tracked_env:
            findings.append(f"{p}: {why}")
    if owner_left:
        notes.append(f"owner's uncommitted files left alone (not part of the change): {', '.join(owner_left)}")

    # -- ignored scratch files (git-ignored logs would never show up in the diff) -------------------
    for entry in sorted(worktree_files(root) - set(untracked)):
        if entry.endswith("/"):
            continue
        base = entry.rsplit("/", 1)[-1]
        if base.lower().endswith(".log") and base not in KNOWN_LOGS:
            findings.append(f"{entry}: a scratch log file (git-ignored, but delete it)")
        elif STRAY_NAME_RE.match(base) or any(fnmatch.fnmatch(base.lower(), g) for g in ("*.tmp", "*.bak", "*.orig", "*.pid", "*.temp")):
            findings.append(f"{entry}: a scratch / temporary file")

    # -- secret patterns in every added line and every untracked text file --------------------------
    secret_values = _env_secret_values()

    def scan_text(path: str, text_line: str) -> list[str]:
        """Finding names for one line (never the matched text)."""
        names = scan_line(path, text_line)
        for v in secret_values:
            if v in text_line and re.search(r"(?<![A-Za-z0-9])" + re.escape(v) + r"(?![A-Za-z0-9])", text_line):
                names.append("a value copied from your local .env")
        return names

    diff = git("diff", "--no-color", "--no-ext-diff", "--no-renames", "-U0", mb, timeout=60, cwd=root)
    if diff.returncode != 0:
        return FAIL, f"git diff {mb[:8]} failed: {one_line(diff.stderr)}", []
    added = parse_added_lines(diff.stdout)
    seen_path_name: set[tuple[str, str]] = set()
    for path, lineno, text_line in added:
        for name in scan_text(path, text_line):
            findings.append(f"{path}:{lineno}: {name}")
            seen_path_name.add((path, name))
    for p in untracked:
        fp = root / p
        if Path(p).suffix.lower() in BINARY_EXT:
            continue
        try:
            if not fp.is_file() or fp.stat().st_size > 2_000_000:
                continue
            data = fp.read_bytes()
        except OSError:
            continue
        if b"\0" in data[:8000]:
            continue
        scanned_files += 1
        for i, text_line in enumerate(data.decode("utf-8", errors="replace").splitlines(), 1):
            for name in scan_text(p, text_line):
                findings.append(f"{p}:{i}: {name}")

    # -- every commit and every commit MESSAGE of the branch: because merges are --no-ff, a secret that was added and
    #    removed again later on the branch still lands in main's history, although the final diff above is clean (L-074)
    commits = branch_history(mb, root)
    history_lines = 0
    for sha, message, cdiff in commits:
        for i, text_line in enumerate(message.splitlines(), 1):
            history_lines += 1
            for name in scan_text("commit-message.txt", text_line):
                findings.append(f"commit {sha[:7]} message line {i}: {name}")
        for path, lineno, text_line in parse_added_lines(cdiff):
            history_lines += 1
            for name in scan_text(path, text_line):
                if (path, name) not in seen_path_name:       # the working-tree diff above already reported this one
                    findings.append(f"commit {sha[:7]} {path}:{lineno}: {name} (in the branch history, even though it "
                                    f"may have been removed later)")

    # findings never carry the matched text, but keep the list short and unique
    seen, uniq = set(), []
    for f in findings:
        if f not in seen:
            seen.add(f)
            uniq.append(f)
    summary = (f"{branch} vs {base_ref} (merge-base {mb[:8]}): {len(set(paths))} changed + {len(untracked)} new files, "
               f"{len(added)} added lines + {len(commits)} commit(s) / {history_lines} history lines scanned")
    if uniq:
        return FAIL, f"{len(uniq)} finding(s): {uniq[0]}" + (" (+more)" if len(uniq) > 1 else ""), uniq + notes
    return PASS, summary + "; no secrets, no stray files", notes


# ============================================================================ check 6: agent system
# ----------------------------------------------------------------------------- typography (part of check 6)
# The project font is Montserrat. The retired names are assembled at run time so this file does not flag itself (L-073).
_OLD_FONTS = re.compile("Fraun" + "ces|Public[ +_-]?" + "Sans|IBM[ +_-]?" + "Plex", re.IGNORECASE)
# Documents where the retired names may legitimately appear - but only on a LINE that is visibly history or a quotation.
# A whole file is never exempt (lesson L-124): a change log is exactly where somebody would paste a live font rule.
_OLD_FONT_DOCS = (".claude/lessons.md", ".claude/journal.md", ".claude/backlog.md", ".claude/skills/", "PROJECT_NOTES.md",
                  "README.md")
# What makes a line history / a quotation rather than a live instruction.
_HISTORY_MARKER = re.compile(r"(?i)\b(?:history|historic|formerly|retired|superseded|obsolete|deprecated|replaced|replacing|"
                             r"was|were|used to|old|previous|before|until|no longer|instead of|change ?log|trio)\b"
                             r"|\b20\d\d-\d\d-\d\d\b|\bL-\d{3}\b")
# A line that actually SETS a font never gets the history exemption, whatever prose surrounds it: the marker list above
# is a heuristic, and a code block pasted into a change log is exactly where a live rule could hide (L-124).
_LIVE_FONT_RULE = re.compile(r"(?i)font-family\s*:|--(?:font|code)\s*:|\bfont\s*[:=]\s*[\"']|[?&]family=|\bfont\.face\b|"
                             r"\bfontFamily\b|\bFONT_(?:NAME|STACK)\s*=")
# A check that had to be skipped (streamlit not importable) explains itself in the check-6 notes.
TYPOGRAPHY_NOTES: list[str] = []
_BINARY_EXT = {".png", ".ico", ".jpg", ".jpeg", ".gif", ".woff", ".woff2", ".ttf", ".pb", ".zip", ".pyc"}


def old_font_problem(path: str, line: str) -> str | None:
    """Reason this LINE must not name a retired font, else None.

    Rule: any tracked line naming one of the retired families (_OLD_FONTS, spelled at run time so this file does not
    flag itself) is a finding, EXCEPT a line in a change-log / decision-log document that also carries a history marker
    on the same line. Exempting whole documents would let a live `font-family:` rule hide in a change log (L-124).
    """
    if not _OLD_FONTS.search(line):
        return None
    if path.startswith(_OLD_FONT_DOCS) and _HISTORY_MARKER.search(line) and not _LIVE_FONT_RULE.search(line):
        return None
    if _LIVE_FONT_RULE.search(line):
        where = "a line that actually SETS a font (no history exemption applies)"
    elif path.startswith(_OLD_FONT_DOCS):
        where = "a change-log line without a history marker"
    else:
        where = "live source"
    return (f"names a retired font in {where} - the project font is Montserrat "
            f"(tokens: erp/typography.py, --font in desktop/static/shell.css)")


def font_consistency_problems(css_text: str, html_text: str, font_name: str, font_stack: str, code_stack: str,
                              fonts_url: str) -> list[str]:
    """The drift that used to fail SILENTLY into the fallback font (L-124): shell.css, shell.html and erp/typography.py
    must agree - same family in all three, and the Google Fonts <link> must be exactly the URL typography.py builds
    (so a weight added in one place cannot go missing in the other)."""
    out: list[str] = []

    def family(decl: str) -> str:
        return (decl or "").split(",")[0].strip().strip("'\"").strip()

    m = re.search(r"--font:\s*([^;]+);", css_text)
    css_font = m.group(1).strip() if m else ""
    m = re.search(r"--code:\s*([^;]+);", css_text)
    css_code = m.group(1).strip() if m else ""
    if not css_font:
        out.append("desktop/static/shell.css: there is no --font token")
    elif css_font != font_stack:
        out.append(f"desktop/static/shell.css: --font is {css_font!r} but erp/typography.py FONT_STACK is {font_stack!r} "
                   "- the two copies of the token must stay identical")
    if not css_code:
        out.append("desktop/static/shell.css: there is no --code token")
    elif css_code != code_stack:
        out.append(f"desktop/static/shell.css: --code is {css_code!r} but erp/typography.py CODE_STACK is {code_stack!r}")

    links = re.findall(r'<link[^>]+href="(https://fonts\.googleapis\.com/css2\?[^"]+)"', html_text)
    if not links:
        out.append("desktop/static/shell.html: no Google Fonts css2 <link> - ERP Desk would silently fall back to Segoe UI")
    else:
        for href in links:
            if href != fonts_url:
                out.append("desktop/static/shell.html: the Google Fonts <link> is not the URL erp/typography.py builds "
                           "(GOOGLE_FONTS_CSS_URL) - family / weights / display have drifted apart")
        html_fam = ""
        m = re.search(r"[?&]family=([^:&\"]+)", links[0])
        if m:
            html_fam = m.group(1).replace("+", " ").strip()
        if html_fam != font_name:
            out.append(f"desktop/static/shell.html: the <link> loads {html_fam or '?'!r} but the project font is "
                       f"{font_name!r} - the page would render in the fallback font with nobody noticing")
        if css_font and family(css_font) != html_fam:
            out.append(f"desktop/static/shell.css --font starts with {family(css_font)!r} but shell.html loads "
                       f"{html_fam or '?'!r}: the CSS asks for a font the page never downloads")
    return out


# --- the weights the project downloads, and the numerals rule -----------------------------------------------------
# Where a CSS rule lives that the weight / numeral scans must read (a .py file carries its CSS in an f-string).
FONT_CSS_SOURCES = ("desktop/static/shell.css", "desktop/static/report.css", "desktop/static/today.css",
                    "desktop/static/leads.css", "desktop/static/channels.css", "desktop/static/flow.css", "desktop/static/flow3d.css",
                    "dashboard/script_center.py")
# Every rule that renders a number which ANIMATES (counts up, or is redrawn by a poll) or stands in a row / column of
# other numbers. Montserrat's proportional figures differ in width by about 2x, so such a number resizes its own tile
# on every frame; tabular figures are written into the rule itself, so no missing class can undo it (lesson L-112).
TABULAR_NUMBER_RULES = (
    ("desktop/static/report.css", ".kpi-value", "Reporting KPI value - report.js tween() counts it up"),
    ("desktop/static/report.css", ".f-nums b", "Reporting funnel count - report.js tween() counts it up"),
    ("desktop/static/report.css", ".bm-value", "Reporting benchmark value - a column of tiles"),
    ("desktop/static/flow.css", ".tally-item b", "Data Flow tally - flow.js tween() on [data-tally]"),
    ("desktop/static/flow.css", ".cov-ring .center b", "Data Flow coverage ring centre - redrawn on refresh"),
    ("desktop/static/today.css", ".t-kpi-value", "Today KPI value - redrawn on every poll, six tiles in a row"),
    ("desktop/static/leads.css", ".ld-tri-c b", "Leads triptych - three big numbers side by side"),
    ("desktop/static/channels.css", ".ch-kpi-value", "Channels KPI value - redrawn on every poll, six tiles in a row"),
    ("desktop/static/channels.css", ".ch-share-t", "Channels share-of-sessions text - a column of percentages beside the bar"),
    ("dashboard/script_center.py", ".sc-kpi-value", "Script Center KPI row"),
    ("dashboard/script_center.py", ".sc-digest-tile-value", "Script Center digest tiles"),
)


def css_declaration_block(text: str, selector: str) -> str | None:
    """The declarations inside the first `selector { ... }` rule of `text`, or None when there is no such rule.

    `text` for a .py file must come through css_source() first: CSS inside an f-string doubles its braces.
    """
    m = re.search(re.escape(selector) + r"\s*\{", text)
    if not m:
        return None
    depth, start, i = 0, m.end(), m.end() - 1
    while i < len(text):
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
            if depth == 0:
                return text[start:i]
        i += 1
    return None


def css_source(path: str, text: str) -> str:
    """CSS written inside a Python f-string doubles every brace; undo that so one scanner reads both kinds of file."""
    return text.replace("{{", "{").replace("}}", "}") if path.endswith(".py") else text


def tabular_number_problems(sources: dict[str, str]) -> list[str]:
    """L-112: each rule in TABULAR_NUMBER_RULES must still ask for tabular figures."""
    out: list[str] = []
    for path, selector, what in TABULAR_NUMBER_RULES:
        text = sources.get(path)
        if text is None:
            out.append(f"{path}: cannot be read - the tabular-figure rule for `{selector}` could not be checked")
            continue
        block = css_declaration_block(css_source(path, text), selector)
        if block is None:
            out.append(f"{path}: there is no `{selector}` rule any more ({what}) - move this entry in "
                       "TABULAR_NUMBER_RULES to wherever that number is rendered now")
        elif "tabular-nums" not in block:
            out.append(f"{path}: `{selector}` ({what}) no longer asks for `font-variant-numeric: tabular-nums` - "
                       "a number that moves would change its own width on every frame (L-112)")
    return out


def font_weight_problems(weights: str, css_sources: dict[str, str], js_text: str) -> list[str]:
    """L-113: a weight is downloaded only if some CSS rule paints it, and the font gate loads exactly those weights."""
    declared = [w.strip() for w in weights.split(";") if w.strip()]
    used: set[str] = set()
    for path, text in css_sources.items():
        used.update(re.findall(r"font-weight\s*:\s*(\d{3})", css_source(path, text)))
    out: list[str] = []
    for w in declared:
        if w not in used:
            out.append(f"erp/typography.py FONT_WEIGHTS lists weight {w}, but no CSS rule uses `font-weight: {w}` - "
                       "it is a font file downloaded for nothing and it delays the first chart draw (L-113)")
    for w in sorted(used - set(declared)):
        out.append(f"a CSS rule asks for `font-weight: {w}`, which FONT_WEIGHTS does not download - the browser would "
                   "synthesise that weight from another one")
    m = re.search(r"var faces = \[([^\]]*)\]", js_text)
    if not m:
        out.append("desktop/static/shell.js: the font gate's weight list (`var faces = [...]`) could not be found - "
                   "charts wait on it, so it must stay checkable")
    elif re.findall(r"\d{3}", m.group(1)) != declared:
        out.append(f"desktop/static/shell.js: the font gate loads {re.findall(r'[0-9]{3}', m.group(1))} but "
                   f"erp/typography.py FONT_WEIGHTS is {declared} - a weight the <link> never fetches makes every "
                   "chart wait for a face that cannot arrive (L-113)")
    return out


def streamlit_font_flag_problems(theme_args: list[str], font_name: str, fonts_url: str,
                                 code_stack: str) -> tuple[list[str], str]:
    """Parse the --theme.* font flags with STREAMLIT'S OWN parser instead of comparing two copies of a string.

    Two byte-equal copies prove consistency, not correctness - they can be wrong together (L-124). Streamlit splits
    `<family>:<url>` on the FIRST colon, so anything appended after the URL is swallowed into it and Google Fonts
    then answers without `display: swap` (L-114). Returns (problems, note); the note explains a skipped check.
    """
    args = {theme_args[i]: theme_args[i + 1] for i in range(0, len(theme_args), 2)}
    root_logger = logging.getLogger()
    before, level = list(root_logger.handlers), root_logger.level
    try:
        from streamlit.runtime.theme_util import _parse_font_config
    except Exception as exc:  # noqa: BLE001 - streamlit missing / broken is a note, not a red gate
        return [], f"streamlit could not be imported ({type(exc).__name__}), so the --theme.* font flags were not parsed"
    finally:
        for h in list(root_logger.handlers):
            if h not in before:
                root_logger.removeHandler(h)
        root_logger.setLevel(level)

    out: list[str] = []
    for flag in ("--theme.font", "--theme.headingFont"):
        value = args.get(flag)
        if value is None:
            out.append(f"erp/typography.py streamlit_theme_args() no longer passes {flag}")
            continue
        try:
            name, url = _parse_font_config(value, flag)
        except Exception as exc:  # noqa: BLE001
            out.append(f"{flag}: Streamlit's own parser rejects the value ({type(exc).__name__}: {one_line(str(exc))})")
            continue
        if name.strip("'\"") != font_name:
            out.append(f"{flag}: Streamlit reads the family as {name!r}, not {font_name!r}")
        if url != fonts_url:
            out.append(f"{flag}: Streamlit reads the source URL as {url!r}, not the GOOGLE_FONTS_CSS_URL this project "
                       "builds - everything after the first colon is the URL, so a fallback stack appended to the flag "
                       "ends up inside it and Google Fonts serves the face without `display: swap` (L-114)")
        elif "display=swap" not in (url or ""):
            out.append(f"{flag}: the URL Streamlit would fetch has no `display=swap` - up to ~3 s of invisible text")
    code = args.get("--theme.codeFont")
    if code != code_stack:
        out.append(f"--theme.codeFont is {code!r}, not erp/typography.py CODE_STACK")
    elif _parse_font_config(code, "--theme.codeFont")[1] is not None:
        out.append("--theme.codeFont: Streamlit reads a source URL out of the monospace stack - it must stay a plain "
                   "CSS font list")
    return out, ""


def typography_files() -> list[str]:
    """TRACKED text files only (L-124: an untracked scratch note must never turn the gate red)."""
    cp = git("ls-files")
    if cp.returncode != 0:
        raise GateError(f"git ls-files failed - the typography check could not list the tracked files: {one_line(cp.stderr)}")
    return sorted({ln.strip() for ln in cp.stdout.splitlines() if ln.strip()})


def typography_problems() -> list[str]:
    """Light and fast (about 70 small tracked files): no retired font name outside a history line, and the three copies
    of the font token (shell.css, shell.html, erp/typography.py) still agree; the weights downloaded are the weights
    painted (L-113); every number that animates keeps tabular figures (L-112); and the Streamlit --theme.* font flags
    parse, through Streamlit's own parser, to exactly the URL this project builds (L-114 / L-124)."""
    problems: list[str] = []
    TYPOGRAPHY_NOTES.clear()
    for f in typography_files():
        p = ROOT / f
        if p.suffix.lower() in _BINARY_EXT or not p.is_file():
            continue
        try:
            if p.stat().st_size > 2_000_000:
                continue
            text = p.read_text(encoding="utf-8", errors="replace")
        except OSError:
            continue
        if not _OLD_FONTS.search(text):
            continue                                   # fast path: whole-file test first, only then line by line
        for i, line in enumerate(text.splitlines(), 1):
            why = old_font_problem(f, line)
            if why:
                problems.append(f"{f}:{i}: {why}")
    try:
        from erp import typography
    except Exception as exc:  # noqa: BLE001
        return problems + [f"erp/typography.py cannot be imported: {type(exc).__name__}"]
    if typography.FONT_NAME != "Montserrat":
        problems.append(f"erp/typography.py: FONT_NAME is {typography.FONT_NAME!r}, expected 'Montserrat' "
                        "(changing the project font means editing this gate too - see README, Typography)")
    css = ROOT / "desktop" / "static" / "shell.css"
    html = ROOT / "desktop" / "static" / "shell.html"
    for p in (css, html):
        if not p.is_file():
            problems.append(f"{rel(p)}: missing - the ERP Desk font tokens cannot be checked")
    if css.is_file() and html.is_file():
        problems += font_consistency_problems(css.read_text(encoding="utf-8", errors="replace"),
                                              html.read_text(encoding="utf-8", errors="replace"),
                                              typography.FONT_NAME, typography.FONT_STACK, typography.CODE_STACK,
                                              typography.GOOGLE_FONTS_CSS_URL)
    # the .bat launcher carries a literal copy of the Streamlit font flags (a .bat cannot import Python)
    want_list = typography.streamlit_theme_args()
    want = {want_list[i]: want_list[i + 1] for i in range(0, len(want_list), 2)}
    for bat in ("run_script_center.bat",):
        bp = ROOT / bat
        if not bp.is_file():
            problems.append(f"{bat}: missing")
            continue
        got = dict(re.findall(r'(--theme\.(?:font|headingFont|codeFont))\s+"([^"]*)"', bp.read_text(encoding="utf-8", errors="replace")))
        for flag, value in want.items():
            if got.get(flag) != value:
                problems.append(f"{bat}: {flag} differs from erp/typography.py streamlit_theme_args() - keep the two copies identical")
    # ... and the values themselves go through Streamlit's own parser, because two equal copies can be wrong together
    # (L-124): this is the check that catches a fallback stack swallowed into the Google Fonts URL (L-114).
    flag_problems, flag_note = streamlit_font_flag_problems(want_list, typography.FONT_NAME,
                                                            typography.GOOGLE_FONTS_CSS_URL, typography.CODE_STACK)
    problems += flag_problems
    if flag_note:
        TYPOGRAPHY_NOTES.append(flag_note)

    # the weights that are downloaded (L-113) and the numbers that must not jitter (L-112)
    sources: dict[str, str] = {}
    for f in FONT_CSS_SOURCES:
        p = ROOT / f
        try:
            sources[f] = p.read_text(encoding="utf-8", errors="replace")
        except OSError:
            problems.append(f"{f}: missing - the font-weight / tabular-figure scan could not read it")
    problems += tabular_number_problems(sources)
    if ".tnum" not in sources.get("desktop/static/shell.css", ""):
        problems.append("desktop/static/shell.css: the `.tnum` utility is gone - hand-written markup has no way left "
                        "to ask for tabular figures (L-112)")
    js = ROOT / "desktop" / "static" / "shell.js"
    problems += font_weight_problems(typography.FONT_WEIGHTS, sources,
                                     js.read_text(encoding="utf-8", errors="replace") if js.is_file() else "")
    return problems


# The two reusable skills must describe the repository as it IS (they rotted once: a deleted file kept being recommended).
SKILL_FILES = (".claude/skills/dashboard-craft/SKILL.md", ".claude/skills/cinematic-3d-flow/SKILL.md")
# A backticked path in a skill that is deliberately not a real file (an illustrative name for ANOTHER project). Keep it tiny.
SKILL_PATH_ALLOW = frozenset({"x_data.py", "x.js", "x.css", "x_scenarios.py"})
# Folders a skill may abbreviate a path from ("static/flow.js", "shell.css"): a path counts as real if it resolves under any.
SKILL_PATH_ROOTS = ("", "desktop", "desktop/static", "desktop/static/flow3d", "erp", "tests", "dashboard", ".claude")
_SKILL_PATH_EXT = (".py", ".js", ".css", ".html", ".md", ".bat", ".sql", ".ps1", ".json")
_BACKTICK = re.compile(r"`([^`\n]+)`")
_PATH_CHARS = re.compile(r"[\w./-]+")


def skill_path_problems(text: str, root: Path = ROOT, allow: frozenset[str] = SKILL_PATH_ALLOW) -> list[str]:
    """Backticked repository paths in a skill's prose that do not exist on disk (fenced code blocks are examples: skipped).

    A token counts as a path when it is made only of path characters and ends in a known source extension (so a token
    with a placeholder such as `<x>` or `*` is never one). It exists if it resolves under the repo root or one of
    SKILL_PATH_ROOTS. `allow` is the explicit list of illustrative names.
    """
    bad: list[str] = []
    fenced = False
    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            fenced = not fenced
            continue
        if fenced:
            continue
        for tok in _BACKTICK.findall(line):
            tok = tok.strip()
            if not _PATH_CHARS.fullmatch(tok) or not tok.endswith(_SKILL_PATH_EXT) or tok.startswith((".", "/")):
                continue
            if tok in allow or any((root / base / tok).exists() for base in SKILL_PATH_ROOTS):
                continue
            bad.append(tok)
    return list(dict.fromkeys(bad))


def check_agent_system() -> tuple[str, str, list[str]]:
    missing = [f for f in AGENT_FILES if not (ROOT / f).is_file()]
    if missing:
        return FAIL, f"missing: {', '.join(missing)}", missing
    problems: list[str] = []
    notes: list[str] = []
    empty = [f for f in AGENT_FILES if (ROOT / f).stat().st_size == 0]
    problems += [f"{f}: empty file" for f in empty]

    def bad_bytes(data: bytes) -> str | None:
        if b"\r" in data:
            return f"{data.count(b'\r')} CR byte(s) (CRLF line endings)"
        ctl = sorted({b for b in data if b < 32 and b not in (9, 10)})
        if ctl:
            return "control character(s) " + ", ".join(f"0x{b:02x}" for b in ctl)
        return None

    targets = list(NO_CR_FILES) + [rel(p) for p in sorted((ROOT / ".claude" / "workflows").glob("*.js"))]
    for f in dict.fromkeys(targets):
        p = ROOT / f
        if p.is_file():
            why = bad_bytes(p.read_bytes())
            if why:
                problems.append(f"{f}: {why} - rewrite it with LF only (L-066 / L-068)")
    # the other pinned files: report, do not fail (git normalises them; only the two above break tooling)
    for f in [f for f in AGENT_FILES if f.endswith(".md") and f.startswith(".claude/") or f == "Law.md"]:
        if f not in NO_CR_FILES and (ROOT / f).is_file() and b"\r" in (ROOT / f).read_bytes():
            notes.append(f"note: {f} has CRLF line endings (pinned to LF by .gitattributes)")

    lessons = (ROOT / ".claude/lessons.md").read_text(encoding="utf-8", errors="replace")
    ids = re.findall(r"^- `(L-\d{3})\b", lessons, flags=re.M)
    dups = sorted({i for i in ids if ids.count(i) > 1})
    if not ids:
        problems.append(".claude/lessons.md: no `L-NNN` lessons found - is the file damaged?")
    if dups:
        problems.append(f".claude/lessons.md: duplicate lesson id(s) {', '.join(dups)} (two agents appended the same number)")
    for f in (".claude/agents/erp-builder.md", ".claude/agents/erp-reviewer.md", ".claude/agents/erp-auditor.md"):
        head = (ROOT / f).read_text(encoding="utf-8", errors="replace").lstrip("﻿").splitlines()[:6]
        if not (head and head[0].strip() == "---" and any(h.startswith("name:") for h in head)):
            problems.append(f"{f}: front matter (--- / name:) missing")
    for f in SKILL_FILES:
        sp = ROOT / f
        if not sp.is_file():
            problems.append(f"{f}: skill file missing")
            continue
        ghosts = skill_path_problems(sp.read_text(encoding="utf-8", errors="replace"))
        if ghosts:
            problems.append(f"{f}: names repository path(s) that do not exist: {', '.join(ghosts[:5])} "
                            f"(fix the skill, or list an illustrative name in SKILL_PATH_ALLOW)")
    for f, skills in ((".claude/agents/erp-builder.md", ("dashboard-craft", "cinematic-3d-flow")),
                      (".claude/agents/erp-reviewer.md", ("dashboard-craft",)), (".claude/agents/erp-auditor.md", ("dashboard-craft",))):
        body = (ROOT / f).read_text(encoding="utf-8", errors="replace")
        problems += [f"{f}: does not point at the {k} skill" for k in skills if k not in body]
    problems += typography_problems()
    notes += TYPOGRAPHY_NOTES
    if (ROOT / ".claude" / "PAUSE").exists():
        notes.append("note: .claude/PAUSE exists (the conductor must not start new work; this does not fail the smoke check)")
    if problems:
        return FAIL, problems[0] + (" (+more)" if len(problems) > 1 else ""), problems + notes
    return PASS, (f"{len(AGENT_FILES)} agent-system files present; lessons.md + {len(list((ROOT / '.claude' / 'workflows').glob('*.js')))} "
                  f"workflow script(s) are CR-free; {len(ids)} lessons, ids unique; typography: no retired font on a live "
                  f"tracked line, shell.css / shell.html / erp/typography.py agree (Streamlit's own parser included), every "
                  f"downloaded weight is painted, {len(TABULAR_NUMBER_RULES)} moving-number rules still tabular; "
                  f"{len(SKILL_FILES)} skills name only paths that exist"), notes


# ============================================================================ check 7: nothing left behind
def check_left_behind(state: dict) -> tuple[str, str, list[str]]:
    if state.get("setup_error"):
        return FAIL, str(state["setup_error"]), []
    problems: list[str] = []
    new_files = sorted(worktree_files() - state["files_before"])
    for f in new_files:
        if not f.endswith("/"):
            problems.append(f"new file left in the work tree: {f}")
    port = state.get("server_port")
    if port and port_open(port):
        problems.append(f"port {port} is still listening after the server was stopped")
    extra = [t.name for t in threading.enumerate() if t is not threading.main_thread() and not t.daemon and t.is_alive()]
    if extra:
        problems.append("non-daemon threads still running: " + ", ".join(extra))
    added = [h for h in logging.getLogger().handlers if h not in state["root_handlers_before"]]
    if added:
        problems.append(f"{len(added)} logging handler(s) added to the root logger")
    if problems:
        return FAIL, problems[0] + (" (+more)" if len(problems) > 1 else ""), problems
    return PASS, "no new files, no listening port, no stray threads or logging handlers", []


# ============================================================================ check 8: branch discipline
PROTECTED_BRANCHES = {"main", "master"}


def branch_discipline_problem(branch: str, modified: set[str]) -> str | None:
    """Reason this work breaks "never edit on main" (Law.md rule 1 / lesson L-053), else None.

    Pure so tests/smoke_selftest.py can drive every combination without checking anything out. `modified` = tracked
    files with staged or unstaged changes; the owner's own dirty files (Law.md rule 9) do not count as agent work.
    """
    head = (branch or "").lstrip("(").split()
    if not head or head[0].lower() not in PROTECTED_BRANCHES:
        return None
    agent_edits = sorted(modified - OWNER_FILES)
    if not agent_edits:
        return None
    return (f"the branch IS '{branch}' and {len(agent_edits)} tracked file(s) are modified there "
            f"({', '.join(agent_edits[:4])}{' ...' if len(agent_edits) > 4 else ''}) - never edit on main "
            f"(Law.md rule 1 / L-053): move the work to a branch `agent/YYYYMMDD-slug`")


def check_branch_discipline() -> tuple[str, str, list[str]]:
    branch = current_branch()
    if branch.startswith("(unknown"):
        return FAIL, f"the current branch cannot be determined: {branch}", []
    if branch == "(not a git repository)":
        return FAIL, "not a git repository: branch discipline cannot be checked", []
    modified: set[str] = set()
    for args in (("diff", "--name-only", "--no-renames", "HEAD"), ("diff", "--cached", "--name-only", "--no-renames")):
        r = git(*args)
        if r.returncode != 0:
            return FAIL, f"git {args[0]} failed while listing modified files: {one_line(r.stderr) or 'no commit yet?'}", []
        modified |= {ln.strip() for ln in r.stdout.splitlines() if ln.strip()}
    why = branch_discipline_problem(branch, modified)
    details = [f"modified tracked file(s): {', '.join(sorted(modified)) or 'none'}"]
    if why:
        return FAIL, why, details
    if branch.lower() in PROTECTED_BRANCHES:
        left = sorted(modified & OWNER_FILES)
        return PASS, (f"on '{branch}' with no agent edits" + (f" (only the owner's {', '.join(left)})" if left else "")), details
    return PASS, f"on '{branch}', not a protected branch ({len(modified)} tracked file(s) modified)", details


# ============================================================================ check 9: gate self-test
def check_gate_selftest(args) -> tuple[str, str, list[str]]:
    """Run tests/smoke_selftest.py: one scenario per protection this gate makes (see its module docstring).

    The module object is handed over rather than imported by name: this file runs as `__main__` under `-m`, so
    `import tests.smoke` inside the self-test would execute a SECOND copy of the gate.
    """
    from tests import smoke_selftest

    rows = smoke_selftest.run(sys.modules[__name__], base_ref=args.base)
    bad = [r for r in rows if not r[1]]
    details = [f"{'ok  ' if ok else 'FAIL'} {name}" + (f"  <- {one_line(str(why), 150)}" if why else "") for name, ok, why in rows]
    if bad:
        return FAIL, (f"{len(bad)} of {len(rows)} gate protection(s) are broken: {bad[0][0]} "
                      f"({one_line(str(bad[0][2]), 110)})"), details
    groups = sorted({n.split(":")[0] for n, _o, _w in rows})
    return PASS, f"{len(rows)} scenarios over {len(groups)} protection(s) pass: {', '.join(groups)}", details


# ============================================================================ runner / output
CHECKS = [
    (1, "Imports & compile", lambda a, s: check_imports()),
    (2, "Database (read-only)", lambda a, s: check_database()),
    (3, "Data Flow map sync", lambda a, s: check_flow_map()),
    (4, "ERP Desk server + pages", lambda a, s: check_server(s)),
    (5, "Secrets & stray files", lambda a, s: check_secrets_and_strays(a.base)),
    (6, "Agent system files + typography", lambda a, s: check_agent_system()),
    (7, "Nothing left behind", lambda a, s: check_left_behind(s)),
    (8, "Branch discipline", lambda a, s: check_branch_discipline()),
    (9, "Gate self-test", lambda a, s: check_gate_selftest(a)),
]


def run_check(n: int, name: str, fn, args, state) -> Result:
    res = Result(n, name)
    RESULTS.append(res)
    t0 = time.monotonic()
    try:
        res.status, res.reason, res.details = fn(args, state)
    except GateError as exc:  # a missing / broken prerequisite (git, the repo, the base branch): a clean FAIL line (L-076)
        res.status, res.reason, res.details = FAIL, one_line(str(exc), 200), []
    except Exception as exc:  # noqa: BLE001 - a crash inside a check is a FAIL line, never a traceback on the console
        res.status = FAIL
        res.reason = f"the check crashed: {type(exc).__name__}: {one_line(str(exc), 140)}"
        res.details = [scrub(ln) for ln in traceback.format_exc().splitlines()[-12:]]
    if n == 2:
        state["db_ok"] = res.status == PASS
    res.seconds = time.monotonic() - t0
    return res


def _use_color() -> bool:
    if os.environ.get("NO_COLOR") or not sys.stdout.isatty():
        return False
    if os.name == "nt":
        os.system("")          # enables ANSI escape processing in the Windows console
    return True


def render(results: list[Result], branch: str, total: float, verbose: bool, partial: bool, note: str = "") -> str:
    color = _use_color()

    def paint(s: str, status: str) -> str:
        if not color:
            return s
        return f"\033[{'32' if status == PASS else '31;1'}m{s}\033[0m"

    width = max(80, min(shutil.get_terminal_size((150, 20)).columns, 160))
    name_w = max(len(r.name) for r in results) if results else 10
    reason_w = max(30, width - name_w - 22)
    lines = [f"ERP smoke check - branch {branch} - {time.strftime('%Y-%m-%d %H:%M:%S')}", ""]
    lines.append(f" #  {'Check'.ljust(name_w)}  Result  Time   Reason")
    lines.append(f"--- {'-' * name_w}  ------  -----  {'-' * min(reason_w, 60)}")
    for r in results:
        reason = one_line(r.reason, reason_w)
        lines.append(f" {r.n}  {r.name.ljust(name_w)}  {paint(r.status.ljust(6), r.status)}  {r.seconds:4.1f}s  {reason}")
    shown = [r for r in results if r.details and (verbose or r.status != PASS or any(d.startswith("note") or "owner" in d for d in r.details))]
    if shown:
        lines.append("")
        for r in shown:
            lines.append(f"[{r.n}] {r.name} - details")
            cap = None if verbose else 12
            flat = [ln.rstrip() for d in r.details for ln in scrub(d).splitlines() if ln.strip() and "sqlalche.me" not in ln]
            for d in (flat if cap is None else flat[:cap]):
                lines.append("    " + d)
            if cap is not None and len(flat) > cap:
                lines.append(f"    ... {len(flat) - cap} more (run with --verbose)")
    failed = [r for r in results if r.status != PASS]
    lines.append("")
    if note:
        lines.append(f"RESULT: FAIL - {note}")
    elif failed:
        lines.append(f"RESULT: FAIL - {len(failed)} of {len(results)} check(s) failed ({', '.join(str(r.n) for r in failed)}) in {total:.1f} s. Do NOT merge.")
    elif partial:
        lines.append(f"RESULT: PASS (partial: only checks {', '.join(str(r.n) for r in results)} ran) in {total:.1f} s. This is NOT the merge gate: run without --only.")
    else:
        lines.append(f"RESULT: PASS - all {len(results)} checks passed in {total:.1f} s. The smoke gate is green.")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m tests.smoke", description="Read-only merge-gate smoke check (see the module docstring).")
    ap.add_argument("--base", default="main", help="branch to diff against for the secret / stray-file scan (default: main)")
    ap.add_argument("--only", default="", help="comma separated check numbers, e.g. 3,5 (a subset is never a valid merge gate)")
    ap.add_argument("--verbose", "-v", action="store_true", help="print the details of every check, not only the failing ones")
    ap.add_argument("--timeout", type=int, default=120, help="give up (exit 2) after this many seconds (default 120)")
    args = ap.parse_args(argv)

    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")         # cp1252 console (L-011)
    if args.verbose:
        h = logging.StreamHandler()
        h.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
        log.addHandler(h)
        log.setLevel(logging.DEBUG)
        log.propagate = False

    wanted = {int(x) for x in re.findall(r"\d+", args.only)}
    selected = [c for c in CHECKS if not wanted or c[0] in wanted]
    if not selected:
        print(f"No check matches --only {args.only!r} (valid: 1-{len(CHECKS)})")
        return 2
    partial = len(selected) < len(CHECKS)
    try:
        os.chdir(ROOT)
    except OSError as exc:
        print(f"RESULT: FAIL - the project folder {ROOT} cannot be entered: {type(exc).__name__}: {one_line(str(exc), 120)}")
        return 2
    started = time.monotonic()
    branch = current_branch()

    def on_timeout() -> None:
        for r in RESULTS:
            if r.reason == "did not run":
                r.status, r.reason = FAIL, f"timed out after {args.timeout} s (still running)"
        print(render(RESULTS, branch, time.monotonic() - started, args.verbose, partial,
                     note=f"the smoke run exceeded {args.timeout} s and was aborted"), file=sys.__stdout__, flush=True)   # sys.stdout may be redirected by an import in progress
        os._exit(2)

    watchdog = threading.Timer(args.timeout, on_timeout)
    watchdog.daemon = True
    watchdog.start()

    # set-up code runs OUTSIDE run_check's try/except, so it must never raise either (L-076)
    state: dict = {"files_before": set(), "root_handlers_before": list(logging.getLogger().handlers), "setup_error": None}
    try:
        state["files_before"] = worktree_files()
    except GateError as exc:
        state["setup_error"] = f"the work-tree snapshot could not be taken: {one_line(str(exc), 140)}"
    except OSError as exc:  # noqa: BLE001
        state["setup_error"] = f"the work-tree snapshot could not be taken: {type(exc).__name__}: {one_line(str(exc), 120)}"
    live = sys.stdout.isatty()
    for n, name, fn in selected:
        if live:
            print(f"  [{n}/{len(CHECKS)}] {name} ...", flush=True)
        r = run_check(n, name, fn, args, state)
        log.debug("check %d %s -> %s (%.1fs)", n, name, r.status, r.seconds)
    watchdog.cancel()
    if live:
        print()
    print(render(RESULTS, branch, time.monotonic() - started, args.verbose, partial))
    if any(r.status != PASS for r in RESULTS):
        return 1
    return 3 if partial else 0


if __name__ == "__main__":
    sys.exit(main())
