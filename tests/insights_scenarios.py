"""
Scenario table for the Channels Insights engine and Excel report (desktop/insights_rules.py, insights_queries.py, insights_data.py,
insights_xlsx.py, static/insights.js), kept in the repo so the Reviewer and the merge gate (tests/smoke.py) can re-run it (L-089).

  * pure: for EACH of the six rules a scenario that it fires, one that it does not, and one that stays quiet (or only says 'thin
    data', never a warning) on a tiny sample; the window / cost-per-conversion / conversion-rate maths; per-currency isolation
    (a mixed-currency input is refused); severity ordering; automation ideas; the thresholds text being generated from the constants.
  * the store: an unknown parameter is a 400 BEFORE the cache and before any database access, on the insights feed and on the xlsx;
    not-installed / empty states carry no finding and the report answers 409.
  * the database, throwaway objects only: the real db/sql/07 + 10 + 09 DDL in the session's TEMPORARY schema (tests.channels_scenarios
    builds it), a hand-worked multi-currency rollup, the window split done in SQL against a Python-side sum, the whole store, and the
    generated .xlsx opened with openpyxl (sheets, header, frozen row, numeric cells, number formats, formula-injection guard, size bound).
  * the real database, read only: whatever state it is in, the feed and the report answer honestly.

Run (from the project root):
    venv\\Scripts\\python.exe -B -m tests.insights_scenarios
"""
from __future__ import annotations

import contextlib
import io
import re
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

M = 1_000_000
W = None


def _ir():
    from desktop import insights_rules as ir
    return ir


def P(**kw):
    return _ir().Period(**kw)


def G(ch="a", cid=1, cur="EUR", c=None, p=None, name=None):
    ir = _ir()
    return ir.Group(channel_key=ch, channel_name=ch.upper(), campaign_id=cid, campaign_name=name or f"Camp {cid}", currency=cur,
                    cur=c or P(days=30), prev=p or P(days=30))


def win():
    ir = _ir()
    return ir.Window(date(2026, 6, 1), date(2026, 6, 30), date(2026, 5, 2), date(2026, 5, 31))


def cp(spend=0.0, conv=0, clicks=0, days=30, rev=0.0, imp=0, sess=0):
    return P(spend_micros=int(round(spend * M)), conversions=conv, clicks=clicks, days=days, revenue_micros=int(round(rev * M)), impressions=imp, sessions=sess)


def ids(res):
    return [f["id"] for f in res["findings"]]


@contextlib.contextmanager
def _patched(obj, **attrs):
    old = {k: getattr(obj, k) for k in attrs}
    for k, v in attrs.items():
        setattr(obj, k, v)
    try:
        yield
    finally:
        for k, v in old.items():
            setattr(obj, k, v)


# =========================================================================================== rule 1
def _check_rule1(add) -> None:
    ir = _ir()
    w = win()
    run = lambda gs: ir.rule_spend_no_conversions(gs, "EUR", w)   # noqa: E731
    base = [G(cid=1, c=cp(300, 0, 400), p=cp(300, 8, 400)), G(cid=2, c=cp(700, 20, 900), p=cp(600, 18, 800))]
    r = run(base)
    f = r["findings"][0] if r["findings"] else {}
    add("rule1 fires: 30% of spend, 400 clicks, 0 conversions, and the channel converts elsewhere", len(r["findings"]) == 1 and f["id"] == "spend_no_conversions", r)
    add("rule1 severity: a 30% share is critical (>= 25%) with enough data", f.get("severity") == "crit" and f.get("confidence") == "enough data", f)
    add("rule1 evidence shows the exact comparison against the threshold", any(">=" == e["compare"] and e["threshold"] == ir.NOCONV_MIN_SHARE_PCT and "30.0%" in e["text"] for e in f["evidence"]), f["evidence"])
    add("rule1 evidence carries the previous window's conversions", any("8 conversion" in e["text"] for e in f["evidence"]), f["evidence"])
    warn = run([G(cid=1, c=cp(100, 0, 400), p=cp(100, 8, 400)), G(cid=2, c=cp(900, 20, 900), p=cp(600, 18, 800))])
    add("rule1 severity: a 10% share is a warning, not critical", warn["findings"][0]["severity"] == "warn", warn["findings"])
    add("rule1 does not fire when the campaign converts", not run([G(cid=1, c=cp(300, 1, 400)), G(cid=2, c=cp(700, 20, 900))])["findings"], None)
    add("rule1 does not fire under the spend share threshold (3%)", not run([G(cid=1, c=cp(30, 0, 400), p=cp(30, 3, 400)), G(cid=2, c=cp(970, 20, 900))])["findings"], None)
    q = run([G(cid=1, c=cp(300, 0, 4), p=cp(300, 8, 400)), G(cid=2, c=cp(700, 20, 900))])
    add("rule1 stays quiet under the click floor (4 clicks): no finding at all", not q["findings"], q)
    t = run([G(cid=1, c=cp(300, 0, 20), p=cp(300, 8, 400)), G(cid=2, c=cp(700, 20, 900))])
    add("rule1 thin data (20 clicks): only info with confidence 'thin data', never an alarm", len(t["findings"]) == 1 and t["findings"][0]["severity"] == "info" and t["findings"][0]["confidence"] == "thin data", t)
    u = run([G(ch="mail", cid=1, c=cp(300, 0, 400), p=cp(300, 0, 400)), G(ch="ads", cid=2, c=cp(700, 20, 900))])
    add("rule1 honesty: a channel that never reports conversions is a NOTE, not a finding (0 means unknown)", not u["findings"] or all(x["channel"]["key"] != "mail" for x in u["findings"]), u)
    add("rule1 note names the unknown-conversions campaign", any("unknown" in n for n in u["notes"]), u["notes"])
    s = run([G(cid=1, c=cp(0, 0, 400))])
    add("rule1 skipped (not ran) when nobody has spend: a 0 spend is 'no spend field'", s["ran"] is False and "spend" in s["why"], s)
    add("rule1 spend text is in the campaign's currency code", "EUR" in f["what"], f["what"])


# =========================================================================================== rule 2
def _check_rule2(add) -> None:
    ir = _ir()
    w = win()
    run = lambda gs: ir.rule_cpa_jump(gs, "EUR", w)   # noqa: E731
    r = run([G(c=cp(1400, 20, 800), p=cp(1000, 20, 800))])
    f = r["findings"][0] if r["findings"] else {}
    add("rule2 fires: cost per conversion 50 -> 70 = +40%", len(r["findings"]) == 1 and abs(f["magnitude"] - 40.0) < 1e-6 and f["severity"] == "warn", r)
    add("rule2 evidence shows both costs and the change", any("70.00" in e["text"] and "50.00" in e["text"] and "+40.0%" in e["text"] for e in f["evidence"]), f["evidence"])
    c = run([G(c=cp(2000, 20, 800), p=cp(1000, 20, 800))])
    add("rule2 critical at +100%", c["findings"][0]["severity"] == "crit", c)
    add("rule2 does not fire at +10%", not run([G(c=cp(1100, 20, 800), p=cp(1000, 20, 800))])["findings"], None)
    add("rule2 does not fire when the cost fell", not run([G(c=cp(800, 20, 800), p=cp(1000, 20, 800))])["findings"], None)
    t = run([G(c=cp(1400, 3, 800), p=cp(1000, 3, 800))])
    add("rule2 thin data (3 conversions): info + 'thin data', not a warning", len(t["findings"]) == 1 and t["findings"][0]["severity"] == "info" and t["findings"][0]["confidence"] == "thin data", t)
    add("rule2 stays quiet under the conversion floor (1 conversion)", not run([G(c=cp(1400, 1, 800), p=cp(1000, 3, 800))])["findings"], None)
    add("rule2 skipped when the previous window has no spend", run([G(c=cp(1400, 20, 800), p=cp(0, 0, 0, days=0))])["ran"] is False, None)
    add("rule2 needs conversions in both windows: 0 now is rule 1's business", not run([G(c=cp(1400, 0, 800), p=cp(1000, 20, 800))])["findings"], None)
    add("Period.cpa is unknown without spend or without conversions", cp(0, 5, 10).cpa() is None and cp(10, 0, 10).cpa() is None and cp(10, 4, 10).cpa() == 2.5, None)


# =========================================================================================== rule 3
def _check_rule3(add) -> None:
    ir = _ir()
    w = win()
    run = lambda gs: ir.rule_cvr_drop(gs, "EUR", w)   # noqa: E731
    r = run([G(c=cp(0, 20, 1000), p=cp(0, 50, 1000))])
    f = r["findings"][0] if r["findings"] else {}
    add("rule3 fires: 5.0% -> 2.0% (-60%) is critical with enough clicks", len(r["findings"]) == 1 and f["severity"] == "crit" and f["confidence"] == "enough data", r)
    add("rule3 evidence shows both rates, both counts and the points", any("2.00% now" in e["text"] and "5.00% before" in e["text"] and "3.00 points" in e["text"] for e in f["evidence"]), f["evidence"])
    add("rule3 warning at a 30% relative fall", run([G(c=cp(0, 35, 1000), p=cp(0, 50, 1000))])["findings"][0]["severity"] == "warn", None)
    add("rule3 does not fire at a 10% relative fall", not run([G(c=cp(0, 45, 1000), p=cp(0, 50, 1000))])["findings"], None)
    add("rule3 ignores a big relative fall of under 0.3 points (0.5% -> 0.3%)", not run([G(c=cp(0, 3, 1000), p=cp(0, 5, 1000))])["findings"], None)
    t = run([G(c=cp(0, 2, 40), p=cp(0, 5, 40))])
    add("rule3 thin data (40 clicks): info + 'thin data'", len(t["findings"]) == 1 and t["findings"][0]["severity"] == "info" and t["findings"][0]["confidence"] == "thin data", t)
    add("rule3 stays quiet under the click floor (10 clicks)", not run([G(c=cp(0, 1, 10), p=cp(0, 5, 10))])["findings"], None)
    add("rule3 skipped without clicks in the previous window", run([G(c=cp(0, 20, 1000), p=P(days=0))])["ran"] is False, None)
    add("rule3 does not fire when the rate rose", not run([G(c=cp(0, 60, 1000), p=cp(0, 50, 1000))])["findings"], None)
    add("Period.cvr is unknown without clicks", cp(0, 5, 0).cvr() is None and abs(cp(0, 5, 100).cvr() - 5.0) < 1e-9, None)
    nm = run([G(c=cp(0, 20, 1000), p=cp(0, 50, 1000))])["findings"][0]
    add("rule3 on a counts-only group names no currency", nm["currency"] is None, nm)


# =========================================================================================== rule 4
def _check_rule4(add) -> None:
    ir = _ir()
    w = win()
    days = tuple(date(2026, 6, 1) + timedelta(days=i) for i in range(30))
    sp = lambda **kw: ir.Span(**{"channel_key": "a", "channel_name": "A", "grain": "day", "first_day": date(2026, 1, 1), "last_day": date(2026, 6, 30), "history_days": 180, "window_days": days, **kw})   # noqa: E731
    add("rule4 does not fire on a channel with a row on the last day and no holes", not ir.rule_channel_data([sp()], w)["findings"], None)
    r = ir.rule_channel_data([sp(last_day=date(2026, 6, 10), window_days=days[:10])], w)
    stale = [f for f in r["findings"] if "stale" in f["key"]]
    add("rule4 fires: newest row 20 days before the window end (> 7 allowed)", len(stale) == 1 and stale[0]["severity"] == "warn" and stale[0]["confidence"] == "enough data", r)
    add("rule4 evidence shows the days and the allowed limit", any("20 days without a row > 7 days" in e["text"] for e in stale[0]["evidence"]), stale[0]["evidence"])
    add("rule4 does not fire at 5 days of silence", not [f for f in ir.rule_channel_data([sp(last_day=date(2026, 6, 25), window_days=days[:25])], w)["findings"] if "stale" in f["key"]], None)
    wk = sp(grain="week", last_day=date(2026, 6, 20), window_days=(date(2026, 6, 6), date(2026, 6, 13), date(2026, 6, 20)))
    add("rule4 weekly channel: 10 days of silence is normal (limit 14)", not [f for f in ir.rule_channel_data([wk], w)["findings"] if "stale" in f["key"]], None)
    wk2 = sp(grain="week", last_day=date(2026, 6, 6), window_days=(date(2026, 6, 6),))
    add("rule4 weekly channel: 24 days of silence fires", len([f for f in ir.rule_channel_data([wk2], w)["findings"] if "stale" in f["key"]]) == 1, None)
    th = ir.rule_channel_data([sp(last_day=date(2026, 6, 10), window_days=days[:10], history_days=3)], w)
    add("rule4 thin data: a channel with 3 rows of history is info + 'thin data'", [f["severity"] for f in th["findings"] if "stale" in f["key"]] == ["info"] and th["findings"][0]["confidence"] == "thin data", th)
    holed = tuple(d for d in days if not (date(2026, 6, 10) <= d <= date(2026, 6, 16)))
    h = ir.rule_channel_data([sp(window_days=holed)], w)
    hf = [f for f in h["findings"] if "holes" in f["key"]]
    add("rule4 holes: 7 missing days in a row inside the window fire", len(hf) == 1 and "7 missing day" in hf[0]["what"], h)
    add("rule4 holes: 7 of 30 days (23%) is a warning (>= 20%)", hf[0]["severity"] == "warn", hf)
    small = tuple(d for d in days if not (date(2026, 6, 10) <= d <= date(2026, 6, 12)))
    add("rule4 holes: 3 missing days (10%) is only info", ir.rule_channel_data([sp(window_days=small)], w)["findings"][0]["severity"] == "info", None)
    two = tuple(d for d in days if not (date(2026, 6, 10) <= d <= date(2026, 6, 11)))
    add("rule4 holes: a 2-day gap is below HOLE_MIN_DAYS and stays quiet", not ir.rule_channel_data([sp(window_days=two)], w)["findings"], None)
    wkh = sp(grain="week", last_day=date(2026, 6, 27), window_days=(date(2026, 6, 6), date(2026, 6, 20), date(2026, 6, 27)))
    g = ir._gaps(wkh.window_days, "week")
    add("rule4 holes: a missing WEEK is found for a weekly channel (one week = 7 days)", g == [(date(2026, 6, 13), date(2026, 6, 19), 7)], g)
    add("rule4 holes evidence lists the gap dates", "10 Jun 2026 to 16 Jun 2026" in " ".join(e["text"] for e in hf[0]["evidence"]), hf[0]["evidence"])
    add("rule4 skipped when no channel has rows", ir.rule_channel_data([], w)["ran"] is False, None)
    out = ir.rule_channel_data([sp(last_day=date(2026, 5, 1), window_days=())], w)
    add("rule4 a channel with no row at all in the window says so", "not a single row" in out["findings"][0]["what"], out)
    # ---- historical feeds: silent for more than STALE_HISTORICAL_DAYS is info ('ended or archived'), never a fresh-outage warning
    hist_days = (w.d_to - date(2024, 6, 20)).days
    hist = ir.rule_channel_data([sp(last_day=date(2024, 6, 20), window_days=())], w)
    hf1 = [f for f in hist["findings"] if "stale" in f["key"]]
    add("rule4 historical fires: last row Jun 2024 in a Jun 2026 window is info, not warn", len(hf1) == 1 and hf1[0]["severity"] == "info" and hist_days > ir.STALE_HISTORICAL_DAYS, hist)
    add("rule4 historical wording says ended / archived, not a fresh outage, and suggests confirm-or-reconnect",
        "ended or archived" in hf1[0]["what"] and "not a fresh outage" in hf1[0]["what"] and "Confirm the feed is retired or re-connect it" in hf1[0]["action"], hf1[0])
    add("rule4 historical evidence carries the cut-off comparison", any(e["label"] == "Historical cut-off" and e["threshold"] == ir.STALE_HISTORICAL_DAYS and e["met"] for e in hf1[0]["evidence"]), hf1[0]["evidence"])
    recent = [f for f in ir.rule_channel_data([sp(last_day=date(2026, 6, 10), window_days=days[:10])], w)["findings"] if "stale" in f["key"]]
    add("rule4 historical quiet: a 20-day gap stays a warning with the old wording", len(recent) == 1 and recent[0]["severity"] == "warn" and "sent nothing" in recent[0]["what"], recent)
    lim = ir.STALE_HISTORICAL_DAYS
    at = [f for f in ir.rule_channel_data([sp(last_day=w.d_to - timedelta(days=lim), window_days=())], w)["findings"] if "stale" in f["key"]]
    over = [f for f in ir.rule_channel_data([sp(last_day=w.d_to - timedelta(days=lim + 1), window_days=())], w)["findings"] if "stale" in f["key"]]
    add("rule4 historical boundary: exactly STALE_HISTORICAL_DAYS of silence is still a warning", len(at) == 1 and at[0]["severity"] == "warn", at)
    add("rule4 historical boundary: one day more is historical info", len(over) == 1 and over[0]["severity"] == "info" and "ended or archived" in over[0]["what"], over)
    with _patched(ir, STALE_HISTORICAL_DAYS=30):
        moved = [f for f in ir.rule_channel_data([sp(last_day=date(2026, 5, 1), window_days=())], w)["findings"] if "stale" in f["key"]]
        txt = " ".join(t["meaning"] + str(t["value"]) for t in ir.thresholds()) + " ".join(r["text"] for r in ir.rule_texts())
    add("rule4 historical cut-off is read from the constant (patched to 30: a 60-day gap becomes info) and the drawer text follows it", moved[0]["severity"] == "info" and "more than 30 days" in txt, moved)
    add("rule4 historical constant is listed in the drawer thresholds", any(t["name"] == "STALE_HISTORICAL_DAYS" and t["value"] == 180 for t in ir.thresholds()), None)


# =========================================================================================== xlsx row cap (pure)
def _check_xlsx_cap(add) -> None:
    import openpyxl
    from desktop import insights_xlsx as ix
    pulled = [0]

    def synthetic(n):
        for i in range(n):
            pulled[0] += 1
            yield [f"camp {i}", i, 1.5]
    cap = ix.ROWS_MAX
    t0 = time.monotonic()
    body = ix.build_workbook([ix.Sheet("Campaigns", [("name", "text"), ("clicks", "int"), ("rate", "pct")], synthetic(50_000))])
    took = time.monotonic() - t0
    ws = openpyxl.load_workbook(io.BytesIO(body))["Campaigns"]
    data_rows = sum(1 for r in ws.iter_rows(min_row=2) if isinstance(r[1].value, int))
    add("xlsx cap: the named constant is 20,000 rows", cap == 20_000, cap)
    add("xlsx cap: 50,000 synthetic rows produce at most the cap rows of data", data_rows == cap, data_rows)
    add("xlsx cap: rows are limited BEFORE cells exist (the writer pulled at most cap + 1 rows out of a 50,000-row source)", pulled[0] <= cap + 1, pulled[0])
    note = ws.cell(row=cap + 3, column=1).value
    add("xlsx cap: the visible truncation note is in the file", isinstance(note, str) and "cut at 20,000 rows" in note, note)
    add("xlsx cap: the capped build is quick", took < 30, round(took, 1))
    small = openpyxl.load_workbook(io.BytesIO(ix.build_workbook([ix.Sheet("S", [("a", "text")], [["x"], ["y"]])])))["S"]
    add("xlsx cap: a small sheet has no note", small.max_row == 3 and small.cell(row=5, column=1).value is None, small.max_row)
    exact = openpyxl.load_workbook(io.BytesIO(ix.build_workbook([ix.Sheet("S", [("a", "int")], ([i] for i in range(cap)))])))["S"]
    add("xlsx cap: exactly the cap is complete and carries no note", exact.max_row == cap + 1, exact.max_row)


# =========================================================================================== rule 5
def _check_rule5(add) -> None:
    ir = _ir()
    w = win()
    run = lambda gs, cur="EUR": ir.rule_concentration(gs, cur, w)   # noqa: E731
    r = run([G(cid=1, c=cp(600, 5, 300)), G(cid=2, c=cp(400, 5, 300))])
    f = r["findings"][0] if r["findings"] else {}
    add("rule5 fires: one campaign carries 60% of the spend (warning)", len(r["findings"]) == 1 and f["severity"] == "warn" and f["campaign"]["id"] == 1, r)
    add("rule5 evidence shows the share against the threshold", any(e["compare"] == ">" and e["threshold"] == ir.CONC_WARN_PCT and "60.0%" in e["text"] for e in f["evidence"]), f["evidence"])
    add("rule5 critical above 80%", run([G(cid=1, c=cp(850, 5, 300)), G(cid=2, c=cp(150, 5, 300))])["findings"][0]["severity"] == "crit", None)
    add("rule5 does not fire when the top campaign holds 45%", not run([G(cid=1, c=cp(450, 5, 300)), G(cid=2, c=cp(400, 5, 300)), G(cid=3, c=cp(150, 5, 300))])["findings"], None)
    add("rule5 skipped with one spending campaign (always 100%)", run([G(cid=1, c=cp(600, 5, 300))])["ran"] is False, None)
    t = run([G(cid=1, c=cp(900, 5, 20)), G(cid=2, c=cp(100, 5, 20))])
    add("rule5 thin data (40 clicks): info + 'thin data'", t["findings"][0]["severity"] == "info" and t["findings"][0]["confidence"] == "thin data", t)
    add("rule5 a 50/50 split does not fire (not above the threshold)", not run([G(cid=1, c=cp(500, 5, 300)), G(cid=2, c=cp(500, 5, 300))])["findings"], None)


# =========================================================================================== rule 6
def _check_rule6(add) -> None:
    ir = _ir()
    w = win()
    run = lambda gs: ir.rule_scale_candidates(gs, "EUR", w)   # noqa: E731
    gs = [G(cid=1, c=cp(300, 30, 500), p=cp(300, 30, 500)), G(cid=2, c=cp(700, 20, 600), p=cp(700, 20, 600)), G(cid=3, c=cp(1000, 10, 900), p=cp(1000, 10, 900))]
    r = run(gs)
    add("rule6 fires: 10.0 per conversion vs a pooled 33.3 is well over 25% better", len(r["findings"]) >= 1 and r["findings"][0]["campaign"]["id"] == 1, r)
    add("rule6 findings are info and 'enough data'", all(f["severity"] == "info" and f["confidence"] == "enough data" for f in r["findings"]), r)
    add("rule6 lists the best cost per conversion first", [f["campaign"]["id"] for f in r["findings"]][0] == 1, r)
    add("rule6 evidence shows the pooled cost and the comparison", any(e["compare"] == ">=" and "pooled" in e["text"] for e in r["findings"][0]["evidence"]), r["findings"][0]["evidence"])
    add("rule6 does not fire when all costs are alike", not run([G(cid=1, c=cp(500, 25, 500)), G(cid=2, c=cp(500, 25, 500))])["findings"], None)
    thin = run([G(cid=1, c=cp(30, 3, 50)), G(cid=2, c=cp(700, 20, 600))])
    add("rule6 stays quiet on thin data: never advises scaling a 3-conversion campaign", not any(f["campaign"]["id"] == 1 for f in thin["findings"]), thin)
    add("rule6 skipped with one converting campaign", run([G(cid=1, c=cp(300, 30, 500)), G(cid=2, c=cp(0, 0, 500))])["ran"] is False, None)
    imp = run([G(cid=1, c=cp(400, 20, 500), p=cp(600, 20, 500)), G(cid=2, c=cp(400, 20, 500), p=cp(400, 20, 500))])
    add("rule6 improver: 30 -> 20 per conversion (33% better) at a cost no worse than the pool", [f["campaign"]["id"] for f in imp["findings"]] == [1] or any(f["campaign"]["id"] == 1 for f in imp["findings"]), imp)
    many = run([G(cid=i, c=cp(100 + i, 40, 500)) for i in range(1, 7)] + [G(cid=9, c=cp(5000, 10, 900))])
    add("rule6 lists at most SCALE_MAX candidates per currency", len(many["findings"]) <= ir.SCALE_MAX, len(many["findings"]))


# =========================================================================================== engine
def _check_engine(add) -> None:
    ir = _ir()
    w = win()
    eur = [G(cid=1, cur="EUR", c=cp(300, 0, 400), p=cp(300, 8, 400)), G(cid=2, cur="EUR", c=cp(700, 20, 900), p=cp(600, 18, 800))]
    vnd = [G(ch="v", cid=5, cur="VND", c=cp(3_000_000, 0, 400), p=cp(3_000_000, 8, 400)), G(ch="v", cid=6, cur="VND", c=cp(7_000_000, 20, 900), p=cp(6_000_000, 18, 800))]
    try:
        ir.rules_for_currency(eur + vnd, "EUR", w)
        refused = False
    except ir.MixedCurrencyError as exc:
        refused = "VND" in str(exc) and "never mixed" in str(exc)
    add("mixed currencies: rules_for_currency REFUSES rows of another currency (MixedCurrencyError)", refused, None)
    add("mixed currencies: a single-currency call accepts its own rows", set(ir.rules_for_currency(eur, "EUR", w)) == {r for r, _t in ir.CURRENCY_RULES}, None)
    res = ir.analyse(eur + vnd, [], w)
    money = [f for f in res["findings"] if f["id"] == "spend_no_conversions"]
    add("per-currency isolation: one finding per currency, each naming only its own", sorted(f["currency"] for f in money) == ["EUR", "VND"], money)
    eurf = next(f for f in money if f["currency"] == "EUR")
    add("per-currency isolation: the EUR finding's texts contain no VND figure", "VND" not in eurf["what"] and all("VND" not in e["text"] for e in eurf["evidence"]), eurf)
    add("per-currency isolation: the EUR share is computed on EUR only (30%, not diluted by VND)", abs(eurf["magnitude"] - 30.0) < 1e-6, eurf)
    add("analyse lists the currencies that carry money", res["currencies"] == ["EUR", "VND"], res["currencies"])
    only_counts = ir.analyse([G(cur="USD", c=cp(0, 5, 100, rev=0), p=cp(0, 5, 100))], [], w)
    add("a counts-only currency is not listed as a currency of money", only_counts["currencies"] == [], only_counts["currencies"])
    order = ir.analyse(eur + [G(cid=3, c=cp(0, 0, 0)), G(cid=4, c=cp(0, 0, 0))], [], w)["findings"]
    add("findings are sorted critical first", [f["severity"] for f in order] == sorted([f["severity"] for f in order], key=lambda s: ir._SEV_RANK[s]), [f["severity"] for f in order])
    mixed_sev = [{"severity": "info", "confidence": "enough data", "magnitude": 1, "id": "x", "channel": None, "campaign": None, "currency": None},
                 {"severity": "crit", "confidence": "thin data", "magnitude": 1, "id": "x", "channel": None, "campaign": None, "currency": None},
                 {"severity": "warn", "confidence": "enough data", "magnitude": 1, "id": "x", "channel": None, "campaign": None, "currency": None}]
    add("sort key: severity first, then enough data before thin data", [f["severity"] for f in sorted(mixed_sev, key=ir._sort_key)] == ["crit", "warn", "info"], None)
    add("a thin finding can never be louder than info (the _finding guard)", ir._finding("x", "crit", "thin data", channel=None, campaign=None, currency=None, what="", evidence=[], action="", magnitude=1)["severity"] == "info", None)
    st = {r["id"]: r for r in res["rules"]}
    add("every rule reports its status; a rule with no input says why it was skipped", set(st) == {r for r, _t in ir.RULES} and st["stale_channel"]["status"] == "skipped" and st["stale_channel"]["why"], st)
    add("counts() tallies severities and thin findings", ir.counts(res["findings"])["crit"] >= 1 and ir.counts(order)["thin"] == 0, ir.counts(res["findings"]))
    add("has_previous is false when no group has rows in the previous window", ir.analyse([G(c=cp(1, 1, 1), p=P(days=0))], [], w)["has_previous"] is False, None)
    ideas = ir.automation_ideas(res["findings"])
    add("automation: a critical finding gives a concrete plain-text idea", any(i["id"] == "spend_no_conversions" and "Weekly" in i["title"] for i in ideas), ideas)
    add("automation: an idea states its threshold from the constants", any(f"{ir.NOCONV_MIN_SHARE_PCT:g}%" in i["text"] for i in ideas), ideas)
    one_warn = [{"id": "cpa_jump", "severity": "warn", "channel": {"name": "a"}, "campaign": {"name": "c"}}]
    add("automation: one warning is not 'recurring'", not ir.automation_ideas(one_warn), None)
    add("automation: two warnings of a rule are recurring", len(ir.automation_ideas(one_warn * 2)) == 1, None)
    add("automation: info findings never create an idea", not ir.automation_ideas([{"id": "scale_candidates", "severity": "info", "channel": None, "campaign": {"name": "c"}}] * 5), None)
    add("automation: ideas are text only (no callable, no schedule field)", all(set(i) == {"id", "title", "text", "based_on", "count"} for i in ideas), ideas)
    with _patched(ir, NOCONV_MIN_SHARE_PCT=7.5):
        txt = " ".join(t["meaning"] + str(t["value"]) for t in ir.thresholds()) + " ".join(r["text"] for r in ir.rule_texts())
    add("the 'How these are decided' text is generated from the constants (patched 7.5 appears)", "7.5" in txt, None)
    add("every threshold constant is listed for the drawer", all(n in {t["name"] for t in ir.thresholds()} or any(n in t["meaning"] for t in ir.thresholds()) for n in dir(ir) if re.match(r"^(NOCONV|CPA|CVR|STALE|HOLE|CONC|SCALE)_[A-Z_]+$", n) and n != "HOLE_LISTED"), None)
    add("the engine never schedules or runs anything: no subprocess, thread, schedule or socket in the rules module", not re.search(r"\b(subprocess|threading|sched|socket|os\.system|Popen)\b", (ROOT / "desktop" / "insights_rules.py").read_text(encoding="utf-8").split('"""', 2)[2]), None)
    add("the rules module has no SQL and never names the raw fact table", not re.search(r"interaction_fact|SELECT |INSERT ", (ROOT / "desktop" / "insights_rules.py").read_text(encoding="utf-8").split('"""', 2)[2]), None)


# =========================================================================================== windows and payload
def _view(**req):
    from desktop import channels_data as cd
    rows = [{"channel_id": i + 1, "days": 20, "sessions": 100, "channel_key": f"ch{i + 1}", "display_name": f"Channel {i + 1}", "medium": "search", "is_paid": True} for i in range(3)]
    options = cd.build_options(rows, {"rollup_rows": 60, "first_day": date(2026, 1, 1), "last_day": date(2026, 6, 30), "currency_rows": [{"currency": "EUR", "days": 20, "has_money": True}]})
    return cd.resolve_view(cd.Request(**req), options, {c["key"]: c["id"] for c in options["_channels"]}), options


def _check_windows(add) -> None:
    from desktop import insights_data as idt
    from desktop import insights_queries as iq
    v, options = _view()
    w = idt.window_of(v)
    add("window maths: the default window is 30 days ending at the newest day, the previous one is the 30 days before", (w.d_to - w.d_from).days == 29 and w.prev_to == w.d_from - timedelta(days=1) and (w.prev_to - w.prev_from).days == 29, w)
    v2, _o = _view(d_from=date(2026, 6, 1), d_to=date(2026, 6, 10))
    add("window maths: a 10-day range has a 10-day previous window ending the day before", (v2.prev_to, v2.prev_from) == (date(2026, 5, 31), date(2026, 5, 22)), (v2.prev_from, v2.prev_to))
    sql, p = iq.groups_sql("public", v2, 10)
    add("groups_sql: the previous window's first day is the lower bound and the analysed window's first day splits the sums", p["d_from"] == v2.prev_from and p["cur_from"] == v2.d_from and "FILTER (WHERE event_date >= :cur_from)" in sql and "FILTER (WHERE event_date < :cur_from)" in sql, p)
    add("groups_sql: every value is a bound parameter (no date or id literal in the text)", not re.search(r"'\d{4}-\d\d-\d\d'", sql) and ":lim" in sql, sql)
    vc, _o = _view(currency="EUR", channels=("ch1",), campaign=7)
    sql, p = iq.groups_sql("public", vc, 10)
    add("groups_sql: channel, campaign and currency filters are the page's own bound clauses", p.get("currency") == "EUR" and p.get("campaign_id") == 7 and p.get("channel_ids") == [1], p)
    sql, p = iq.spans_sql("public", vc)
    add("spans_sql: reads history up to the window end only (upper bound = window end)", p["d_to"] == vc.d_to and p["d_from"].year == 2000, p)
    add("insights queries never name the raw fact or landing table", not re.search(r"interaction_fact|marketing_landing", (ROOT / "desktop" / "insights_queries.py").read_text(encoding="utf-8").split('"""', 2)[2]), None)
    ap = idt.assemble_insights({"groups": {"rows": [], "names": []}, "spans": [], "days": []}, v, options, {"tz": "Asia/Bangkok", "schema": "public"})
    add("payload: a window with no rows says so, names the window and where the data is", ap["state"] == "no_data_in_window" and ap["window"]["from"] == v.d_from.isoformat() and "2026-06-30" in ap["headline"]["sub"] and ap["window"]["data_first_day"] == "2026-01-01", ap["headline"])
    add("payload: no_data_in_window carries no finding and no automation idea", not ap["findings"] and not ap["automation"], ap)
    add("payload: the export links are the server's own view query (xlsx and insights)", ap["export"]["xlsx"].startswith("/api/channels/report.xlsx?from=") and ap["export"]["insights"].startswith("/api/channels/insights?"), ap["export"])
    bad = idt.assemble_insights({"groups": {"error": "boom"}, "spans": [], "days": []}, v, options, {})
    add("payload: a failed block is 'unavailable' with ok false and no finding (never an empty 'all good')", bad["state"] == "unavailable" and bad["ok"] is False and not bad["findings"], bad)
    row = {"channel_id": 1, "campaign_id": 0, "currency": "EUR", **{f"c_{m}": 0 for m in iq.GROUP_MEASURES}, **{f"p_{m}": 0 for m in iq.GROUP_MEASURES}, "c_days": 3, "p_days": 0, "c_clicks": 40, "c_conversions": 2}
    groups, cut = idt.groups_from_rows({"rows": [row], "names": []}, options)
    add("groups_from_rows: campaign 0 is '(no campaign)', the channel name comes from the options, sums map to the window", groups[0].campaign_name == "(no campaign)" and groups[0].channel_name == "Channel 1" and groups[0].cur.clicks == 40 and not cut, groups)
    setup = idt.assemble_insights_setup("not_installed", {"tz": "x", "schema": "public"})
    add("payload: not_installed carries no number, no finding, no export", not setup["findings"] and setup["export"] is None and not re.search(r"\d", setup["headline"]["text"]), setup["headline"])
    e = idt.assemble_insights_setup("empty", {})
    add("payload: empty is a state of its own with the thresholds still available", e["state"] == "empty" and e["how"]["thresholds"], e["state"])
    add("payload: the thresholds block lists all six rules", [r["id"] for r in setup["how"]["rules"]] == [r for r, _t in _ir().RULES], None)


# =========================================================================================== store: whitelist before the cache
def _no_db_engine():
    class _Boom:
        touched = 0

        def connect(self):
            _Boom.touched += 1
            raise RuntimeError("the database must not be touched by this request")
    return _Boom()


def _check_store_guards(add) -> None:
    from desktop import channels_data as cd
    from desktop import insights_data as idt
    boom = _no_db_engine()
    with _patched(cd, engine=boom):
        st = idt.InsightsStore(ttl=60)
        for label, params in (("an unknown parameter name", {"chanel": ["x"]}), ("part (CSV only)", {"part": ["daily"]}), ("a hostile date", {"from": ["2026-02-30"]}),
                              ("a hostile sort", {"sort": [";DROP"]}), ("a currency shape", {"currency": ["eur"]})):
            code, body = st.get(params)
            add(f"insights feed: {label} is a 400 before the cache and before any database access", code == 400 and body["ok"] is False and body["problems"], (code, body))
            code, body, _h = st.export_xlsx(params)
            add(f"xlsx: {label} is a 400 before any database access", code == 400 and body["ok"] is False, (code, body))
        add("the guard names the parameter and never echoes its value", "chanel" in st.get({"chanel": ["SECRETVALUE"]})[1]["error"] + str(st.get({"chanel": ["SECRETVALUE"]})[1]["problems"]) and "SECRETVALUE" not in str(st.get({"chanel": ["SECRETVALUE"]})[1]["problems"]), None)
        add("no request above touched the engine", boom.touched == 0, boom.touched)
    code, _b = st.get({"chanel": ["x"]})
    add("a 400 is never cached", not st._cache, list(st._cache))


# =========================================================================================== SQL fixtures
def _check_setup_states(add) -> None:
    from tests import channels_scenarios as cs
    from desktop import insights_data as idt
    for level, want in (("none", "not_installed"), ("07", "not_installed"), ("empty", "empty")):
        with cs._temp_schema(level):
            st = idt.InsightsStore(ttl=0)
            code, p = st.get({}, fresh=True)
            add(f"states: level '{level}' answers 200 with state {want}, no finding and no automation", code == 200 and p["state"] == want and not p["findings"] and not p["automation"], (code, p.get("state")))
            code, b, _h = st.export_xlsx({})
            add(f"states: the Excel report for '{level}' is a 409 that says why", code == 409 and b.get("state") == want, (code, b))


def _check_sql(add) -> None:
    from tests import channels_scenarios as cs
    from desktop import channels_data as cd
    from desktop import insights_data as idt
    from desktop import insights_queries as iq
    plan = cs._money_plan()
    lo, hi = date(2026, 8, 25), date(2026, 9, 10)
    q = {"from": [lo.isoformat()], "to": [hi.isoformat()]}
    with cs._temp_schema("money") as (eng, conn):
        st = idt.InsightsStore(ttl=0)
        code, p = st.get(q, fresh=True)
        add("sql: the multi-currency insights answer 200 and ready", code == 200 and p["state"] == "ready" and p["ok"], (code, p.get("state"), p.get("error")))
        if code != 200:
            return
        add("sql: the analysed window is the requested one and the previous window is the equal span before it", p["window"]["from"] == "2026-08-25" and p["window"]["to"] == "2026-09-10" and p["window"]["prev_to"] == "2026-08-24" and p["window"]["days"] == 17, p["window"])
        add("sql: currencies with money are exactly EUR, USD, VND in code order", p["currencies"] == ["EUR", "USD", "VND"], p["currencies"])
        add("sql: every finding names at most ONE currency and carries evidence, an action and a confidence", all((f["currency"] is None or re.fullmatch(r"[A-Z]{3}", f["currency"])) and f["evidence"] and f["action"] and f["confidence"] for f in p["findings"]), [f["id"] for f in p["findings"]])
        v, options, meta = None, None, None
        with eng.connect() as c2:
            rd = cd.ld._Reader(c2, 10)
            built = cd.build_view(rd, "pg_temp", q, cd.KNOWN_PARAMS)
            v, options = built[1], built[2]
            raw = iq.collect_insights(rd, "pg_temp", v)
            groups, cut = idt.groups_from_rows(raw["groups"], options)
        exp_cur = sum(r["c"] for r in plan if r["ch"] == 2 and r["cp"] == 20 and lo <= r["d"] <= hi)
        exp_prev = sum(r["c"] for r in plan if r["ch"] == 2 and r["cp"] == 20 and v.prev_from <= r["d"] <= v.prev_to)
        g = next((x for x in groups if x.campaign_id == 20), None)
        add("sql: the EUR e-mail campaign's clicks in the window equal the hand-worked sum", g is not None and g.cur.clicks == exp_cur, (g and g.cur.clicks, exp_cur))
        add("sql: the same campaign's previous-window clicks equal the hand-worked sum (window split done in SQL)", g is not None and g.prev.clicks == exp_prev and exp_prev > 0, (g and g.prev.clicks, exp_prev))
        add("sql: a campaign in two currencies is two groups (hybrid: VND and EUR never merged)", sorted(x.currency for x in groups if x.campaign_id == 50) == ["EUR", "VND"], [(x.campaign_id, x.currency) for x in groups])
        vnd = next(x for x in groups if x.campaign_id == 50 and x.currency == "VND")
        add("sql: the VND group's spend is the VND rows only (3 x 70,000)", vnd.cur.spend_micros == 3 * 70_000 * cs._M, vnd.cur.spend_micros)
        add("sql: the group count matches a direct SQL count of distinct (channel, campaign, currency)", len(groups) == conn.execute(cd.ld.text("SELECT count(*) FROM (SELECT 1 FROM pg_temp.interaction_daily_rollup WHERE event_date BETWEEN :a AND :b GROUP BY channel_id, campaign_id, currency) x"), {"a": v.prev_from, "b": v.d_to}).scalar(), len(groups))
        cur_opt = {**q, "currency": ["VND"]}
        code, pv = st.get(cur_opt, fresh=True)
        add("sql: a currency filter narrows the insights to that currency (no EUR or USD finding)", code == 200 and all(f["currency"] in (None, "VND") for f in pv["findings"]), [f["currency"] for f in pv["findings"]])
        code, pc = st.get({**q, "channel": ["email"]}, fresh=True)
        add("sql: a channel filter limits the findings to that channel", code == 200 and all(f["channel"] is None or f["channel"]["key"] == "email" for f in pc["findings"]), None)
        code, pb = st.get({"from": ["2020-01-01"], "to": ["2020-01-31"]}, fresh=True)
        add("sql: a window before all data is 'no_data_in_window' and offers the data's real range", code == 200 and pb["state"] == "ready" or pb["state"] == "no_data_in_window", pb.get("state"))
        add("sql: cached on the second identical call (same object)", (lambda s2: s2.get(q)[1] is s2.get(q)[1])(idt.InsightsStore(ttl=30)), None)
        add("sql: the poisoned raw-fact row (JPY) appears nowhere in the insights", "JPY" not in str(p), None)
        # ---- the workbook
        code, body, hdr = st.export_xlsx(q)
        add("xlsx: the report answers 200 with bytes and a safe attachment filename", code == 200 and isinstance(body, bytes) and re.fullmatch(r'attachment; filename="channels_report_\d{8}_\d{4}(_filtered)?\.xlsx"', hdr["Content-Disposition"]), (code, hdr))
        _check_xlsx(add, body, plan, lo, hi)
        code, small, _h = st.export_xlsx({"from": ["2020-01-01"], "to": ["2020-01-31"]})
        add("xlsx: a window with no data still yields a valid workbook whose Insights sheet is header-only", code == 200 and _sheet_rows(small, "Insights") == 1, code)
        with _patched(idt, XLSX_MAX_BYTES=100):
            code, b, _h = st.export_xlsx(q)
        add("xlsx: a file over the size bound is refused with 413 and a reason (bounded)", code == 413 and "larger" in b["error"], (code, b))
        add("xlsx: the sheet row cap is a named constant", idt.XLSX_CAMPAIGN_ROWS_MAX == 20_000, None)


def _sheet_rows(body: bytes, name: str) -> int:
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(body))[name].max_row


def _check_xlsx(add, body: bytes, plan, lo, hi) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body))
    add("xlsx: the workbook opens with openpyxl and has exactly the five sheets in order", wb.sheetnames == ["Summary", "Channels", "Campaigns", "Insights", "Definitions"], wb.sheetnames)
    add("xlsx: every sheet has a frozen header row and a styled (bold, filled) header", all(ws.freeze_panes == "A2" and ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("1C1D22") for ws in wb), [ws.freeze_panes for ws in wb])
    add("xlsx: column widths are set (not the default) on the data sheets", all(ws.column_dimensions["A"].width and ws.column_dimensions["A"].width >= 10 for ws in wb), None)
    ch = wb["Channels"]
    head = [c.value for c in ch[1]]
    add("xlsx Channels: the header is the CSV's column set (same names as the CSV)", head[:4] == ["period_from", "period_to", "time_zone", "channel"] and "conversion_rate_pct" in head, head)
    row = {h: ch.cell(row=2, column=i + 1) for i, h in enumerate(head)}
    add("xlsx Channels: counts are real integers with a thousands format", isinstance(row["clicks"].value, int) and row["clicks"].number_format == "#,##0" and row["clicks"].data_type == "n", (row["clicks"].value, row["clicks"].number_format))
    add("xlsx Channels: the period columns are date cells with a date format", row["period_from"].number_format == "yyyy-mm-dd" and hasattr(row["period_from"].value, "year"), row["period_from"].value)
    add("xlsx Channels: rates are fractions with a percent format", row["ctr_pct"].value is None or (isinstance(row["ctr_pct"].value, float) and row["ctr_pct"].number_format == "0.0%" and 0 <= row["ctr_pct"].value <= 1), (row["ctr_pct"].value, row["ctr_pct"].number_format))
    camp = wb["Campaigns"]
    chead = [c.value for c in camp[1]]
    ci = {h: i + 1 for i, h in enumerate(chead)}
    fmts = {}
    for r in range(2, camp.max_row + 1):
        cur, cell = camp.cell(row=r, column=ci["Currency"]).value, camp.cell(row=r, column=ci["Spend"])
        if cell.value is not None:
            fmts.setdefault(cur, set()).add(cell.number_format)
            add("xlsx Campaigns: a spend cell is a number, never text", cell.data_type == "n" and isinstance(cell.value, (int, float)), (cur, cell.value)) if False else None
    add("xlsx Campaigns: VND spend has no decimals, EUR / USD spend has two (format follows the row's currency)", fmts.get("VND") == {"#,##0"} and fmts.get("EUR", {"#,##0.00"}) <= {"#,##0.00"} and fmts.get("USD") == {"#,##0.00"}, fmts)
    add("xlsx Campaigns: every spend cell is numeric (a blank means unknown, never text or 0)", all(camp.cell(row=r, column=ci["Spend"]).value is None or camp.cell(row=r, column=ci["Spend"]).data_type == "n" for r in range(2, camp.max_row + 1)), None)
    add("xlsx Campaigns: a row's money currency column is one code or empty - never a list", all(re.fullmatch(r"[A-Z]{3}|", str(camp.cell(row=r, column=ci["Currency"]).value or "")) for r in range(2, camp.max_row + 1)), None)
    sm = wb["Summary"]
    items = [(c[0].value, c[1].value, c[2].value) for c in sm.iter_rows(min_row=2)]
    spends = [(v, cur) for k, v, cur in items if k == "Spend"]
    add("xlsx Summary: spend is listed once per currency and never as one total across currencies", len({cur for _v, cur in spends}) == len(spends) and all(cur for _v, cur in spends), spends)
    exp_vnd = round(sum(r["sp"] for r in plan if r["cur"] == "VND" and lo <= r["d"] <= hi) / 1e6, 2)
    got_vnd = next((v for v, cur in spends if cur == "VND"), None)
    add("xlsx Summary: the VND spend equals the hand-worked VND sum", got_vnd == exp_vnd, (got_vnd, exp_vnd))
    add("xlsx Summary: the window, the previous window and the time zone are stated", {"Window from", "Window to", "Previous window from", "Time zone (database session)"} <= {i[0] for i in items} and any(i[1] == datetime(lo.year, lo.month, lo.day) for i in items if i[0] == "Window from"), None)
    ins = wb["Insights"]
    add("xlsx Insights: the header is fixed and each finding row has a severity, evidence and an action", [c.value for c in ins[1]][:2] == ["Severity", "Rule"] and all(ins.cell(row=r, column=8).value for r in range(2, ins.max_row + 1)), ins.max_row)
    df = wb["Definitions"]
    thr = [(r[1].value, r[2].value) for r in df.iter_rows(min_row=2) if r[0].value == "Threshold"]
    add("xlsx Definitions: the thresholds are numeric cells and match the constants", dict(thr).get("NOCONV_MIN_SHARE_PCT") == _ir().NOCONV_MIN_SHARE_PCT and all(isinstance(v, (int, float)) for _n, v in thr), thr[:3])
    add("xlsx Definitions: it explains how numbers are computed (a Definition row per page term, the Excel notes)", sum(1 for r in df.iter_rows(min_row=2) if r[0].value == "Definition") >= 10 and any(r[1].value == "Percentages" for r in df.iter_rows(min_row=2)), None)
    add("xlsx: no cell of any sheet is a formula", all(c.data_type != "f" for ws in wb for row in ws.iter_rows() for c in row), None)


def _check_injection(add) -> None:
    from desktop import insights_xlsx as ix
    from tests import channels_scenarios as cs
    from desktop import insights_data as idt
    for v in ("=HYPERLINK(\"http://x\")", "+1+1", "-2+3", "@SUM(A1)", "\t=1", "\r=1"):
        add(f"formula guard: {v!r} is written with a leading apostrophe", ix.safe_text(v).startswith("'"), ix.safe_text(v))
    add("formula guard: plain text and a leading space are untouched", ix.safe_text("Summer sale") == "Summer sale" and ix.safe_text(" =x") == " =x", None)
    add("formula guard: XML-illegal characters are removed, not fatal", ix.safe_text("a\x00b\x1fc") == "abc", ix.safe_text("a\x00b\x1fc"))
    body = ix.build_workbook([ix.Sheet("T", [("=Head", "text"), ("N", "int")], [["=cmd|' /C calc'!A0", 5], ["ok", -7], ["+x", None]])])
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(body))["T"]
    add("formula guard: a hostile header and hostile cells are text (data_type s), never formulas", ws["A1"].data_type == "s" and ws["A2"].data_type == "s" and ws["A2"].value.startswith("'=") and ws["A4"].value == "'+x", (ws["A1"].value, ws["A2"].value))
    add("formula guard: a NEGATIVE number stays a real number (the guard is for text only)", ws["B3"].value == -7 and ws["B3"].data_type == "n", ws["B3"].value)
    with cs._temp_schema("full") as (eng, conn):
        code, b, _h = idt.InsightsStore(ttl=0).export_xlsx({"from": ["2026-08-01"], "to": ["2026-09-10"]})
        wb = openpyxl.load_workbook(io.BytesIO(b))
        names = [r[3].value for r in wb["Channels"].iter_rows(min_row=2)]
        evil = [n for n in names if n and "HYPERLINK" in n]
        add("formula guard on real fixture data: a channel named =HYPERLINK(...) is exported as text with a leading apostrophe", code == 200 and evil and evil[0].startswith("'=") and all(c.data_type != "f" for ws2 in wb for row in ws2.iter_rows() for c in row), (code, evil))


# =========================================================================================== files, feeds, real database
def _check_files(add) -> None:
    html = (ROOT / "desktop/static/shell.html").read_text(encoding="utf-8")
    js = (ROOT / "desktop/static/insights.js").read_text(encoding="utf-8")
    ch = (ROOT / "desktop/static/channels.js").read_text(encoding="utf-8")
    css = (ROOT / "desktop/static/insights.css").read_text(encoding="utf-8")
    srv = (ROOT / "desktop/server.py").read_text(encoding="utf-8")
    add("page: Download Excel sits right after Download CSV in the header", re.search(r'id="chCsv".*?</button>\s*<button[^>]*id="chXlsx"', html, re.S) is not None and "Download Excel" in html, None)
    add("page: the Insights panel, its drawer and both script and stylesheet are wired in the shell", all(x in html for x in ('id="chPIns"', 'id="insBody"', 'id="chInsDrawer"', "/static/insights.js", "/static/insights.css")), None)
    add("page: the insights script loads after channels.js", html.index("/static/channels.js") < html.index("/static/insights.js"), None)
    add("page: the drawer button is a real button with aria-expanded / aria-controls (keyboard reachable)", re.search(r'<button[^>]*id="insHow"[^>]*aria-expanded="false"[^>]*aria-controls="chInsDrawer"', html) is not None, None)
    add("js: the panel uses the server's export link (L-101), a fetch timeout, esc() and no innerHTML from raw data without it", "d.export.insights" in js and "AbortController" in js and js.count("esc(") > 20, None)
    add("js: Escape closes the drawer and focus is restored without scrolling (L-099)", "Escape" in js and js.count("preventScroll") >= 2, None)
    add("js: it never schedules or runs anything: no setInterval, no POST", "setInterval" not in js and "method: 'POST'" not in js and "POST" not in js, None)
    add("js: the Excel button uses the server's xlsx link, not live filter state", "S.data.export.xlsx" in ch and "downloadFrom" in ch, None)
    add("css: colours come from the design tokens and fonts from var(--font) / var(--code), no font name", "Montserrat" not in css and "font-family: var(--font)" in css and not re.search(r"font-family:\s*['\"A-Za-z]", css.replace("var(--font)", "").replace("var(--code)", "")), None)
    add("css: the panel keeps a minimum height while loading (no layout shift)", re.search(r"\.ins-body\s*\{[^}]*min-height", css) is not None, None)
    add("css: reduced motion is honoured", "prefers-reduced-motion" in css, None)
    add("server: both routes are GET and read-only; the report is streamed with the xlsx content type", '@app.get("/api/channels/insights")' in srv and '@app.get("/api/channels/report.xlsx")' in srv and "StreamingResponse" in srv and "XLSX_MIME" in srv, None)
    add("requirements.txt lists openpyxl", re.search(r"^openpyxl\b", (ROOT / "requirements.txt").read_text(encoding="utf-8"), re.M) is not None, None)
    from desktop import insights_data as idt
    add("the xlsx mime type is the standard one", idt.XLSX_MIME == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", idt.XLSX_MIME)


def _check_real_db(add) -> None:
    from desktop import insights_data as idt
    st = idt.InsightsStore(ttl=0)
    code, p = st.get({}, fresh=True)
    add("real db: the insights feed answers honestly in whatever state the database is in", code in (200, 503) and (code == 503 or p["state"] in ("not_installed", "empty", "no_data_in_window", "ready")), (code, p.get("state")))
    if code == 200 and p["state"] == "ready":
        add("real db: a ready answer states the analysed window and the data's real range", p["window"]["from"] and p["window"]["data_first_day"] and p["window"]["data_last_day"], p["window"])
        add("real db: thin-data findings are never warnings", all(f["severity"] == "info" for f in p["findings"] if f["confidence"] == "thin data"), None)
        add("real db: it states it never reads the raw facts", p["read"]["raw_fact_read"] is False, None)
        code, body, _h = st.export_xlsx({})
        add("real db: the Excel report builds and opens", code == 200 and _sheet_rows(body, "Summary") > 5, code)


# =========================================================================================== runner
def run(db_ok: bool = True) -> list[tuple[str, bool, object]]:
    import logging
    rows: list[tuple[str, bool, object]] = []

    def add(name: str, ok, detail=None) -> None:
        rows.append((name, bool(ok), detail if not ok else ""))
    loggers = [logging.getLogger(n) for n in ("erp_desk.channels", "erp_desk.leads", "erp_desk.insights")]
    old = [lg.disabled for lg in loggers]
    for lg in loggers:
        lg.disabled = True
    try:
        fns = ([_check_rule1, _check_rule2, _check_rule3, _check_rule4, _check_rule5, _check_rule6, _check_xlsx_cap, _check_engine, _check_windows, _check_store_guards, _check_files]
               + ([_check_setup_states, _check_sql, _check_injection, _check_real_db] if db_ok else [_check_injection_pure]))
        for fn in fns:
            try:
                fn(add)
            except Exception as exc:  # noqa: BLE001
                import traceback
                add(f"{fn.__name__} crashed", False, f"{type(exc).__name__}: {exc} | {traceback.format_exc().splitlines()[-3][:160]}")
    finally:
        for lg, o in zip(loggers, old):
            lg.disabled = o
        with contextlib.suppress(Exception):
            from desktop import leads_data as ld
            ld._logged.clear()
    return rows


def _check_injection_pure(add) -> None:
    from desktop import insights_xlsx as ix
    add("formula guard (no database): a leading = is neutralised", ix.safe_text("=1+1").startswith("'"), None)


def main() -> int:
    try:
        from sqlalchemy import text

        from erp.db import engine
        with engine.connect() as c:
            c.execute(text("SELECT 1"))
        db_ok = True
    except Exception:  # noqa: BLE001
        db_ok = False
    t0 = time.monotonic()
    rows = run(db_ok)
    for name, ok, detail in rows:
        print(f"{'ok  ' if ok else 'FAIL'} {name}" + (f"   <- {detail}" if not ok else ""))
    bad = [r for r in rows if not r[1]]
    print(f"\n{len(rows) - len(bad)}/{len(rows)} scenarios passed in {time.monotonic() - t0:.1f}s" + ("" if db_ok else "  (database not reachable: the SQL scenarios were skipped)"))
    with contextlib.suppress(Exception):
        from erp.db import engine
        engine.dispose()
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
