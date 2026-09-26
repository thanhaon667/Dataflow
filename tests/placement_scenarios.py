"""
Scenario table for the PLACEMENT feature (display-network placements: WHERE the ads ran and what each place cost and produced), dump-and-count
style like the other scenario files: every scenario is one line "ok / FAIL <name>" and the run ends with "N/M scenarios passed". Kept in the repo so
the Reviewer and the merge gate (tests/smoke.py) can re-run it (L-089).

Everything is SYNTHETIC and generated here (sites under the reserved .example domain, campaigns called "Synthetic ..."); no real export, company or
brand is read or named, and the owner's real files are never touched.

  * pure, no database: the placement_performance connector (strict header, sanitised placement text, occurrence-numbered duplicates, money and currency
    like ad_performance, the never-0 rule for unknown values), the detection order of the inbox processor (no connector steals another's file), the
    five placement rules (fire / quiet / thin data / floor / per-currency isolation / the labelled estimate), the row builders, the request whitelist
    and the 400 before the cache and before any database access, the static files, the Data Flow map;
  * the database, throwaway objects only: a SCHEMA created for the run (perf_plc_<hex>) with the real DDL 07 + 10 + 09 + 11 installed, the synthetic
    files loaded through the REAL autorun -> pipeline -> rollup chain from a temp inbox, and every number compared with a Python-side recomputation.
    `public` is never written: its tables and its interaction_fact row count are compared before and after;
  * the store, the CSV, the Excel Placements sheet and the Insights placement findings read from that schema, and the proof that the page never reads
    interaction_fact (every statement the store sends is captured).

Run (from the project root):
    venv\\Scripts\\python.exe -B -m tests.placement_scenarios
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import logging
import os
import random
import re
import subprocess
import sys
import tempfile
import time
import uuid
from datetime import date, timedelta
from pathlib import Path

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

M = 1_000_000


@contextlib.contextmanager
def _patched(obj, **attrs):
    old = {k: getattr(obj, k) for k in attrs}
    for k, v in attrs.items():
        setattr(obj, k, v)
    try:
        yield
    finally:
        for k, v in old.items():
            setattr(obj, k, v)


# =========================================================================================== connector
def _pc():
    from erp.marketing.connectors import placement_performance as pc
    return pc


def R(**over) -> dict:
    """One good synthetic export row (keys are the export's column names); over = {'Clicks': '5'} replaces cells, None drops the key."""
    row = {"Date": "2026-09-01", "Campaign": "Synthetic Alpha", "Ad group": "Group 1", "Placement": "news-01.example", "Placement type": "website",
           "Ad size": "300x250", "Position": "above_fold", "Device": "mobile", "Impressions": "1000", "Clicks": "10", "Cost": "5.00 EUR",
           "Conversions": "1", "Conversion value": "20.00 EUR", "Viewable impressions": "700"}
    row.update(over)
    return {k: v for k, v in row.items() if v is not None}


def _connector(rows, currency=None, channel=None, date_format="auto"):
    c = _pc().PlacementPerformanceConnector.from_rows(rows, channel=channel, date_format=date_format)
    if currency:
        c.currency = currency
    return c


def _map_all(rows, currency=None, channel=None):
    c = _connector(rows, currency=currency, channel=channel)
    return [c.map_record(r) for r in c.records()]


def _map(row, currency=None, channel=None):
    return _map_all([row], currency, channel)[0]


def _rejection(row, currency=None):
    """The InvalidRecord message a row is refused with, or None when it maps."""
    from erp.marketing.model import InvalidRecord
    try:
        _map(row, currency)
    except InvalidRecord as exc:
        return str(exc)
    return None


def _map_dmy():
    c = _connector([R(Date="03/04/2026")], date_format="dmy")
    return c.map_record(next(iter(c.records())))["event_date"]


def _check_connector(add) -> None:
    pc = _pc()
    from erp.marketing.connectors.base import ConnectorError
    H = list(pc.HEADER)
    add("connector: the export has exactly 14 columns", len(H) == 14 and len(set(H)) == 14, H)
    add("connector: the exact header (any order) is accepted", _connector([R()]) is not None and _connector([dict(reversed(list(R().items())))]) is not None)
    for label, mutate in (("a missing column", lambda r: r.pop("Device") and r), ("an extra column", lambda r: {**r, "Extra": "1"}), ("a renamed column", lambda r: {("Cost " if k == "Cost" else k): v for k, v in r.items()})):
        try:
            _connector([mutate(R())])
            add(f"connector: {label} in the header is refused", False)
        except ConnectorError as exc:
            add(f"connector: {label} in the header is refused, naming it", True if str(exc) else False, str(exc)[:100])
    try:
        pc.PlacementPerformanceConnector.check_header(H + ["Date"])
        add("connector: a repeated column name is refused", False)
    except ConnectorError as exc:
        add("connector: a repeated column name is refused", "repeated" in str(exc), str(exc)[:100])

    m = _map(R())
    add("map: the channel defaults to display", m["channel_key"] == "display", m["channel_key"])
    add("map: --channel overrides the default", _map(R(), channel="Programmatic")["channel_key"] == "programmatic")
    add("map: the source key is the connector's own name", m["source_key"] == "placement_performance")
    add("map: the day, campaign and counts are mapped", (m["event_date"], m["campaign_name"], m["impressions"], m["clicks"], m["conversions"]) == (date(2026, 9, 1), "Synthetic Alpha", 1000, 10, 1), m)
    add("map: money is exact integer micros with the currency written on the cell", (m["spend_micros"], m["revenue_micros"], m["currency"]) == (5 * M, 20 * M, "EUR"), m)
    a = m["attrs"]
    add("map: the placement, its type, size, position, ad group and device are kept as attributes", (a["placement"], a["placement_type"], a["ad_size"], a["position"], a["ad_group"], a["device"]) == ("news-01.example", "website", "300x250", "above_fold", "Group 1", "mobile"), a)
    add("map: the row is flagged as a placement row (the rollup selects on it)", a["placement_row"] is True and a["grain"] == "day")
    add("map: the placement becomes the creative dimension row (name = the text)", m["creative_name"] == "news-01.example" and re.fullmatch(r"news_01_example_[0-9a-f]{8}", m["creative_key"]) is not None, m["creative_key"])
    add("map: reported conversion value is labelled reported", a["revenue_source"].startswith("reported") and a["revenue_derived"] is False)
    add("map: viewable impressions are kept as a number", a["viewable_impressions"] == 700)

    # ---- unknown is never 0
    b = _map(R(**{"Conversion value": "", "Viewable impressions": "", "Ad size": "", "Position": "", "Device": "", "Ad group": ""}))
    add("unknown: a blank conversion value is stored as 0 but labelled NONE (the summary counts it unknown, not zero)", b["revenue_micros"] == 0 and b["attrs"]["revenue_source"].startswith("none"), b["attrs"])
    add("unknown: a blank viewable cell leaves the attribute absent, never 0", b["attrs"]["viewable_impressions"] is None)
    add("unknown: a blank ad size is None", b["attrs"]["ad_size"] is None)
    add("unknown: a blank position is 'unknown'", b["attrs"]["position"] == "unknown")
    add("unknown: a blank device and ad group are None, not guessed", b["attrs"]["device"] is None and b["attrs"]["ad_group"] is None)
    add("unknown: an explicit 0 viewable impressions IS zero, different from blank", _map(R(**{"Viewable impressions": "0"}))["attrs"]["viewable_impressions"] == 0)
    add("unknown: an 'unknown' ad size reads as blank", _map(R(**{"Ad size": "Unknown"}))["attrs"]["ad_size"] is None)
    from erp.marketing import model
    full, _idx = model.split_attrs(b["attrs"])
    add("unknown: after the pipeline's attribute clean-up the blank viewable key is really gone", "viewable_impressions" not in full and "ad_size" not in full, full)

    # ---- sanitised text
    add("text: a placement is stripped and its inner whitespace collapsed", pc.sanitise_placement("  news   site\t.example \n") == ("news site .example", False))
    add("text: an embedded line break is a space, not a control character", pc.sanitise_placement("a\nb")[0] == "a b")
    for label, text in (("NUL", "a\x00b.example"), ("ESC", "a\x1bb.example"), ("a zero-width space", "a​b.example"), ("a right-to-left override", "a‮b.example"), ("a BOM inside", "a﻿b.example"),
                        ("a DEL", "a\x7fb.example"), ("a soft hyphen", "a­b.example")):
        add(f"text: a placement holding {label} is REFUSED, not cleaned", "control or invisible" in (_rejection(R(Placement=text)) or ""), _rejection(R(Placement=text)))
    for label, text in (("empty", ""), ("spaces", "   "), ("null", "null"), ("a dash", "-")):
        add(f"text: an {label} placement is refused" if label == "empty" else f"text: a placement of '{label}' is refused", "Placement" in (_rejection(R(Placement=text)) or ""), _rejection(R(Placement=text)))
    long = "x" * 260 + ".example"
    clean, cut = pc.sanitise_placement(long)
    add("text: a placement over 200 characters is cut and flagged", len(clean) <= pc.MAX_PLACEMENT and cut is True and _map(R(Placement=long))["attrs"]["placement_truncated"] is True)
    add("text: a normal placement is not flagged as cut", _map(R())["attrs"]["placement_truncated"] is None)
    add("text: a formula-looking placement is kept literally (the sheet and the CSV neutralise it later)", _map(R(Placement="=HYPERLINK(\"x\")"))["creative_name"].startswith("=HYPERLINK"))
    add("text: unicode letters are kept", pc.sanitise_placement("tin-tuc.example/bài-viết")[0] == "tin-tuc.example/bài-viết")

    # ---- creative keys
    k1, k2, k3 = pc.creative_key_of("a.b.com"), pc.creative_key_of("a-b.com"), pc.creative_key_of("A.B.COM")
    add("creative key: two placements that differ only in punctuation never share a key", k1 != k2, (k1, k2))
    add("creative key: case does not split a placement", k1 == k3)
    add("creative key: at most 200 wide (the column width) and ends in 8 hex digits", len(pc.creative_key_of("y" * 400)) <= 200 and re.search(r"_[0-9a-f]{8}$", pc.creative_key_of("x")) is not None)
    add("creative key: a placement with no letters still gets a key", re.fullmatch(r"placement_[0-9a-f]{8}", pc.creative_key_of("...")) is not None, pc.creative_key_of("..."))

    # ---- type / size / position
    for label, cell, want in (("Website", "Website", "website"), ("APP", "APP", "app"), ("video", " video ", "video"), ("other", "Other", "other")):
        add(f"type: '{label}' reads as {want}", _map(R(**{"Placement type": cell}))["attrs"]["placement_type"] == want)
    for cell in ("banner", "", "web site"):
        add(f"type: {cell!r} is refused", "Placement type" in (_rejection(R(**{"Placement type": cell})) or ""), _rejection(R(**{"Placement type": cell})))
    for cell, want in (("300x250", "300x250"), ("300X250", "300x250"), ("320 x 50", "320x50"), ("0300x0250", "300x250")):
        add(f"size: {cell!r} reads as {want}", _map(R(**{"Ad size": cell}))["attrs"]["ad_size"] == want, _rejection(R(**{"Ad size": cell})))
    for cell in ("300*250", "0x50", "300x", "big", "３００x２５０", "300x250x2"):
        add(f"size: {cell!r} is refused (fullwidth digits included: L-220)", "Ad size" in (_rejection(R(**{"Ad size": cell})) or ""), _rejection(R(**{"Ad size": cell})))
    for cell, want in (("above_fold", "above_fold"), ("Above fold", "above_fold"), ("below-fold", "below_fold"), ("", "unknown"), ("unknown", "unknown")):
        add(f"position: {cell!r} reads as {want}", _map(R(Position=cell))["attrs"]["position"] == want, _rejection(R(Position=cell)))
    add("position: 'top' is refused", "Position" in (_rejection(R(Position="top")) or ""))

    # ---- numbers
    for label, over, needle in (("clicks above impressions", {"Clicks": "2000"}, "more than Impressions"), ("viewable above impressions", {"Viewable impressions": "1001"}, "more than Impressions"),
                                ("a fractional conversion count", {"Conversions": "2.5"}, "Conversions"), ("a negative click count", {"Clicks": "-1"}, "Clicks"),
                                ("fullwidth digits in impressions", {"Impressions": "１０００"}, "Impressions"), ("text in impressions", {"Impressions": "many"}, "Impressions"),
                                ("an empty impressions cell", {"Impressions": ""}, "Impressions"), ("an impossible date", {"Date": "2026-02-30"}, "Date"), ("an empty campaign", {"Campaign": " "}, "Campaign")):
        add(f"reject: {label}", needle in (_rejection(R(**over)) or ""), _rejection(R(**over)))
    add("accept: a whole conversion count written 3.0", _map(R(Conversions="3.0"))["conversions"] == 3)
    add("accept: thousands separators in a count", _map(R(Impressions="1,500"))["impressions"] == 1500)

    # ---- money and currency, like ad_performance
    for cell, want in (("12.50 EUR", ("EUR", 12_500_000)), ("USD 12.50", ("USD", 12_500_000)), ("€12.50", ("EUR", 12_500_000)), ("475.401 ₫", ("VND", 475_401 * M))):
        try:
            got = _map(R(Cost=cell, **{"Conversion value": ""}))
            ok = (got["currency"], got["spend_micros"]) == want
        except Exception as exc:  # noqa: BLE001
            ok = False
            got = str(exc)
        add(f"money: cost {cell!r} names its currency: {want[0]}", ok, got if not ok else "")
    add("money: a bare dollar sign is ambiguous (USD, AUD, CAD ...) and refused with the reason, never guessed", "ambiguous" in (_rejection(R(Cost="$12.50", **{"Conversion value": ""})) or ""), _rejection(R(Cost="$12.50")))
    add("money: a cell naming no currency is refused when no --currency is given", "names no currency" in (_rejection(R(Cost="12.50", **{"Conversion value": ""})) or ""), _rejection(R(Cost="12.50", **{"Conversion value": ""})))
    add("money: the --currency supplies the currency of a bare cell", _map(R(Cost="12.50", **{"Conversion value": ""}), currency="EUR")["currency"] == "EUR")
    add("money: a cell in another currency than --currency is refused (one file, one currency)", "one file, one currency" in (_rejection(R(Cost="12.50 USD", **{"Conversion value": ""}), currency="EUR") or ""), _rejection(R(Cost="12.50 USD"), currency="EUR"))
    add("money: a conversion value in another currency than the cost is refused", "one row, one currency" in (_rejection(R(**{"Conversion value": "20.00 USD"})) or ""), _rejection(R(**{"Conversion value": "20.00 USD"})))
    add("money: a bare conversion value takes the cost's currency", _map(R(**{"Conversion value": "20.00"}))["currency"] == "EUR")
    add("money: 1.249 in EUR is ambiguous and refused, never read as 1249 or 1.249", "Cost" in (_rejection(R(Cost="1.249 EUR", **{"Conversion value": ""})) or ""), _rejection(R(Cost="1.249 EUR")))
    add("money: an unreadable cost is refused", "Cost" in (_rejection(R(Cost="cheap")) or ""))
    try:
        pc.PlacementPerformanceConnector([R()], currency="EU")
        add("money: the constructor refuses a currency that is not ISO 4217", False)
    except ValueError:
        add("money: the constructor refuses a currency that is not ISO 4217", True)

    # ---- duplicates: occurrence numbering, never a silent overwrite
    recs = _map_all([R(Clicks="10"), R(Clicks="12"), R(Clicks="14")])
    ids = [r["external_id"] for r in recs]
    add("duplicates: three rows under one natural key are numbered 0, 1, 2", [r["occurrence"] for r in recs] == [0, 1, 2], [r["occurrence"] for r in recs])
    add("duplicates: ... and get three different identities (all three are kept)", len(set(ids)) == 3, ids)
    other = _map_all([R(), R(Device="desktop"), R(**{"Ad group": "Group 2"}), R(Position="below_fold"), R(**{"Ad size": "728x90"}), R(Placement="news-02.example"), R(Date="2026-09-02"), R(Campaign="Synthetic Beta")])
    add("duplicates: a row that differs in ANY key part is a different identity, occurrence 0", len({r["external_id"] for r in other}) == 8 and all(r["occurrence"] == 0 for r in other))
    add("duplicates: the same key in another letter case and spacing is the same natural key", _map_all([R(Placement="News-01.example"), R(Placement=" news-01.example ")])[1]["occurrence"] == 1)
    add("duplicates: the natural key is stored readably on the record", "news-01.example" in recs[0]["natural_key"] and recs[0]["natural_key"].startswith("display|2026-09-01|"), recs[0]["natural_key"])
    add("duplicates: an identity longer than the column is replaced by a digest, never truncated", _map(R(Placement="p" * 190, Campaign="c" * 250))["external_id"].startswith("h:"))

    # ---- dates
    add("dates: ISO dates need no format", _map(R(Date="2026-09-03"))["event_date"] == date(2026, 9, 3))
    try:
        _connector([R(Date="03/04/2026"), R(Date="05/06/2026")])
        add("dates: a slash file whose order cannot be proven is refused", False)
    except ConnectorError:
        add("dates: a slash file whose order cannot be proven is refused", True)
    add("dates: --date-format dmy settles an ambiguous slash file and the day is read day-first", _map_dmy() == date(2026, 4, 3))

    # ---- files
    with tempfile.TemporaryDirectory(prefix="plc_csv_") as t:
        p = Path(t) / "x.csv"
        p.write_text(",".join(pc.HEADER) + "\n2026-09-01,Synthetic Alpha,G,a.example,website,300x250,above_fold,mobile,100,1,1.00 EUR,0,,\n", encoding="utf-8")
        c = pc.PlacementPerformanceConnector.from_csv(p, currency="EUR")
        got = [c.map_record(r) for r in c.records()]
        add("file: a CSV with blank trailing cells maps", len(got) == 1 and got[0]["attrs"]["viewable_impressions"] is None and got[0]["revenue_micros"] == 0)
        c._rows.close()
        p.write_text("a,b\n1,2\n", encoding="utf-8")
        try:
            pc.PlacementPerformanceConnector.from_csv(p)
            add("file: a CSV with another header is refused before any row is read", False)
        except ConnectorError as exc:
            add("file: a CSV with another header is refused before any row is read", "placement_performance" in str(exc))
        try:
            pc.PlacementPerformanceConnector.from_csv(p, currency="EURO")
            add("file: an invalid --currency is refused by from_csv", False)
        except (ValueError, ConnectorError):
            add("file: an invalid --currency is refused by from_csv", True)
        p.write_text(",".join(pc.HEADER) + "\n2026-09-01,Synthetic Alpha,G,a.example,website,300x250,above_fold,mobile,100,1,1.00 EUR,0,,,EXTRA\n", encoding="utf-8")
        c = pc.PlacementPerformanceConnector.from_csv(p)
        try:
            [c.map_record(r) for r in c.records()]
            add("file: a row with more cells than the header is rejected by name", False)
        except model.InvalidRecord as exc:
            add("file: a row with more cells than the header is rejected by name", "cells" in str(exc), str(exc))
        c._rows.close()


# =========================================================================================== detection, CLI
def _check_detection(add) -> None:
    from erp.marketing import autorun, ingest
    from erp.marketing.connectors.ad_performance import AdPerformanceConnector
    from erp.marketing.connectors.base import ConnectorError
    from erp.marketing.connectors.email_campaign import EmailCampaignConnector
    from tests import marketing_autorun_scenarios as mas
    pc = _pc()
    H = list(pc.HEADER)
    for label, cls, head in (("ad_performance", AdPerformanceConnector, mas.AD_HEADER.split(",")), ("email_campaign", EmailCampaignConnector, mas.EMAIL_HEADER.split(","))):
        try:
            pc.PlacementPerformanceConnector.check_header(head)
            add(f"no theft: the {label} header is refused by the placement connector", False)
        except ConnectorError:
            add(f"no theft: the {label} header is refused by the placement connector", True)
        try:
            cls.check_header(H)
            add(f"no theft: the placement header is refused by {label}", False)
        except ConnectorError:
            add(f"no theft: the placement header is refused by {label}", True)
    body = ",".join(H) + "\n2026-09-01,Synthetic Alpha,G,a.example,website,300x250,above_fold,mobile,100,1,1.00 EUR,0,,\n"
    with tempfile.TemporaryDirectory(prefix="plc_det_") as t:
        def det(name, text):
            p = Path(t) / name
            p.write_bytes(text.encode("utf-8"))
            return autorun.detect(p)
        d = det("export.csv", body)
        add("detect: the exact placement header -> placement_performance", d.connector == "placement_performance" and d.channel is None and d.currency is None, d)
        d = det("display__eur__q3.csv", body)
        add("detect: a placement file may carry its channel and currency in the name", (d.connector, d.channel, d.currency) == ("placement_performance", "display", "EUR"), d)
        d = det("data1.csv", mas.AD_CSV)
        add("detect: the ad_performance file still goes to ad_performance", d.connector == "ad_performance", d)
        d = det("data2.csv", mas.EMAIL_CSV)
        add("detect: the email_campaign file still goes to email_campaign", d.connector == "email_campaign", d)
        d = det("facebook__usd__q3.csv", mas.FLAT_CSV)
        add("detect: a flat file with a filename hint still goes to flat_file", d.connector == "flat_file" and d.channel == "facebook", d)
        d = det("mystery.csv", mas.UNKNOWN_CSV)
        add("detect: an unknown header is still refused with reasons that name the placement export", d.connector is None and "placement_performance" in d.reason, d.reason[:200])
        d = det("display__eur__extra.csv", body.replace("Viewable impressions", "Viewable impressions,Extra").replace(",,\n", ",,,\n"))
        add("detect: a placement header with an extra column is not placement_performance (flat_file by the hint, never a guess)", d.connector != "placement_performance", d)
        d = det("wide.csv", body.replace("Device,", "Device ,"))
        add("detect: a header with a padded name is refused", d.connector is None, d)

    p = ingest.build_parser()
    a = p.parse_args(["--connector", "placement_performance", "--csv", "x.csv", "--currency", "EUR", "--channel", "display", "--dry-run"])
    add("cli: --connector placement_performance is understood", a.connector == "placement_performance" and a.currency == "EUR")
    add("cli: the connector list has the four connectors, flat_file stays the default", set(ingest.CONNECTORS) == {"flat_file", "ad_performance", "email_campaign", "placement_performance"} and p.parse_args(["--csv", "x"]).connector == "flat_file")
    with tempfile.TemporaryDirectory(prefix="plc_cli_") as t:
        f = Path(t) / "p.csv"
        f.write_text(body, encoding="utf-8")
        args = p.parse_args(["--connector", "placement_performance", "--csv", str(f), "--currency", "EUR", "--source", "My Placement Feed"])
        c = ingest.build_connector(args, f)
        add("cli: --currency reaches the placement connector and --source renames it (slug)", c.currency == "EUR" and c.source_key == "my_placement_feed", (c.currency, c.source_key))
        c._rows.close()
        bad = p.parse_args(["--connector", "ad_performance", "--csv", str(f), "--currency", "EUR"])
        try:
            ingest.build_connector(bad, f)
            add("cli: --currency is still refused for ad_performance", False)
        except ValueError as exc:
            add("cli: --currency is still refused for ad_performance", "placement_performance" in str(exc), str(exc)[:120])
        worse = p.parse_args(["--connector", "placement_performance", "--csv", str(f), "--currency", "EURO"])
        try:
            ingest.build_connector(worse, f)
            add("cli: an invalid --currency is refused for the placement connector", False)
        except (ValueError, ConnectorError):
            add("cli: an invalid --currency is refused for the placement connector", True)

        # a dry run opens no database connection at all (L-235)
        from erp.marketing.pipeline import Pipeline

        class Boom:
            touched = 0

            def connect(self, *a, **k):
                Boom.touched += 1
                raise RuntimeError("no database in a dry run")

            begin = connect
        conn = ingest.build_connector(p.parse_args(["--connector", "placement_performance", "--csv", str(f), "--currency", "EUR"]), f)
        res = Pipeline(engine=Boom(), schema="perf_none").run(conn, dry_run=True)
        conn._rows.close()
        add("dry run: the placement file is read and validated with an engine whose connect() raises", res.rows_read == 1 and res.rows_loaded == 0 and Boom.touched == 0, (res.rows_read, Boom.touched))


# =========================================================================================== rules (pure)
def _ir():
    from desktop import insights_rules as ir
    return ir


def PS(name, key=None, impressions=0, clicks=0, conversions=0, spend=0.0, viewable=0, base=0, ptype="website"):
    ir = _ir()
    return ir.PlacementStat(name=name, key=key if key is not None else name.lower().replace(" ", "_"), ptype=ptype, impressions=impressions, clicks=clicks,
                            conversions=conversions, spend_micros=int(round(spend * M)), viewable=viewable, viewable_base=base)


def PG(*stats, cur="EUR", camp=1, ch="display"):
    return _ir().PlacementGroup(channel_key=ch, channel_name=ch.title(), campaign_id=camp, campaign_name=f"Camp {camp}", currency=cur, placements=tuple(stats))


def WIN():
    return _ir().Window(date(2026, 9, 1), date(2026, 9, 30), date(2026, 8, 2), date(2026, 8, 31))


def _rule(fn, *stats, **kw):
    return fn(PG(*stats, **kw), WIN())


def _base_placements(n=6, conv=2, clicks=200, spend=100.0, impressions=40000):
    return [PS(f"base {i}", clicks=clicks, conversions=conv, spend=spend, impressions=impressions) for i in range(n)]


def _check_rules(add) -> None:
    ir = _ir()
    fn = ir
    # ---------------- exclude: no conversions
    base = _base_placements()
    waste = PS("waste", clicks=300, spend=200.0, impressions=60000)
    r = _rule(ir.rule_placement_exclude, waste, *base)
    f = [x for x in r["findings"] if x["placement"]["key"] == "waste"]
    add("exclude: a big spender with no conversions in a converting campaign fires", len(f) == 1 and r["ran"], r)
    if f:
        f = f[0]
        share = 200 / 800 * 100
        add("exclude: it is enough data (300 clicks) and a warning at 25% share becomes critical", f["confidence"] == "enough data" and f["severity"] == "crit", (f["confidence"], f["severity"], share))
        add("exclude: it carries evidence with the numbers, an action and never claims to have changed anything", len(f["evidence"]) >= 4 and "Nothing has been excluded" in f["action"] and f["estimate"] is False)
        add("exclude: the finding names ONE currency and its placement", f["currency"] == "EUR" and f["placement"]["name"] == "waste")
    small = PS("waste", clicks=300, spend=40.0, impressions=60000)
    r = _rule(ir.rule_placement_exclude, small, *base)
    f = [x for x in r["findings"] if x["placement"]["key"] == "waste"]
    add("exclude: a 6% spend share is a warning, not critical", len(f) == 1 and f[0]["severity"] == "warn", [x["severity"] for x in f])
    tiny = PS("waste", clicks=300, spend=10.0, impressions=60000)
    add("exclude: a share under 5% is quiet", not [x for x in _rule(ir.rule_placement_exclude, tiny, *base)["findings"] if x["placement"]["key"] == "waste"])
    thin = PS("waste", clicks=50, spend=200.0, impressions=20000)
    rich = _base_placements(conv=20)
    f = [x for x in _rule(ir.rule_placement_exclude, thin, *rich)["findings"] if x["placement"]["key"] == "waste"]
    add("exclude: 50 clicks is thin data: reported, but as info only - never a warning or critical", len(f) == 1 and f[0]["confidence"] == "thin data" and f[0]["severity"] == "info", [(x["confidence"], x["severity"]) for x in f])
    floor = PS("waste", clicks=10, spend=200.0, impressions=5000)
    add("exclude: under 20 clicks nothing is said", not [x for x in _rule(ir.rule_placement_exclude, floor, *base)["findings"] if x["placement"]["key"] == "waste"])
    lucky = PS("waste", clicks=25, spend=200.0, impressions=5000)
    rest_low = [PS(f"base {i}", clicks=200, conversions=1, spend=100.0, impressions=40000) for i in range(6)]
    add("exclude: a small sample that had no luck is not accused (the chance of zero is above 5%)", not [x for x in _rule(ir.rule_placement_exclude, lucky, *rest_low)["findings"] if x["placement"]["key"] == "waste"])
    few = _base_placements(conv=0)
    few[0] = PS("base 0", clicks=200, conversions=2, spend=100.0, impressions=40000)
    r = _rule(ir.rule_placement_exclude, waste, *few)
    add("exclude: a campaign with fewer than 5 conversions gets a note (missing tracking?) instead of a verdict", not [x for x in r["findings"] if x["placement"]["key"] == "waste"] and any("no verdict" in n for n in r["notes"]), r["notes"])
    other = PS("(other placements)", key="", clicks=300, spend=200.0, impressions=60000)
    add("exclude: the '(other placements)' bucket is never judged", not [x for x in _rule(ir.rule_placement_exclude, other, *base)["findings"] if x["placement"] and x["placement"]["key"] == ""])
    add("exclude: no spend at all is skipped, with the reason", _rule(ir.rule_placement_exclude, PS("a", clicks=5), PS("b", clicks=6))["ran"] is False)
    lowctr = PS("lowctr", impressions=200000, clicks=100, conversions=2, spend=60.0)
    ctr_base = [PS(f"c{i}", impressions=100000, clicks=1000, conversions=10, spend=100.0) for i in range(4)]
    f = [x for x in _rule(ir.rule_placement_exclude, lowctr, *ctr_base)["findings"] if x["placement"]["key"] == "lowctr"]
    add("exclude: a CTR under a quarter of the rest of the campaign fires as a warning", len(f) == 1 and f[0]["severity"] == "warn" and f[0]["confidence"] == "enough data", f)
    lowctr_thin = PS("lowctr", impressions=8000, clicks=4, conversions=0, spend=10.0)
    f = [x for x in _rule(ir.rule_placement_exclude, lowctr_thin, *ctr_base)["findings"] if x["placement"]["key"] == "lowctr"]
    add("exclude: the same CTR gap on 8,000 impressions is thin data", len(f) == 1 and f[0]["confidence"] == "thin data", f)
    fine = PS("fine", impressions=200000, clicks=1500, conversions=8, spend=60.0)
    add("exclude: a normal CTR is quiet", not [x for x in _rule(ir.rule_placement_exclude, fine, *ctr_base)["findings"] if x["placement"]["key"] == "fine"])
    fewi = PS("fewi", impressions=3000, clicks=0, conversions=0, spend=60.0)
    add("exclude: under 5,000 impressions the CTR test says nothing", not [x for x in _rule(ir.rule_placement_exclude, fewi, *ctr_base)["findings"] if x["placement"]["key"] == "fewi" and "lowctr" in str(x["id"])])

    # ---------------- scale
    star = PS("star", clicks=100, conversions=20, spend=100.0, impressions=30000)
    mid = [PS(f"m{i}", clicks=200, conversions=10, spend=300.0, impressions=40000) for i in range(3)]
    r = _rule(ir.rule_placement_scale, star, *mid)
    f = [x for x in r["findings"] if x["placement"]["key"] == "star"]
    add("scale: a placement 30%+ cheaper per conversion with enough conversions fires, always info", len(f) == 1 and f[0]["severity"] == "info" and f[0]["confidence"] == "enough data", r["findings"])
    add("scale: its action asks for small steps and does not promise a result", "small steps" in f[0]["action"] and "watch" in f[0]["action"])
    star3 = PS("star", clicks=100, conversions=3, spend=10.0, impressions=30000)
    f = [x for x in _rule(ir.rule_placement_scale, star3, *mid)["findings"] if x["placement"]["key"] == "star"]
    add("scale: 3 to 4 conversions is thin data", len(f) == 1 and f[0]["confidence"] == "thin data", f)
    star2 = PS("star", clicks=100, conversions=2, spend=4.0, impressions=30000)
    add("scale: 2 conversions says nothing", not [x for x in _rule(ir.rule_placement_scale, star2, *mid)["findings"] if x["placement"]["key"] == "star"])
    same = PS("same", clicks=100, conversions=10, spend=290.0, impressions=30000)
    add("scale: a cost per conversion close to the rest is quiet", not [x for x in _rule(ir.rule_placement_scale, same, *mid)["findings"] if x["placement"]["key"] == "same"])
    add("scale: a campaign with one converting placement is skipped (nothing to compare)", _rule(ir.rule_placement_scale, star)["ran"] is False)
    many = [PS(f"s{i}", clicks=100, conversions=20, spend=50.0 + i, impressions=30000) for i in range(6)]
    got = [x for x in _rule(ir.rule_placement_scale, *many, PS("z", clicks=100, conversions=10, spend=900.0, impressions=1000))["findings"]]
    add("scale: at most 3 candidates per campaign and currency", len(got) == 3, len(got))
    add("scale: the best (cheapest) come first", [x["placement"]["key"] for x in sorted(got, key=lambda x: -x["magnitude"])][0] == "s0", [x["placement"]["key"] for x in got])

    # ---------------- viewability
    seen = PS("seen", impressions=50000, clicks=200, conversions=5, spend=100.0, viewable=30000, base=50000)
    unseen = PS("unseen", impressions=50000, clicks=200, conversions=5, spend=100.0, viewable=10000, base=50000)
    r = _rule(ir.rule_placement_viewability, unseen, seen)
    f = [x for x in r["findings"] if x["placement"]["key"] == "unseen"]
    add("viewability: 20% viewable on 50,000 measured impressions fires as a warning, enough data", len(f) == 1 and f[0]["severity"] == "warn" and f[0]["confidence"] == "enough data", r["findings"])
    add("viewability: a placement at 60% is quiet", not [x for x in r["findings"] if x["placement"]["key"] == "seen"])
    thinv = PS("unseen", impressions=8000, clicks=20, conversions=1, spend=100.0, viewable=800, base=8000)
    f = [x for x in _rule(ir.rule_placement_viewability, thinv, seen)["findings"] if x["placement"]["key"] == "unseen"]
    add("viewability: 8,000 measured impressions is thin data", len(f) == 1 and f[0]["confidence"] == "thin data", f)
    floorv = PS("unseen", impressions=3000, clicks=20, conversions=1, spend=100.0, viewable=100, base=3000)
    add("viewability: under 5,000 measured impressions says nothing", not [x for x in _rule(ir.rule_placement_viewability, floorv, seen)["findings"] if x["placement"]["key"] == "unseen"])
    partial = PS("unseen", impressions=200000, clicks=200, conversions=5, spend=100.0, viewable=2000, base=10000)
    r = _rule(ir.rule_placement_viewability, partial, seen)
    add("viewability: only 5% of its impressions measured: not judged, a note says why", not [x for x in r["findings"] if x["placement"]["key"] == "unseen"] and any("measured" in n for n in r["notes"]), r["notes"])
    add("viewability: no placement reporting viewability skips the rule (a blank is unknown, not 0%)", _rule(ir.rule_placement_viewability, PS("a", impressions=9000, clicks=9, spend=5.0), PS("b", impressions=9000, clicks=9, spend=5.0))["ran"] is False)
    tinyshare = PS("unseen", impressions=50000, clicks=200, conversions=5, spend=1.0, viewable=1000, base=50000)
    add("viewability: a placement with under 2% of the spend is not judged", not [x for x in _rule(ir.rule_placement_viewability, tinyshare, seen, PS("big", impressions=9000, clicks=90, conversions=2, spend=500.0))["findings"] if x["placement"]["key"] == "unseen"])

    # ---------------- concentration
    five = [PS(f"top{i}", clicks=300, conversions=6, spend=s, impressions=50000) for i, s in enumerate((300.0, 250.0, 200.0, 10.0, 10.0, 10.0))]
    r = _rule(ir.rule_placement_concentration, *five)
    add("concentration: the top 3 carrying about 95% of the spend fires", len(r["findings"]) == 1 and r["findings"][0]["placement"] is None, r["findings"])
    if r["findings"]:
        add("concentration: it is a warning at 80%+ and critical at 95%+", r["findings"][0]["severity"] == "crit", (r["findings"][0]["severity"], r["findings"][0]["magnitude"]))
    spread = [PS(f"p{i}", clicks=300, conversions=6, spend=100.0, impressions=50000) for i in range(6)]
    add("concentration: an even spread over 6 placements is quiet", not _rule(ir.rule_placement_concentration, *spread)["findings"])
    add("concentration: fewer than 6 placements is skipped (few places are always concentrated)", _rule(ir.rule_placement_concentration, *five[:5])["ran"] is False)
    warn = [PS(f"w{i}", clicks=300, conversions=6, spend=s, impressions=50000) for i, s in enumerate((300.0, 250.0, 200.0, 50.0, 50.0, 50.0))]
    fw = _rule(ir.rule_placement_concentration, *warn)["findings"]
    add("concentration: 83% is a warning", len(fw) == 1 and fw[0]["severity"] == "warn", fw)
    lowclicks = [PS(f"l{i}", clicks=20, conversions=1, spend=s, impressions=5000) for i, s in enumerate((300.0, 250.0, 200.0, 10.0, 10.0, 10.0))]
    fl = _rule(ir.rule_placement_concentration, *lowclicks)["findings"]
    add("concentration: under 200 clicks in the campaign is thin data", len(fl) == 1 and fl[0]["confidence"] == "thin data", fl)

    # ---------------- reallocation estimate
    w2 = PS("waste", clicks=300, spend=200.0, impressions=60000)
    b2 = [PS(f"good{i}", clicks=100, conversions=20, spend=100.0, impressions=30000) for i in range(2)]
    rest = [PS(f"rest{i}", clicks=200, conversions=10, spend=300.0, impressions=40000) for i in range(3)]
    r = _rule(ir.rule_placement_reallocation, w2, *b2, *rest)
    add("estimate: with a clearly wasteful and a clearly efficient group it fires", len(r["findings"]) == 1, r)
    if r["findings"]:
        f = r["findings"][0]
        add("estimate: it is flagged estimate, info, enough data and says ESTIMATE in words", f["estimate"] is True and f["severity"] == "info" and f["what"].startswith("Estimate"), f["what"])
        assum = [e for e in f["evidence"] if e["label"] == "Assumption"] if f["evidence"] and "label" in f["evidence"][0] else [e for e in f["evidence"] if "Assumption" in str(e)]
        add("estimate: the assumption is written on the finding (constant cost per conversion, no saturation)", bool(assum) and "constant" in str(assum[0]).lower() or "CURRENT cost per conversion" in str(f["evidence"]), f["evidence"])
        add("estimate: the gain is the moved spend divided by the receiving group's own cost per conversion", abs(f["magnitude"] - 200.0 / (200.0 / 40)) < 1e-6, f["magnitude"])
        add("estimate: it changes nothing by itself and is not a promise", "changes nothing" in f["action"] and "not as a promise" in f["action"])
    only_waste = _rule(ir.rule_placement_reallocation, w2, *rest)
    add("estimate: without an efficient group there is no estimate", not only_waste["findings"] and only_waste["ran"] is False, only_waste)
    only_good = _rule(ir.rule_placement_reallocation, *b2, *rest)
    add("estimate: without a wasteful placement there is no estimate", not only_good["findings"] and only_good["ran"] is False)
    weak = [PS("g0", clicks=100, conversions=3, spend=10.0, impressions=30000), PS("g1", clicks=100, conversions=3, spend=10.0, impressions=30000)]
    add("estimate: an efficient group with fewer than 10 conversions gives no estimate", not _rule(ir.rule_placement_reallocation, w2, *weak, *rest)["findings"])
    thin_waste = PS("waste", clicks=50, spend=200.0, impressions=20000)
    add("estimate: a sending group with fewer than 100 clicks (thin data) gives no estimate", not _rule(ir.rule_placement_reallocation, thin_waste, *b2, *rest)["findings"])
    huge = PS("waste", clicks=300, spend=5000.0, impressions=60000)
    rr = _rule(ir.rule_placement_reallocation, huge, *b2, *rest)
    add("estimate: the amount moved is capped at the receiving group's own spend (it can at most double)", bool(rr["findings"]) and abs(rr["findings"][0]["magnitude"] - 200.0 / 5.0) < 1e-6, rr["findings"] and rr["findings"][0]["magnitude"])

    # ---------------- engine: currency isolation, ordering, statuses
    eur = PG(w2, *b2, *rest, cur="EUR", camp=1)
    usd_only_good = PG(*b2, *rest, cur="USD", camp=1)
    out = ir.analyse_placements([eur, usd_only_good], WIN())
    add("engine: findings of a EUR group carry EUR only", all(f["currency"] == "EUR" for f in out["findings"] if f["placement"] and f["placement"]["key"] == "waste"))
    realloc = [f for f in out["findings"] if f["id"] == "placement_reallocation"]
    add("engine: a waste in EUR and efficiency in USD of the same campaign never make an estimate (currencies are separate groups)", len(realloc) == 1 and realloc[0]["currency"] == "EUR", [(f["currency"], f["what"]) for f in realloc])
    mixed_split = ir.analyse_placements([PG(w2, *rest, cur="EUR", camp=5), PG(*b2, cur="USD", camp=5)], WIN())
    add("engine: waste in one currency and efficiency in the other: no estimate at all", not [f for f in mixed_split["findings"] if f["id"] == "placement_reallocation"])
    add("engine: every placement rule gets a status (ran or skipped with a reason)", [s["id"] for s in out["rules"]] == [k for k, _t in ir.PLACEMENT_RULES] and all(s["status"] in ("ran", "skipped") for s in out["rules"]))
    add("engine: the findings are severity-sorted", [f["severity"] for f in out["findings"]] == sorted((f["severity"] for f in out["findings"]), key=lambda s: {"crit": 0, "warn": 1, "info": 2}[s]))
    add("engine: no data yields no finding and skipped rules with a reason", ir.analyse_placements([], WIN())["findings"] == [] and all(s["status"] == "skipped" and s["why"] for s in ir.analyse_placements([], WIN())["rules"]))
    add("engine: a thin-data finding is never louder than warn, and an estimate is never louder than info", all(f["severity"] != "crit" for f in out["findings"] if f["confidence"] == "thin data") and all(f["severity"] == "info" for f in out["findings"] if f["estimate"]))
    add("engine: every finding has evidence, an action and a confidence", all(f["evidence"] and f["action"] and f["confidence"] for f in out["findings"]))
    add("engine: the six campaign rules are untouched (their list still has six rules and no placement rule)", len(ir.RULES) == 6 and not any(k.startswith("placement") for k, _t in ir.RULES), [k for k, _t in ir.RULES])

    # ---------------- thresholds are named constants and the drawer text is generated from them
    thr = ir.placement_thresholds()
    consts = {n: getattr(ir, n) for n in dir(ir) if n.startswith("PLC_")}
    add("thresholds: every PLC_ constant is listed once with its live value", {t["name"] for t in thr} == set(consts) and all(t["value"] == consts[t["name"]] for t in thr), sorted(set(consts) ^ {t["name"] for t in thr}))
    add("thresholds: each belongs to one of the five placement rules", {t["rule"] for t in thr} <= {k for k, _t in ir.PLACEMENT_RULES})
    texts = ir.placement_rule_texts()
    with _patched(ir, PLC_EXCL_MIN_SHARE_PCT=7.5):
        add("thresholds: patching a constant changes the generated text and the behaviour", "7.5%" in ir.placement_rule_texts()[0]["text"] and not [x for x in _rule(ir.rule_placement_exclude, PS("waste", clicks=300, spend=45.0, impressions=60000), *base)["findings"] if x["placement"]["key"] == "waste"])
    add("thresholds: ... and the same 45 of spend (7%) fires again at the real 5% threshold", bool([x for x in _rule(ir.rule_placement_exclude, PS("waste", clicks=300, spend=45.0, impressions=60000), *base)["findings"] if x["placement"]["key"] == "waste"]))
    add("thresholds: the estimate text states its assumption", "assumption" in texts[4]["text"].lower() and "not a forecast" in texts[4]["text"])
    add("thresholds: five texts, one per rule, in the rule order", [t["id"] for t in texts] == [k for k, _t in ir.PLACEMENT_RULES])


# =========================================================================================== request / builders / files (pure)
def _check_pure_feed(add) -> None:
    from desktop import channels_data as cd
    from desktop import placements_data as pld
    from tests.insights_scenarios import _no_db_engine
    K = pld.KNOWN_PARAMS
    add("params: the placement feed accepts the channel page's names plus psort and pdir", set(K) == set(cd.KNOWN_PARAMS) | {"psort", "pdir"}, K)
    add("params: psort defaults to spend, descending", pld.parse_psort({})[:2] == ("spend", "desc"))
    add("params: cpa's first direction is ascending (cheaper first)", pld.parse_psort({"psort": ["cpa"]})[:2] == ("cpa", "asc"))
    for k in pld.PSORTS:
        add(f"params: psort={k} is accepted", pld.parse_psort({"psort": [k]})[2] == [])
    for label, params in (("psort not in the list", {"psort": ["nope"]}), ("psort with SQL", {"psort": ["spend;DROP"]}), ("pdir not asc/desc", {"pdir": ["up"]}), ("psort twice", {"psort": ["spend", "clicks"]})):
        _k, _d, probs = pld.parse_psort(params)
        add(f"params: {label} is a problem naming the parameter", len(probs) >= 1 and all(p["param"] in ("psort", "pdir") for p in probs), probs)
    add("params: a hostile psort value is echoed at most truncated to 60 characters, next to the parameter's name", all(len(p["value"]) <= 60 and p["param"] == "psort" for p in pld.parse_psort({"psort": ["x" * 200]})[2]))

    boom = _no_db_engine()
    with _patched(cd, engine=boom):
        st = pld.PlacementsStore(ttl=60)
        for label, params in (("an unknown parameter name", {"chanel": ["x"]}), ("part (CSV of another page)", {"part": ["daily"]}), ("a hostile date", {"from": ["2026-02-30"]}),
                              ("a bad psort", {"psort": ["nope"]}), ("a bad pdir", {"pdir": ["sideways"]}), ("a currency shape", {"currency": ["eur"]}), ("a hostile sort", {"sort": [";DROP"]})):
            code, body = st.get(params)
            add(f"guard: {label} is a 400 before the cache and before any database access", code == 400 and body["ok"] is False and body.get("problems"), (code, body))
            code, body, _h = st.export_csv(params)
            add(f"guard: {label} is a 400 on the CSV too", code == 400 and body["ok"] is False, (code, body))
        code, body = st.get({"chanel": ["SECRETVALUE"]})
        add("guard: the 400 names the parameter and never echoes its value", "chanel" in str(body) and "SECRETVALUE" not in str(body), body)
        add("guard: no request above touched the engine and nothing was cached", boom.touched == 0 and not st._cache, (boom.touched, list(st._cache)))
        for known in ("sort", "dir"):
            add(f"guard: the channel table's own '{known}' is a known name here (the page address carries it)", pld.cd.unknown_params({known: ["x"]}, K) == [], None)

    # ---- static files, wiring, honesty of the JS
    static = ROOT / "desktop" / "static"
    js = (static / "placements.js").read_text(encoding="utf-8")
    css = (static / "placements.css").read_text(encoding="utf-8")
    shell = (static / "shell.html").read_text(encoding="utf-8")
    chjs = (static / "channels.js").read_text(encoding="utf-8")
    add("static: shell.html loads placements.js and placements.css after their siblings", "/static/placements.js" in shell and "/static/placements.css" in shell and shell.index("insights.js") < shell.index("placements.js"))
    add("static: the panel has its host, its heading and a disabled-until-ready CSV button", 'id="plBody"' in shell and 'id="plTitle"' in shell and re.search(r'id="plCsv"[^>]*disabled', shell) is not None)
    add("static: channels.js calls the panel after every render and on refresh", "channelsPlacements.update(d)" in chjs and "channelsPlacements.invalidate()" in chjs)
    add("static: the URL comes from the channels payload's export links (L-101)", "d.export.placements" in js and "d.export.csv" in js)
    add("static: the CSV button is disabled while a request is pending (L-101) and its handle is not a stale timer", "S.busy" in js and "S.loadedUrl === S.url" in js)
    add("static: every server string goes through esc() (no raw payload text in an innerHTML)", js.count("esc(") > 30 and "innerHTML = d." not in js)
    add("static: unknown values are drawn as a muted dash with a reason, never as 0", "ch-no" in js and "not measured" in js and "unknown" in js)
    add("static: the file computes no rate or ratio (the server decides every number)", not re.search(r"\bclicks\s*/\s*impressions|\bspend\s*/\s*conversions|\bviewable\s*/", js, re.I))
    add("static: money is drawn with its currency code, never converted", "ch-mny" in js and "toFixed" not in js)
    add("static: the file parses (node --check)", _node_check(static / "placements.js"), None)
    add("static: no hardcoded font name (the font is one token per surface)", not re.search(r"font-family\s*:\s*(?!var\()\S", css) and not re.search(r"Montserrat|Segoe|Consolas|Arial|Helvetica", css))
    add("static: the CSS uses only design tokens for colour where a token exists", "var(--line)" in css and "var(--surface)" in css)
    add("static: the mobile layout stacks the table like the channels table", "@media (max-width: 720px)" in css)
    add("static: server.py routes both endpoints", '"/api/channels/placements"' in (ROOT / "desktop" / "server.py").read_text(encoding="utf-8") and '"/api/channels/placements.csv"' in (ROOT / "desktop" / "server.py").read_text(encoding="utf-8"))
    add("static: the Insights cards mark an estimate", "f.estimate" in (static / "insights.js").read_text(encoding="utf-8"))

    # ---- builders
    slots = [{"placement": "a.example", "key": "a_x", "type": "website", "ad_size": "", "position": "unknown", "currency": "EUR", "campaign_ids": {1}, "impressions": 1000, "clicks": 3,
              "conversions": 0, "spend_micros": 0, "revenue_micros": 0, "viewable": 0, "viewable_base": 0}]
    c = pld.slot_cells(slots[0])
    add("cells: no ad size and position are None (blank), not 0 or 'unknown'", c["ad_size"] is None and c["position"] is None)
    add("cells: no viewability measured is state none with no pct", c["viewability"]["state"] == "none" and c["viewability"]["pct"] is None and c["viewable_impressions"] is None)
    add("cells: 3 clicks of 1000 impressions is a percentage only above LOW_N; here 1000 impressions is enough", c["ctr"]["pct"] is not None and c["ctr"]["low_n"] is False)
    add("cells: no spend is state none, no currency shown, cost per conversion unknown", c["spend"]["state"] == "none" and c["currency"] is None and c["cpa"]["value"] is None, c)
    small = pld.slot_cells({**slots[0], "impressions": 3, "clicks": 1})
    add("cells: under LOW_N impressions the CTR reads as counts (1 of 3), no percentage", small["ctr"]["pct"] is None and small["ctr"]["low_n"] is True and (small["ctr"]["n"], small["ctr"]["of"]) == (1, 3), small["ctr"])
    none = pld.slot_cells({**slots[0], "impressions": 0, "clicks": 0})
    add("cells: no impressions means no CTR (why: no impressions)", none["ctr"]["pct"] is None and none["ctr"]["why"] == "no impressions", none["ctr"])
    v = pld.slot_cells({**slots[0], "impressions": 1000, "viewable": 300, "viewable_base": 500, "spend_micros": 5 * M})
    add("cells: viewability is viewable / MEASURED impressions, with the measured share", v["viewability"]["pct"] == 60.0 and v["viewability"]["coverage_pct"] == 50.0, v["viewability"])
    add("cells: a row with spend shows its currency; cost per conversion needs conversions", v["currency"] == "EUR" and v["cpa"]["value"] is None)
    add("cells: cost per conversion is spend / conversions in the row's currency", pld.slot_cells({**slots[0], "spend_micros": 10 * M, "conversions": 4})["cpa"]["value"] == 2.5)

    rows = [pld.slot_cells({**slots[0], "placement": n, "spend_micros": s * M, "clicks": c_, "impressions": 1000, "currency": cur, "conversions": cv, "key": n})
            for n, s, c_, cur, cv in (("a", 5, 10, "EUR", 1), ("b", 3, 20, "EUR", 2), ("c", 900, 5, "VND", 1), ("d", 0, 7, "EUR", 0), ("e", 100, 30, "USD", 5))]
    order = [r["placement"] for r in pld.sort_rows(rows, "spend", "desc")]
    add("sort: spend orders WITHIN each currency (EUR, USD, VND in code order), never across", order == ["a", "b", "e", "c", "d"], order)
    add("sort: a row with no spend sorts last, not first", order[-1] == "d")
    add("sort: clicks descending", [r["placement"] for r in pld.sort_rows(rows, "clicks", "desc")] == ["e", "b", "a", "d", "c"])
    add("sort: placement name ascending is case-insensitive and stable", [r["placement"] for r in pld.sort_rows(rows, "placement", "asc")] == ["a", "b", "c", "d", "e"])
    add("sort: cost per conversion ascending puts unknown (no conversions) last", pld.sort_rows(rows, "cpa", "asc")[-1]["placement"] == "d")
    tot = pld.totals_by_currency([{**slots[0], "currency": cur, "spend_micros": s * M, "clicks": 1, "impressions": 100} for cur, s in (("EUR", 5), ("EUR", 7), ("USD", 3), ("VND", 0))])
    add("totals: one line per currency, sorted by code, money never added across currencies", [t["code"] for t in tot] == ["EUR", "USD", "VND"] and tot[0]["spend"]["value"] == 12.0 and tot[1]["spend"]["value"] == 3.0, [(t["code"], t["spend"]) for t in tot])
    add("totals: a currency with no money is counts only (no currency shown)", tot[2]["currency"] is None and tot[2]["spend"]["state"] == "none")
    add("csv rows: the header set is the documented one and ends with the source table", pld.PLACEMENT_CSV[-1] == "source_table" and "viewability_measured_impressions" in pld.PLACEMENT_CSV)
    sr = pld.sheet_rows([pld.slot_cells({**slots[0], "impressions": 1000, "clicks": 50, "spend_micros": 5 * M, "viewable": 300, "viewable_base": 500})], type("V", (), {"d_from": date(2026, 9, 1), "d_to": date(2026, 9, 2)})(), "UTC")[0]
    hi = {h: i for i, h in enumerate(pld.PLACEMENT_CSV)}
    add("sheet: percentages are fractions for Excel (L-236: 5% is 0.05), counts and money are not touched", abs(sr[hi["ctr_pct"]] - 0.05) < 1e-9 and abs(sr[hi["viewability_pct"]] - 0.6) < 1e-9 and sr[hi["clicks"]] == 50 and sr[hi["spend"]] == 5.0, sr)
    add("sheet kinds: date, int, pct, money, bool, text are assigned to every column", {pld.sheet_kind(h) for h in pld.PLACEMENT_CSV} <= {"date", "int", "pct", "money", "bool", "text"})
    add("headline: every state has a sentence and none of the setup states carries a digit", all(pld._headline(s, [], 0, 0, None)["text"] for s in ("not_installed", "placements_not_installed", "empty", "no_match", "outside_filter", "unavailable")) and not re.search(r"\d", pld._headline("empty", [], 0, 0, None)["text"]))
    setup = pld.assemble_setup("placements_not_installed", {"tz": "UTC", "schema": "x"})
    add("setup: not installed carries the exact psql command and no number, no export", setup["install"]["steps"][0]["command"] == pld.rp.PSQL_11 and setup["export"] is None and setup["table"]["rows"] == [])
    add("setup: empty carries the load, rollup and synthetic-file commands", len(pld.assemble_setup("empty", {})["install"]["steps"]) == 3)
    add("definitions: the drawer text states blank = unknown, per-currency money and the cap from the constants", any(str(pld.rp.MAX_PLACEMENTS_PER_CAMPAIGN) in d["text"] for d in pld.build_definitions()) and any("Blank means unknown" in d["term"] for d in pld.build_definitions()))


def _node_check(path: Path) -> bool:
    try:
        r = subprocess.run(["node", "--check", str(path)], capture_output=True, text=True, timeout=30)
        return r.returncode == 0
    except (OSError, subprocess.SubprocessError):
        return True                      # no node on this machine: the browser check covers it


# =========================================================================================== the database fixture
DDL = ("db/sql/07_marketing_schema.sql", "db/sql/10_marketing_currency.sql", "db/sql/09_marketing_rollup.sql")
MIG_11 = "db/sql/11_marketing_placements.sql"


def _install(engine, schema: str, with_11: bool) -> None:
    """The REAL migrations into a throwaway schema, through search_path (public is not on it: an unqualified name cannot reach it)."""
    from sqlalchemy import text
    from erp.marketing import sample_check
    if with_11:
        with engine.connect() as c, contextlib.redirect_stdout(io.StringIO()):
            sample_check.install(c, schema, sample_check.Report())
        with engine.connect() as c:
            c.exec_driver_sql(f"SET search_path TO {schema}")
            c.exec_driver_sql((ROOT / MIG_11).read_text(encoding="utf-8"))
            c.exec_driver_sql("RESET search_path")
            c.commit()
    else:
        with engine.connect() as c, contextlib.redirect_stdout(io.StringIO()):
            sample_check.install(c, schema, sample_check.Report())
    with engine.connect() as c:
        assert c.execute(text("SELECT to_regclass(:q) IS NOT NULL"), {"q": f"{schema}.interaction_fact"}).scalar()


@contextlib.contextmanager
def _throwaway(engine, tag: str, with_11: bool = True):
    from sqlalchemy import text
    schema = f"perf_plc_{tag}_{uuid.uuid4().hex[:6]}"
    try:
        _install(engine, schema, with_11)
        yield schema
    finally:
        with engine.connect() as c:
            c.exec_driver_sql(f"DROP SCHEMA IF EXISTS {schema} CASCADE")
            c.commit()
        with engine.connect() as c:
            left = c.execute(text("SELECT count(*) FROM pg_namespace WHERE nspname = :s"), {"s": schema}).scalar()
        if left:
            raise RuntimeError(f"the throwaway schema {schema} was not dropped")


PLACES = ([(f"site-{i:02d}.example", "website") for i in range(7)] + [(f"app:com.example.app{i}", "app") for i in range(3)] + [("video-hub.example/embed", "video")])
DAYS = [date(2026, 9, 1) + timedelta(days=d) for d in range(14)]
SIZES = ("300x250", "728x90", "320x50", "")
POSITIONS = ("above_fold", "below_fold", "unknown")


def _row(day, camp, group, place, ptype, size, pos, device, impr, clicks, cost, conv, value, viewable, cur):
    return {"Date": day.isoformat(), "Campaign": camp, "Ad group": group, "Placement": place, "Placement type": ptype, "Ad size": size, "Position": pos, "Device": device,
            "Impressions": str(impr), "Clicks": str(clicks), "Cost": f"{cost:.2f} {cur}", "Conversions": str(conv),
            "Conversion value": "" if value is None else f"{value:.2f} {cur}", "Viewable impressions": "" if viewable is None else str(viewable)}


def make_plan(seed: int = 11) -> dict:
    """Synthetic EUR and USD rows with designed findings: a wasteful placement, an efficient one and a low-viewability one in campaign Alpha."""
    rng = random.Random(seed)
    eur, usd = [], []
    for d in DAYS:
        for camp in ("Synthetic Alpha", "Synthetic Beta"):
            for i, (name, ptype) in enumerate(PLACES):
                if rng.random() < 0.45:
                    continue
                impr = rng.randrange(1500, 9000)
                clicks = rng.randrange(5, max(6, impr // 60))
                cost = round(rng.uniform(4, 16), 2)
                conv = rng.randrange(0, 3)
                pos = rng.choice(POSITIONS)
                viewable = None if (camp == "Synthetic Beta" or pos == "unknown") else int(impr * rng.uniform(0.5, 0.8))
                eur.append(_row(d, camp, f"Group {i % 2}", name, ptype, rng.choice(SIZES), pos, rng.choice(("mobile", "desktop", "")), impr, clicks, cost, conv,
                                round(conv * rng.uniform(20, 40), 2) if rng.random() < 0.7 else None, viewable, "EUR"))
        eur.append(_row(d, "Synthetic Alpha", "Group 0", "waste.example", "website", "300x250", "above_fold", "mobile", 20000, 60, 30.0, 0, None, 14000, "EUR"))
        eur.append(_row(d, "Synthetic Alpha", "Group 0", "star.example", "website", "300x250", "above_fold", "mobile", 3000, 90, 20.0, 5, 150.0, 2400, "EUR"))
        eur.append(_row(d, "Synthetic Alpha", "Group 0", "unseen.example", "website", "728x90", "below_fold", "desktop", 30000, 150, 60.0, 3, 90.0, 6000, "EUR"))
    for d in DAYS[:3]:                                    # the same placement in a second currency: two rows, never one
        usd.append(_row(d, "Synthetic Alpha", "Group 9", "shared.example", "website", "300x250", "above_fold", "desktop", 4000, 40, 25.0, 2, 80.0, 3000, "USD"))
        usd.append(_row(d, "Synthetic Alpha", "Group 9", "usd-only.example", "app", "320x50", "", "mobile", 2500, 25, 12.0, 1, None, None, "USD"))
    eur.append(_row(DAYS[0], "Synthetic Alpha", "Group 0", "shared.example", "website", "300x250", "above_fold", "mobile", 4000, 40, 25.0, 2, 80.0, 3000, "EUR"))
    eur.append(_row(DAYS[0], "Synthetic Alpha", "Group 0", "shared.example", "website", "300x250", "above_fold", "mobile", 4100, 41, 26.0, 2, 82.0, 3100, "EUR"))   # a same-key duplicate: both kept
    old = [_row(date(2025, 3, 10) + timedelta(days=k), "Synthetic Alpha", "Group 0", "old.example", "website", "300x250", "above_fold", "mobile", 2000, 20, 9.0, 1, 30.0, 1500, "EUR") for k in range(3)]
    bad = [dict(_row(DAYS[0], "Synthetic Alpha", "Group 0", "bad1.example", "website", "300x250", "above_fold", "mobile", 100, 200, 1.0, 0, None, None, "EUR")),        # clicks > impressions
           dict(_row(DAYS[0], "Synthetic Alpha", "Group 0", "bad\x00.example", "website", "300x250", "above_fold", "mobile", 100, 2, 1.0, 0, None, None, "EUR")),         # NUL
           dict(_row(DAYS[0], "Synthetic Alpha", "Group 0", "bad3.example", "banner", "300x250", "above_fold", "mobile", 100, 2, 1.0, 0, None, None, "EUR")),            # type
           dict(_row(DAYS[0], "Synthetic Alpha", "Group 0", "bad4.example", "website", "300x250", "above_fold", "mobile", 100, 2, 1.0, 2, None, None, "EUR"), Conversions="2.5"),
           dict(_row(DAYS[0], "Synthetic Alpha", "Group 0", "bad5.example", "website", "300x250", "above_fold", "mobile", 100, 2, 1.0, 0, None, None, "EUR"), Cost="1.00")]   # no currency
    search = [_row(DAYS[k], "Synthetic Search", "Group 0", "search-partner.example", "website", "728x90", "above_fold", "desktop", 3000, 30, 14.0, 1, 40.0, 2100, "EUR") for k in range(3)]
    return {"eur": eur + old, "usd": usd, "bad": bad, "search": search}


def _csv_text(rows) -> str:
    import csv
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\n")
    w.writerow(_pc().HEADER)
    for r in rows:
        w.writerow([r.get(h, "") for h in _pc().HEADER])
    return out.getvalue()


def _expected(rows) -> dict:
    """{(day, campaign, placement, ptype, size, position, currency): sums} recomputed from the generated rows with no help from the code under test."""
    out = {}
    for r in rows:
        cur = r["Cost"].split()[-1]
        k = (r["Date"], r["Campaign"], r["Placement"], r["Placement type"], r["Ad size"].lower().replace(" ", ""), r["Position"] or "unknown", cur)
        g = out.setdefault(k, {"n": 0, "impr": 0, "clicks": 0, "conv": 0, "spend": 0, "rev": 0, "view": 0, "base": 0})
        g["n"] += 1
        g["impr"] += int(r["Impressions"])
        g["clicks"] += int(r["Clicks"])
        g["conv"] += int(r["Conversions"])
        g["spend"] += round(float(r["Cost"].split()[0]) * 100) * 10_000
        if r["Conversion value"]:
            g["rev"] += round(float(r["Conversion value"].split()[0]) * 100) * 10_000
        if r["Viewable impressions"]:
            g["view"] += int(r["Viewable impressions"])
            g["base"] += int(r["Impressions"])
    return out


def _rollup_rows(conn, schema: str) -> dict:
    from sqlalchemy import text
    rows = conn.execute(text(f"""
        SELECT r.event_date, cp.name AS campaign, COALESCE(cr.name, '(other)') AS placement, r.placement_type, r.ad_size, r.position, r.currency,
               r.fact_rows, r.impressions, r.clicks, r.conversions, r.spend_micros, r.revenue_micros, r.viewable_impressions, r.viewability_base_impressions
          FROM {schema}.interaction_placement_rollup r
          LEFT JOIN {schema}.marketing_campaign cp ON cp.campaign_id = r.campaign_id
          LEFT JOIN {schema}.marketing_creative cr ON cr.creative_id = r.placement_id""")).all()
    return {(x[0].isoformat(), x[1], x[2], x[3], x[4], x[5], x[6]): {"n": x[7], "impr": x[8], "clicks": x[9], "conv": x[10], "spend": x[11], "rev": x[12], "view": x[13], "base": x[14]} for x in rows}


class Box:
    def __init__(self, tmp: Path, name: str) -> None:
        self.root = tmp / name
        self.inbox = self.root / "incoming"
        self.inbox.mkdir(parents=True)
        self.processed = self.root / "processed"
        self.failed = self.root / "failed"

    def put(self, name: str, content, age: float = 60.0) -> Path:
        p = self.inbox / name
        p.write_bytes(content if isinstance(content, bytes) else content.encode("utf-8"))
        t = time.time() - age
        os.utime(p, (t, t))
        return p


def _table_hash(conn, schema: str, table: str, skip=("refreshed_at",)) -> str:
    from sqlalchemy import text
    cols = [r[0] for r in conn.execute(text("SELECT column_name FROM information_schema.columns WHERE table_schema = :s AND table_name = :t ORDER BY ordinal_position"), {"s": schema, "t": table}) if r[0] not in skip]
    rows = conn.execute(text(f"SELECT {', '.join(cols)} FROM {schema}.{table} ORDER BY {', '.join(cols)}")).all()
    return hashlib.sha256(repr(rows).encode("utf-8")).hexdigest()


def _check_database(add, tmp: Path) -> None:
    from sqlalchemy import text

    from desktop import channels_data as cd
    from desktop import insights_data as idt
    from desktop import placements_data as pld
    from desktop import placements_queries as pq
    from erp import config as cfg
    from erp.db import engine
    from erp.marketing import autorun, rollup
    from erp.marketing import rollup_placements as rp
    from tests.insights_scenarios import _no_db_engine

    with engine.connect() as c:
        pub_tables = sorted(r[0] for r in c.execute(text("SELECT table_name FROM information_schema.tables WHERE table_schema = 'public'")))
        pub_facts = c.execute(text("SELECT count(*) FROM public.interaction_fact")).scalar() if "interaction_fact" in pub_tables else None
        pub_placement = c.execute(text("SELECT to_regclass('public.interaction_placement_rollup') IS NOT NULL")).scalar()

    plan = make_plan()
    eur_csv, usd_csv = _csv_text(plan["eur"] + plan["bad"]), _csv_text(plan["usd"])
    good_all = plan["eur"] + plan["usd"] + plan["search"]

    with _throwaway(engine, "a") as schema, _throwaway(engine, "b", with_11=False) as schema_no11:
        box = Box(tmp, "e2e")
        box.put("display_eur.csv", eur_csv)
        box.put("display_usd.csv", usd_csv)
        box.put("search__eur__partner.csv", _csv_text(plan["search"]))
        put_hash = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in box.inbox.glob("*.csv")}

        # ---------------- end to end: the REAL autorun -> pipeline -> rollup on a temp inbox
        dry = autorun.run_inbox(box.inbox, schema=schema, dry_run=True)
        got = {o.name: (o.connector, o.rows_read, o.rows_loaded, o.status) for o in dry.outcomes}
        add("e2e dry run: both synthetic files are detected as placement_performance and validated (bad rows counted) without touching the database",
            got["display_eur.csv"][0] == "placement_performance" and got["display_eur.csv"][1] == len(plan["eur"]) + 5 and got["display_usd.csv"][:2] == ("placement_performance", len(plan["usd"])) and got["search__eur__partner.csv"][:2] == ("placement_performance", 3), got)
        with engine.connect() as c:
            add("e2e dry run: the throwaway schema still has no fact row and no ingest run", c.execute(text(f"SELECT count(*) FROM {schema}.interaction_fact")).scalar() == 0 and c.execute(text(f"SELECT count(*) FROM {schema}.marketing_ingest_run")).scalar() == 0)
        rep = autorun.run_inbox(box.inbox, schema=schema)
        st = {o.name: (o.status, o.connector, o.rows_read, o.rows_loaded, o.rows_rejected) for o in rep.outcomes}
        add("e2e: both files load through the existing pipeline and go to processed/", all(v[0] in ("ok", "partial") for v in st.values()) and len(list(box.processed.glob("*.csv"))) == 3, st)
        e = st["display_eur.csv"]
        add("e2e: exactly the five planted bad rows are rejected and every good row is loaded", (e[3], e[4]) == (len(plan["eur"]), 5), e)
        add("e2e: the processed files are byte-identical to what was dropped", {p.name.split("_", 1)[1]: hashlib.sha256(p.read_bytes()).hexdigest() for p in box.processed.glob("*.csv")} == {k: v for k, v in put_hash.items()}, None)
        add("e2e: the rollup ran once and reported ok", rep.rollup == "ok", (rep.rollup, rep.rollup_detail))
        with engine.connect() as c:
            n_fact = c.execute(text(f"SELECT count(*) FROM {schema}.interaction_fact")).scalar()
            add("e2e: the fact table holds exactly the good rows (a same-key duplicate is kept, both rows)", n_fact == len(good_all), (n_fact, len(good_all)))
            add("e2e: every fact row is a placement row of the placement_performance source", c.execute(text(f"SELECT count(*) FROM {schema}.interaction_fact f JOIN {schema}.marketing_source s ON s.source_id = f.source_id WHERE s.source_key = 'placement_performance' AND f.attrs->>'placement_row' = 'true'")).scalar() == n_fact)
            add("e2e: the channel is display (or the one the filename named) and each row keeps its own currency", c.execute(text(f"SELECT count(DISTINCT currency) FROM {schema}.interaction_fact")).scalar() == 2 and {r[0] for r in c.execute(text(f"SELECT channel_key FROM {schema}.marketing_channel"))} == {"display", "search"})
            add("e2e: the placements are creative rows, one per campaign x placement", c.execute(text(f"SELECT count(*) FROM {schema}.marketing_creative WHERE creative_key ~ '_[0-9a-f]{{8}}$'")).scalar() >= 12)
            add("e2e: no ingest run of the throwaway schema was left failed", c.execute(text(f"SELECT count(*) FROM {schema}.marketing_ingest_run WHERE status NOT IN ('ok', 'partial')")).scalar() == 0)

            # ---------------- rollup: every group equals the Python recomputation
            exp = _expected(good_all)
            got_r = _rollup_rows(c, schema)
            add("rollup: the group keys equal an independent recomputation from the generated rows", set(got_r) == set(exp), (len(got_r), len(exp), list(set(got_r) ^ set(exp))[:2]))
            diffs = [(k, exp[k], got_r[k]) for k in exp if k in got_r and exp[k] != got_r[k]]
            add("rollup: every group's fact rows, impressions, clicks, conversions, spend, revenue, viewable and measured base equal the recomputation", not diffs, diffs[:1])
            for label, col, key in (("impressions", "impressions", "impr"), ("clicks", "clicks", "clicks"), ("conversions", "conversions", "conv"), ("spend", "spend_micros", "spend"), ("revenue", "revenue_micros", "rev")):
                tot = c.execute(text(f"SELECT sum({col}) FROM {schema}.interaction_placement_rollup")).scalar()
                fact = c.execute(text(f"SELECT sum({col}) FROM {schema}.interaction_fact")).scalar()
                add(f"rollup: the total {label} equals the fact table's and the generated rows'", tot == fact == sum(g[key] for g in exp.values()), (tot, fact))
            add("rollup: the same-key duplicate rows are both in the rollup (2 fact rows in one group)", any(g["n"] == 2 for g in got_r.values()) and sum(g["n"] for g in got_r.values()) == n_fact)
            add("rollup: unknown ad size is stored '' and unknown position 'unknown'; nothing is invented", c.execute(text(f"SELECT count(*) FROM {schema}.interaction_placement_rollup WHERE ad_size = '' OR position = 'unknown'")).scalar() > 0)
            add("rollup: a measured base only counts the rows that reported viewability (never all impressions)", all(g["base"] <= g["impr"] for g in got_r.values()) and any(g["base"] < g["impr"] for g in got_r.values()))
            add("rollup: the placement table is empty of the cardinality bucket at this size (12 placements)", c.execute(text(f"SELECT count(*) FROM {schema}.interaction_placement_rollup WHERE placement_id = 0")).scalar() == 0)
            add("rollup: it never holds a row for a source that is not a placement source", c.execute(text(f"SELECT sum(fact_rows) FROM {schema}.interaction_placement_rollup")).scalar() == n_fact)

            # ---------------- existing rollups are identical with and without migration 11
        box2 = Box(tmp, "e2e_no11")
        box2.put("display_eur.csv", eur_csv)
        box2.put("display_usd.csv", usd_csv)
        box2.put("search__eur__partner.csv", _csv_text(plan["search"]))
        rep2 = autorun.run_inbox(box2.inbox, schema=schema_no11)
        add("no migration 11: the same files load and the rollup reports ok (the placement step is skipped, nothing fails)", rep2.rollup == "ok" and all(o.status in ("ok", "partial") for o in rep2.outcomes), (rep2.rollup, rep2.rollup_detail))
        with engine.connect() as c:
            add("no migration 11: the placement table does not exist and nothing created it", c.execute(text("SELECT to_regclass(:q) IS NULL"), {"q": f"{schema_no11}.interaction_placement_rollup"}).scalar())
            for table in ("interaction_daily_rollup", "interaction_daily_channel_rollup"):
                a, b = _table_hash(c, schema, table), _table_hash(c, schema_no11, table)
                add(f"existing rollup output is identical with and without migration 11: {table}", a == b, (a[:12], b[:12]))
            add("existing rollup: the fact tables of the two schemas hold the same numbers", c.execute(text(f"SELECT sum(impressions), sum(spend_micros), count(*) FROM {schema}.interaction_fact")).one() == c.execute(text(f"SELECT sum(impressions), sum(spend_micros), count(*) FROM {schema_no11}.interaction_fact")).one())
        summary = rollup.refresh(engine, schema_no11)
        add("no migration 11: a rollup run says the placement rollup is not installed in its summary path", summary.placement is not None and summary.placement.installed is False and "not installed" in summary.placement.summary(), summary.placement)
        with engine.connect() as c:
            try:
                c.exec_driver_sql(f"SET search_path TO {schema}")
                c.exec_driver_sql((ROOT / MIG_11).read_text(encoding="utf-8"))
                c.exec_driver_sql("RESET search_path")
                c.commit()
                again_ok = True
            except Exception as exc:  # noqa: BLE001
                again_ok = str(exc)
            add("rollup migration: running db/sql/11 a second time is a no-op (IF NOT EXISTS, no error)", again_ok is True, again_ok)
            add("rollup migration: it created only its own table, index and constraints (the existing tables are the ones 07 + 09 + 10 made)", c.execute(text("SELECT count(*) FROM pg_tables WHERE schemaname = :s AND tablename = 'interaction_placement_rollup'"), {"s": schema}).scalar() == 1)

        # ---------------- idempotence and the incremental refresh
        with engine.connect() as c:
            c.exec_driver_sql(f"UPDATE {schema}.interaction_fact SET loaded_at = now() - interval '3 days'")
            c.exec_driver_sql(f"UPDATE {schema}.marketing_rollup_run SET started_at = now() - interval '2 days'")
            c.exec_driver_sql(f"UPDATE {schema}.interaction_placement_rollup SET refreshed_at = now() - interval '2 days'")
            c.commit()
        with engine.connect() as c:
            h1 = _table_hash(c, schema, "interaction_placement_rollup")
            t_old = c.execute(text(f"SELECT min(refreshed_at) FROM {schema}.interaction_placement_rollup WHERE event_date = DATE '2025-03-10'")).scalar()
        again = rollup.refresh(engine, schema)
        with engine.connect() as c:
            h2 = _table_hash(c, schema, "interaction_placement_rollup")
            t_old2 = c.execute(text(f"SELECT min(refreshed_at) FROM {schema}.interaction_placement_rollup WHERE event_date = DATE '2025-03-10'")).scalar()
        add("incremental: a second refresh with nothing new leaves the placement rollup identical", h1 == h2, (h1[:10], h2[:10]))
        add("incremental: an old day that did not change is NOT recomputed (its refreshed_at is unchanged)", t_old == t_old2 and t_old is not None, (t_old, t_old2))
        add("incremental: the run reports its placement step", again.placement is not None and again.placement.installed and again.placement.bootstrapped is False, again.placement)
        restate = [dict(r) for r in plan["eur"] if r["Date"] == "2025-03-10"][:1]
        restate[0]["Impressions"] = "2500"
        restate[0]["Viewable impressions"] = "1600"
        box3 = Box(tmp, "restate")
        box3.put("restate.csv", _csv_text(restate))
        rep3 = autorun.run_inbox(box3.inbox, schema=schema)
        with engine.connect() as c:
            row = c.execute(text(f"SELECT impressions, viewable_impressions FROM {schema}.interaction_placement_rollup WHERE event_date = DATE '2025-03-10'")).one()
            t_other = c.execute(text(f"SELECT min(refreshed_at) FROM {schema}.interaction_placement_rollup WHERE event_date = DATE '2025-03-11'")).scalar()
            t_new = c.execute(text(f"SELECT min(refreshed_at) FROM {schema}.interaction_placement_rollup WHERE event_date = DATE '2025-03-10'")).scalar()
        add("incremental: a restated OLD day (same key, new numbers) is picked up by the next run", rep3.rollup == "ok" and row == (2500, 1600), (rep3.rollup, row))
        add("incremental: only the restated day was recomputed (a neighbouring old day kept its timestamp)", t_other == t_old2 and t_new > t_old2, (t_other, t_old2, t_new))

        # ---------------- bootstrap: table installed after the data
        with engine.connect() as c:
            c.exec_driver_sql(f"TRUNCATE {schema}.interaction_placement_rollup")
            c.commit()
        boot = rollup.refresh(engine, schema, days=1, use_touched=False)
        with engine.connect() as c:
            n_boot = c.execute(text(f"SELECT count(*) FROM {schema}.interaction_placement_rollup")).scalar()
            new_h = _table_hash(c, schema, "interaction_placement_rollup")
        add("bootstrap: an EMPTY placement table with placement facts on file is built from scratch on the next run (installing 11 after the data works)", boot.placement.bootstrapped and n_boot > 0, (boot.placement, n_boot))
        add("bootstrap: the rebuilt table has every group again (the restated old day now carries its new numbers, so the hash differs from before)", n_boot == len(exp) and new_h != h1, (n_boot, len(exp)))

        # ---------------- cardinality guard
        with engine.connect() as c:
            camp = c.execute(text(f"SELECT campaign_id FROM {schema}.marketing_campaign WHERE name = 'Synthetic Alpha'")).scalar()
            before_sum = c.execute(text(f"SELECT sum(impressions), sum(clicks), sum(spend_micros), sum(fact_rows) FROM {schema}.interaction_placement_rollup WHERE campaign_id = :c"), {"c": camp}).one()
        conn = engine.connect()
        try:
            with conn.begin():
                conn.execute(text("SELECT set_config('statement_timeout', '60000', true)"))
                rp.refresh_range(conn, schema, (date(2026, 9, 1), None), cap=3)
            capped = _rollup_rows(conn, schema)
            after_sum = conn.execute(text(f"SELECT sum(impressions), sum(clicks), sum(spend_micros), sum(fact_rows) FROM {schema}.interaction_placement_rollup WHERE campaign_id = :c AND event_date >= DATE '2026-09-01'"), {"c": camp}).one()
            base_sum = conn.execute(text(f"SELECT sum(impressions), sum(clicks), sum(spend_micros), count(*) FROM {schema}.interaction_fact WHERE campaign_id = :c AND event_date >= DATE '2026-09-01'"), {"c": camp}).one()
            other_rows = [(k, g) for k, g in capped.items() if k[2] == "(other)"]
            named = {k[2] for k in capped if k[1] == "Synthetic Alpha" and k[2] != "(other)" and k[0] >= "2026-09-01"}
            add("cardinality: with a cap of 3 each campaign keeps 3 named placements and folds the rest into one bucket", len(named) == 3 and len(other_rows) > 0, (len(named), len(other_rows)))
            add("cardinality: folding loses nothing - the campaign's sums equal the fact table's", tuple(int(x) for x in after_sum) == tuple(int(x) for x in base_sum), (after_sum, base_sum))
            add("cardinality: the kept placements are the FIRST ones by creative id (a rule that does not move when a later day is loaded)", named == {r[0] for r in conn.execute(text(f"SELECT name FROM {schema}.marketing_creative WHERE campaign_id = :c AND creative_key ~ '_[0-9a-f]{{8}}$' ORDER BY creative_id LIMIT 3"), {"c": camp})})
            folded = conn.execute(text(f"SELECT max(folded_placements) FROM {schema}.interaction_placement_rollup WHERE placement_id = 0")).scalar()
            add("cardinality: the bucket records how many distinct placements it folded that day", folded and folded >= 1, folded)
            add("cardinality: folded_total counts the placements beyond the cap for the whole campaign", rp.folded_total(conn, schema, cap=3) >= 8, rp.folded_total(conn, schema, cap=3))
            add("cardinality: at the real cap (500) nothing is folded in this fixture", rp.folded_total(conn, schema) == 0)
            conn.rollback()
            with conn.begin():
                conn.execute(text("SELECT set_config('statement_timeout', '60000', true)"))
                rp.refresh_range(conn, schema, (date(2025, 1, 1), None))
            restored = _rollup_rows(conn, schema)
            add("cardinality: the next refresh at the real cap restores every named placement", set(restored) == set(exp), (len(restored), len(exp)))
        finally:
            conn.close()
        with engine.connect() as c:
            tot_fact = c.execute(text(f"SELECT count(*) FROM {schema}.interaction_fact")).scalar()
        add("cardinality: the guard test wrote nothing to the fact table", tot_fact == n_fact, (tot_fact, n_fact))

        # ---------------- the store (page) on this schema
        with _patched(cfg, MARKETING_SCHEMA=schema):
            captured: list[str] = []
            from sqlalchemy import event

            def grab(conn_, cursor, statement, parameters, context, executemany):
                captured.append(statement)
            event.listen(cd.engine, "before_cursor_execute", grab)
            try:
                store = pld.PlacementsStore(ttl=0)
                code, p = store.get({}, fresh=True)
            finally:
                event.remove(cd.engine, "before_cursor_execute", grab)
            add("store: the payload is 200 and ready", code == 200 and p["state"] == "ready" and p["ok"], (code, p.get("state"), p.get("error")))
            add("store: it never read interaction_fact, the landing table or any raw table (every statement captured)", captured and not any(re.search(r"interaction_fact|marketing_landing", s) for s in captured), [s[:80] for s in captured if re.search(r"interaction_fact|marketing_landing", s)][:1])
            add("store: it says so in the payload", p["read"]["raw_fact_read"] is False and p["read"]["tables"] == ["interaction_placement_rollup"])
            add("store: the numbers come from the placement rollup table (it is among the captured statements)", any("interaction_placement_rollup" in s for s in captured))
            in_view = {k: g for k, g in exp.items() if p["range"]["from"] <= k[0] <= p["range"]["to"]}
            add("store: the default range is the last 30 days ending at the newest day of placement data, so the 2025 rows are outside it", len(in_view) < len(exp) and p["range"]["to"] == "2026-09-14", (p["range"], len(in_view), len(exp)))
            exp_by_cur = {}
            for k, g in in_view.items():
                e_ = exp_by_cur.setdefault(k[6], {"impr": 0, "clicks": 0, "conv": 0, "spend": 0})
                for m_ in ("impr", "clicks", "conv", "spend"):
                    e_[m_] += g[m_]
            tot = {t["code"]: t for t in p["totals"]}
            add("store: the totals are one line per currency (EUR and USD), never added", sorted(tot) == ["EUR", "USD"], sorted(tot))
            add("store: each currency's totals equal the recomputation from the generated rows", all(tot[cur]["impressions"] == e_["impr"] and tot[cur]["clicks"] == e_["clicks"] and tot[cur]["conversions"] == e_["conv"] and round(tot[cur]["spend"]["value"] * M) == e_["spend"] for cur, e_ in exp_by_cur.items()), [(c_, tot[c_]["spend"], exp_by_cur[c_]["spend"]) for c_ in tot])
            rows = p["table"]["rows"]
            all_rows = pld.build_slots(_raw_block(store, schema, {}))[0]
            add("store: the table lists rows of both currencies and the shared placement appears twice, once per currency", sum(1 for r in all_rows if r["placement"] == "shared.example") >= 2 and {r["currency"] for r in all_rows if r["placement"] == "shared.example"} == {"EUR", "USD"})
            add("store: the table shows at most the first 100 rows and says how many more", len(rows) <= pld.ROWS_SHOWN_MAX and p["table"]["total_rows"] == len(all_rows) and p["table"]["more"] == max(0, len(all_rows) - pld.ROWS_SHOWN_MAX))
            add("store: the rows are ordered by spend within each currency (EUR block, then USD)", [r["currency"] for r in rows if r["currency"]] == sorted(r["currency"] for r in rows if r["currency"]))
            eur_spend = [r["spend"]["value"] for r in rows if r["currency"] == "EUR"]
            add("store: within EUR the spend is descending", eur_spend == sorted(eur_spend, reverse=True), eur_spend[:5])
            blank = [r for r in all_rows if not r["ad_size"]]
            cells = [pld.slot_cells(r) for r in blank]
            add("store: rows with no ad size show it blank (None), never 0", blank and all(c_["ad_size"] is None for c_ in cells))
            add("store: rows with an unknown position show it blank (None)", any(pld.slot_cells(r)["position"] is None for r in all_rows))
            add("store: a placement that never reported viewability shows 'not measured', not 0%", any(pld.slot_cells(r)["viewability"]["state"] == "none" for r in all_rows) and all(pld.slot_cells(r)["viewability"]["pct"] is None for r in all_rows if r["viewable_base"] == 0))
            add("store: the export links carry the resolved range and the sort (L-101)", p["export"]["csv"].startswith("/api/channels/placements.csv?") and "from=" in p["export"]["csv"] and "psort=spend" in p["export"]["csv"], p["export"])
            add("store: the definitions and thresholds text are present", len(p["definitions"]) >= 8 and p["config"]["max_placements_per_campaign"] == 500)

            for key, d in (("spend", "desc"), ("spend", "asc"), ("clicks", "desc"), ("impressions", "asc"), ("ctr", "desc"), ("conversions", "desc"), ("cpa", "asc"), ("viewability", "desc"), ("placement", "asc")):
                code, ps = store.get({"psort": [key], "pdir": [d]}, fresh=True)
                order = [(r["placement"], r["ad_size"], r["position"], r["currency"]) for r in ps["table"]["rows"]]
                want = [(r["placement"], r["ad_size"], r["position"], r["currency"]) for r in pld.sort_rows([pld.slot_cells(s) for s in all_rows], key, d)][:pld.ROWS_SHOWN_MAX]
                add(f"store: psort={key} pdir={d} answers 200 and the order equals the independent sort", code == 200 and order == want and ps["sort"] == {"key": key, "dir": d}, (code, order[:2], want[:2]))
            code, ps = store.get({"psort": ["cpa"]}, fresh=True)
            eur_cpa = [r["cpa"]["value"] for r in ps["table"]["rows"] if r["currency"] == "EUR"]
            known = [x for x in eur_cpa if x is not None]
            add("store: cost per conversion ascending puts the known values in order and the unknown ones (no conversion) after them", known == sorted(known) and eur_cpa[:len(known)] == known, eur_cpa[:6])
            code, pe = store.get({"currency": ["EUR"]}, fresh=True)
            add("filters: a currency filter leaves only EUR rows and one total line", code == 200 and {r["currency"] for r in pe["table"]["rows"] if r["currency"]} == {"EUR"} and [t["code"] for t in pe["totals"]] == ["EUR"], [t["code"] for t in pe["totals"]])
            code, pc_ = store.get({"campaign": [str(_id(schema, "Synthetic Beta"))]}, fresh=True)
            beta = {k[2] for k in exp if k[1] == "Synthetic Beta"}
            add("filters: a campaign filter (Beta) lists exactly the placements Beta ran, and not the wasteful one of Alpha", code == 200 and {r["placement"] for r in pc_["table"]["rows"]} == beta and "waste.example" not in beta, (sorted(beta)[:3], code))
            code, pw = store.get({"from": ["2025-03-10"], "to": ["2025-03-12"]}, fresh=True)
            add("filters: a range holding only the old placement lists only it", code == 200 and [r["placement"] for r in pw["table"]["rows"]] == ["old.example"], [r["placement"] for r in pw["table"]["rows"]])
            code, pn = store.get({"from": ["2020-01-01"], "to": ["2020-01-05"]}, fresh=True)
            add("state: a range with no placement row says no_match, has no export link and no number in the headline", code == 200 and pn["state"] == "no_match" and not re.search(r"\d", pn["headline"]["text"]), (code, pn.get("state")))
            code, pf = store.get({"channel": ["search"], "focus": ["display"]}, fresh=True)
            add("state: a drill-down channel the channel filter excludes is outside_filter, not 'no data' (L-209)", code == 200 and pf["state"] == "outside_filter" and not pf["table"]["rows"], (code, pf.get("state")))
            c_x, _b, _h = store.export_csv({"channel": ["search"], "focus": ["display"]})
            add("state: ... and the CSV answers 409 for it, like the page (L-209)", c_x == 409, c_x)
            code, ps_ = store.get({"channel": ["search"]}, fresh=True)
            add("filters: a channel filter (search) lists only the placement that channel ran", code == 200 and [r["placement"] for r in ps_["table"]["rows"]] == ["search-partner.example"], [r["placement"] for r in ps_["table"]["rows"]])
            code, pb = store.get({"channel": ["display"], "focus": ["display"]}, fresh=True)
            add("filters: a channel filter and the same focus channel is fine", code == 200 and pb["state"] == "ready" and all(r["placement"] != "search-partner.example" for r in pb["table"]["rows"]), pb.get("state"))
            c2, csvb, hdr = store.export_csv({})
            add("csv: 200 with bytes, a safe attachment name and the row count header", c2 == 200 and isinstance(csvb, bytes) and re.fullmatch(r'attachment; filename="channels_placements_\d{8}_\d{4}(_filtered)?\.csv"', hdr["Content-Disposition"]) and hdr["X-Row-Count"].isdigit(), (c2, hdr))
            text_ = csvb.decode("utf-8-sig")
            lines = text_.split("\r\n")
            add("csv: UTF-8 with a BOM and CRLF, the documented header, one line per row", csvb.startswith(b"\xef\xbb\xbf") and lines[0] == ",".join(pld.PLACEMENT_CSV) and len([x for x in lines[1:] if x]) == int(hdr["X-Row-Count"]) == len(all_rows), (lines[0][:60], hdr["X-Row-Count"], len(all_rows)))
            import csv as _csv
            parsed = list(_csv.DictReader(io.StringIO(text_)))
            add("csv: blank means unknown - an unknown ad size and position are empty cells, not 0", any(r["ad_size"] == "" for r in parsed) and any(r["position"] == "" for r in parsed) and not any(r["ad_size"] == "0" for r in parsed))
            add("csv: viewability is empty when it was not measured", any(r["viewability_pct"] == "" for r in parsed) and any(r["viewability_pct"] != "" for r in parsed))
            add("csv: every money row names its currency (one code per row)", all(re.fullmatch(r"[A-Z]{3}|", r["currency"]) for r in parsed))
            add("csv: the same rows and order as the panel's sort (spend, currency blocks)", [r["placement"] for r in parsed][:5] == [r["placement"] for r in pld.sort_rows([pld.slot_cells(s) for s in all_rows], "spend", "desc")][:5])
            c3, csv_f, hdr_f = store.export_csv({"currency": ["USD"], "psort": ["clicks"]})
            add("csv: the same filters and sort as a link made from the payload (USD only, by clicks)", c3 == 200 and set(r["currency"] for r in _csv.DictReader(io.StringIO(csv_f.decode("utf-8-sig")))) == {"USD"} and hdr_f["Content-Disposition"].endswith('_filtered.csv"'), (c3, hdr_f))
            c4, body, _h = store.export_csv({"from": ["2020-01-01"], "to": ["2020-01-05"]})
            add("csv: nothing to export is a 409 with the reason, not an empty file", c4 == 409 and body["state"] == "no_match", (c4, body))
            with _patched(pld, EXPORT_MAX_ROWS=5):
                c5, cut, hdr_c = store.export_csv({})
            add("csv: a cut at the row cap is announced in a header", c5 == 200 and hdr_c["X-Export-Truncated"] == "true" and hdr_c["X-Row-Count"] == "5", hdr_c)

            # ---- channels payload carries the placements link (the L-101 source)
            cs = cd.ChannelsStore(ttl=0)
            code, chp = cs.get({}, fresh=True)
            add("channels page: its payload carries the validated placements link with the resolved range", code == 200 and chp["export"]["placements"].startswith("/api/channels/placements?") and "from=" in chp["export"]["placements"], chp.get("export"))
            add("channels page: nothing else in the channels payload changed shape (the same export keys plus placements)", set(chp["export"]) == {"channels", "daily", "campaigns", "xlsx", "insights", "placements"}, set(chp["export"]))
            add("channels page: the placement facts are ordinary rollup rows, so the page is ready with them", chp["state"] == "ready")

            # ---- insights: the placement rules on the throwaway data
            ins = idt.InsightsStore(ttl=0)
            code, ip = ins.get({}, fresh=True)
            fnd = [f for f in ip["findings"] if f["id"].startswith("placement")]
            add("insights: the placement rules ran on the synthetic data and are listed among the rules", code == 200 and ip["state"] == "ready" and {r["id"] for r in ip["rules"]} >= {k for k, _t in _ir().PLACEMENT_RULES}, (code, ip.get("state")))
            add("insights: the wasteful placement is named as an exclusion candidate with its evidence", any(f["id"] == "placement_exclude" and f["placement"]["name"] == "waste.example" for f in fnd), [(f["id"], f["placement"] and f["placement"]["name"]) for f in fnd][:6])
            add("insights: the efficient placement is a scale candidate, info", any(f["id"] == "placement_scale" and f["placement"]["name"] == "star.example" and f["severity"] == "info" for f in fnd))
            realloc = [f for f in fnd if f["id"] == "placement_reallocation"]
            add("insights: the reallocation is an ESTIMATE, flagged and carrying its assumption", bool(realloc) and realloc[0]["estimate"] is True and any("Assumption" in str(e) for e in realloc[0]["evidence"]), realloc and realloc[0]["what"])
            add("insights: every placement finding names one currency and its text never names the other", all(f["currency"] in ("EUR", "USD") and not ("USD" in str(f["what"]) and "EUR" in str(f["what"])) for f in fnd))
            add("insights: exclusion, viewability and estimate findings say that nothing is changed, and no finding claims that something already was", all("nothing" in f["action"].lower() for f in fnd if f["id"] in ("placement_exclude", "placement_viewability", "placement_reallocation")) and not any(" has been " in re.sub(r"[Nn]othing has been \w+", "", f["action"] + " " + f["what"]) for f in fnd))
            add("insights: the drawer thresholds and rule texts now include the placement rules (generated from the constants)", any(t["name"] == "PLC_EXCL_MIN_SHARE_PCT" for t in ip["how"]["thresholds"]) and any(r["id"] == "placement_exclude" for r in ip["how"]["rules"]))
            add("insights: the read list names the placement table", "interaction_placement_rollup" in ip["read"]["tables"] and ip["read"]["raw_fact_read"] is False)
            add("insights: the six campaign rules are still listed, first and in their old order, before the placement rules", [r["id"] for r in ip["rules"]][:6] == [k for k, _t in _ir().RULES] and [r["id"] for r in ip["rules"]][6:] == [k for k, _t in _ir().PLACEMENT_RULES], [r["id"] for r in ip["rules"]])
            code, ic = ins.get({"currency": ["USD"]}, fresh=True)
            add("insights: a currency filter limits the placement findings to that currency", all(f["currency"] in (None, "USD") for f in ic["findings"]), [f["currency"] for f in ic["findings"]])

            # ---- Excel
            c6, xb, xh = ins.export_xlsx({})
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(xb))
            add("xlsx: the workbook has a Placements sheet between Campaigns and Insights", c6 == 200 and wb.sheetnames == ["Summary", "Channels", "Campaigns", "Placements", "Insights", "Definitions"], wb.sheetnames)
            ws = wb["Placements"]
            head = [c_.value for c_ in ws[1]]
            add("xlsx Placements: the header is the CSV's column set and the header row is frozen and styled", head == pld.PLACEMENT_CSV and ws.freeze_panes == "A2" and ws["A1"].font.bold)
            add("xlsx Placements: one row per table row (every row, not only the first 100)", ws.max_row - 1 == len(all_rows), (ws.max_row - 1, len(all_rows)))
            ix = {h: i + 1 for i, h in enumerate(head)}
            r2 = {h: ws.cell(row=2, column=ix[h]) for h in head}
            add("xlsx Placements: counts are integers with a thousands format, dates are date cells", isinstance(r2["clicks"].value, int) and r2["clicks"].number_format == "#,##0" and r2["period_from"].number_format == "yyyy-mm-dd")
            add("xlsx Placements: rates are FRACTIONS with a percent format (L-236)", all(ws.cell(row=r, column=ix["ctr_pct"]).value is None or (0 <= ws.cell(row=r, column=ix["ctr_pct"]).value <= 1 and ws.cell(row=r, column=ix["ctr_pct"]).number_format == "0.0%") for r in range(2, ws.max_row + 1)))
            add("xlsx Placements: an unknown ad size / viewability is an EMPTY cell, not 0", any(ws.cell(row=r, column=ix["ad_size"]).value in (None, "") for r in range(2, ws.max_row + 1)) and any(ws.cell(row=r, column=ix["viewability_pct"]).value is None for r in range(2, ws.max_row + 1)))
            add("xlsx Placements: spend follows the row's currency format (EUR / USD two decimals)", all(ws.cell(row=r, column=ix["spend"]).value is None or ws.cell(row=r, column=ix["spend"]).number_format == "#,##0.00" for r in range(2, ws.max_row + 1)))
            add("xlsx Placements: every money row names one currency code", all(re.fullmatch(r"[A-Z]{3}|", str(ws.cell(row=r, column=ix["currency"]).value or "")) for r in range(2, ws.max_row + 1)))
            add("xlsx: the Definitions sheet lists the placement thresholds as numeric cells", any(r_[0].value == "Threshold" and str(r_[1].value).startswith("PLC_") and isinstance(r_[2].value, (int, float)) for r_ in wb["Definitions"].iter_rows(min_row=2)))
            add("xlsx: no cell of any sheet is a formula", all(c_.data_type != "f" for w in wb for row in w.iter_rows() for c_ in row))
            with _patched(pld, XLSX_ROWS_MAX=7):
                c7, xb2, _h2 = ins.export_xlsx({})
            wb2 = openpyxl.load_workbook(io.BytesIO(xb2))
            add("xlsx: the Placements sheet is bounded BEFORE it is built (7 rows + the header)", c7 == 200 and wb2["Placements"].max_row == 8, wb2["Placements"].max_row)
            add("xlsx: ... and the cut is announced in the workbook (L-239)", any("Placements sheet was cut" in str(cell) for row in wb2["Summary"].iter_rows(values_only=True) for cell in row), None)
            add("xlsx: an unknown parameter is a 400 on the report (name only)", ins.export_xlsx({"chanel": ["x"]})[0] == 400)

            # ---- injection: a hostile placement text survives the CSV and the sheet only as text
            add("csv: formula guard - a leading = + - @ in any text cell is neutralised", pld.cd.to_csv(["a"], [["=1+1"], ["+cmd"], ["-x"], ["@y"]]).decode("utf-8-sig").count("'=") >= 1)

    with engine.connect() as c:
        pub_tables2 = sorted(r[0] for r in c.execute(text("SELECT table_name FROM information_schema.tables WHERE table_schema = 'public'")))
        pub_facts2 = c.execute(text("SELECT count(*) FROM public.interaction_fact")).scalar() if "interaction_fact" in pub_tables2 else None
        pub_placement2 = c.execute(text("SELECT to_regclass('public.interaction_placement_rollup') IS NOT NULL")).scalar()
        left = c.execute(text("SELECT count(*) FROM pg_namespace WHERE nspname LIKE 'perf_plc_%'")).scalar()
    add("public: untouched - the same tables, the same interaction_fact row count (the owner's real fact table)", pub_tables == pub_tables2 and pub_facts == pub_facts2, (pub_facts, pub_facts2))
    add("public: this run did not create the placement rollup table there", pub_placement == pub_placement2)
    add("the throwaway schemas were dropped", left == 0, left)


def _id(schema: str, name: str) -> int:
    from sqlalchemy import text
    from erp.db import engine
    with engine.connect() as c:
        return c.execute(text(f"SELECT campaign_id FROM {schema}.marketing_campaign WHERE name = :n"), {"n": name}).scalar()


def _raw_block(store, schema: str, params: dict) -> dict:
    """The raw placement block the store read (for the independent comparisons above): every group of the default view."""
    from desktop import channels_data as cd
    from desktop import leads_data as ld
    with store._connect() as conn:
        rd = ld._Reader(conn, cd.QUERY_BUDGET_SECONDS)
        built = store._read(rd, schema, params, __import__("desktop.placements_data", fromlist=["x"]).KNOWN_PARAMS)
        conn.rollback()
    return built[4]["placements"]


# =========================================================================================== autorun end to end on the sample generator
def _check_sample_and_autorun(add, tmp: Path) -> None:
    from sqlalchemy import text

    from erp.db import engine
    from erp.marketing import autorun, sample_placements
    rows_a = sample_placements.make_rows(400, 20, 3, "EUR", date(2026, 9, 20))
    rows_b = sample_placements.make_rows(400, 20, 3, "EUR", date(2026, 9, 20))
    rows_c = sample_placements.make_rows(400, 20, 4, "EUR", date(2026, 9, 20))
    add("sample generator: the same arguments give the same rows (deterministic)", rows_a == rows_b and rows_a != rows_c)
    add("sample generator: exactly the requested number of rows and the documented header", len(rows_a) == 400 and sample_placements.HEADER == list(_pc().HEADER))
    flat = " ".join(" ".join(map(str, r)) for r in rows_a)
    add("sample generator: everything is synthetic - only .example sites and app:com.example ids, campaigns start with 'Synthetic'", all(r[3].endswith(".example") or r[3].startswith("app:com.example.") or ".example/" in r[3] for r in rows_a) and all(r[1].startswith("Synthetic") for r in rows_a))
    add("sample generator: it plants wasteful placements (real spend, no conversion)", any(int(r[11]) == 0 and float(r[10].split()[0]) > 0 for r in rows_a))
    add("sample generator: viewability is blank for some rows (unknown), not 0", any(r[13] == "" for r in rows_a) and any(r[13] != "" for r in rows_a))
    add("sample generator: it writes no brand or personal data (no @, no http)", "@" not in flat and "http" not in flat)
    with tempfile.TemporaryDirectory(prefix="plc_gen_") as t:
        out = Path(t) / "s.csv"
        rc = sample_placements.main(["--rows", "50", "--days", "5", "--end-date", "2026-09-20", "--out", str(out)])
        conn = _pc().PlacementPerformanceConnector.from_csv(out)
        got = [conn.map_record(r) for r in conn.records()]
        conn._rows.close()
        add("sample generator: its file is accepted by the placement connector with no rejected row", rc == 0 and len(got) == 50 and all(g["currency"] == "EUR" for g in got))
        add("sample generator: bad arguments are refused with exit code 2", sample_placements.main(["--rows", "0", "--out", str(out)]) == 2 and sample_placements.main(["--currency", "EURO", "--out", str(out)]) == 2 and sample_placements.main(["--end-date", "nope", "--out", str(out)]) == 2)

    with engine.connect() as c:
        pub = c.execute(text("SELECT count(*) FROM public.interaction_fact")).scalar()
    with _throwaway(engine, "c") as schema:
        b = Box(tmp, "sample")
        body = ",".join(sample_placements.HEADER) + "\n" + "\n".join(",".join(str(x) for x in r) for r in rows_a) + "\n"
        b.put("synthetic_placements.csv", body)
        b.put("display__usd__nocurrency.csv", ",".join(sample_placements.HEADER) + "\n2026-09-01,Synthetic Alpha,G,a.example,website,300x250,above_fold,mobile,1000,10,5.00,1,,\n")
        rep = autorun.run_inbox(b.inbox, schema=schema)
        st = {o.name: (o.connector, o.status, o.rows_loaded) for o in rep.outcomes}
        add("autorun e2e: a synthetic placement CSV in a temp inbox is detected, loaded and moved to processed/", st["synthetic_placements.csv"] == ("placement_performance", "ok", 400), st)
        add("autorun e2e: a placement file whose Cost cells name no currency takes it from the filename (display__usd__...)", st["display__usd__nocurrency.csv"] == ("placement_performance", "ok", 1), st)
        add("autorun e2e: the rollup ran and reported ok", rep.rollup == "ok", (rep.rollup, rep.rollup_detail))
        with engine.connect() as c:
            add("autorun e2e: the placement rollup holds rows and its impressions equal the fact table's", c.execute(text(f"SELECT sum(impressions) FROM {schema}.interaction_placement_rollup")).scalar() == c.execute(text(f"SELECT sum(impressions) FROM {schema}.interaction_fact")).scalar() > 0)
            add("autorun e2e: the filename's currency reached the fact (USD row) beside the EUR rows", {r[0] for r in c.execute(text(f"SELECT DISTINCT currency FROM {schema}.interaction_fact"))} == {"EUR", "USD"})
        b.put("again.csv", (next(b.processed.glob("*_synthetic_placements.csv"))).read_bytes())
        n0 = 0
        with engine.connect() as c:
            n0 = c.execute(text(f"SELECT count(*) FROM {schema}.interaction_fact")).scalar()
        rep2 = autorun.run_inbox(b.inbox, schema=schema)
        with engine.connect() as c:
            n1 = c.execute(text(f"SELECT count(*) FROM {schema}.interaction_fact")).scalar()
        add("autorun e2e: the same bytes dropped again are a duplicate by hash and load nothing (idempotent)", [o.status for o in rep2.outcomes] == ["duplicate"] and n0 == n1, [o.status for o in rep2.outcomes])
        # the other two connectors in the same inbox
        from tests import marketing_autorun_scenarios as mas
        b2 = Box(tmp, "mixed")
        b2.put("data1.csv", mas.AD_CSV)
        b2.put("data2.csv", mas.EMAIL_CSV)
        b2.put("placements.csv", body.replace("Synthetic", "Synthetic2"))
        rep3 = autorun.run_inbox(b2.inbox, schema=schema)
        st3 = {o.name: (o.connector, o.status) for o in rep3.outcomes}
        add("autorun e2e: in one inbox each file goes to ITS connector - none is stolen", st3 == {"data1.csv": ("ad_performance", "ok"), "data2.csv": ("email_campaign", "ok"), "placements.csv": ("placement_performance", "ok")}, st3)
        with engine.connect() as c:
            src = {r[0]: r[1] for r in c.execute(text(f"SELECT s.source_key, count(*) FROM {schema}.interaction_fact f JOIN {schema}.marketing_source s ON s.source_id = f.source_id GROUP BY 1"))}
            add("autorun e2e: three sources in the fact table, and the placement rollup holds ONLY placement rows", set(src) == {"ad_performance", "email_campaign", "placement_performance"} and c.execute(text(f"SELECT sum(fact_rows) FROM {schema}.interaction_placement_rollup")).scalar() == src["placement_performance"], src)
    with engine.connect() as c:
        pub2 = c.execute(text("SELECT count(*) FROM public.interaction_fact")).scalar()
    add("autorun e2e: public.interaction_fact is untouched", pub == pub2, (pub, pub2))


# =========================================================================================== Data Flow map, docs
def _check_flow_and_docs(add) -> None:
    from desktop import flow_data
    from desktop import flow_definition as fd
    from desktop.flow_mapcheck import map_integrity
    chk = flow_data.check_map(ROOT)
    add("flow map: the map is in sync - no unmapped script, no stale node, no stale IGNORE entry", chk["in_sync"] and not chk["unmapped"] and not chk["stale_nodes"] and not chk["stale_ignores"], (chk["unmapped"], chk["stale_ignores"]))
    add("flow map: no inconsistent node or edge", map_integrity() == [], map_integrity())
    ingest = next(n for n in fd.NODES if n["id"] == "marketing_ingest")
    roll = next(n for n in fd.NODES if n["id"] == "marketing_rollup")
    live = next(n for n in fd.NODES if n["id"] == "live_report")
    add("flow map: the ingest node lists the placement connector and names four connectors", "erp/marketing/connectors/placement_performance.py" in ingest["files"] and "Four connectors" in ingest["summary"] and "placement_performance" in ingest["inputs"][0])
    add("flow map: the rollup node lists rollup_placements.py and migration 11 and the new output table", "erp/marketing/rollup_placements.py" in roll["files"] and "db/sql/11_marketing_placements.sql" in roll["files"] and any("interaction_placement_rollup" in o for o in roll["outputs"]))
    add("flow map: the Live Reporting node covers the placements feed and says which table it reads and that db/sql/11 is optional", "desktop/placements_data.py" in live["file"] and "interaction_placement_rollup" in " ".join(live["inputs"]) and "db/sql/11" in " ".join(live["inputs"]))
    add("flow map: the rollup edge names the placement table", any("interaction_placement_rollup" in e["data"] for e in fd.EDGES if e["id"] == "e_rollup_db"))
    add("flow map: every automated edge still declares armed_by", all(e.get("armed_by") for e in fd.EDGES if e["trigger"] in ("scheduled", "background", "webhook")), [e["id"] for e in fd.EDGES if e["trigger"] in ("scheduled", "background", "webhook") and not e.get("armed_by")])
    add("flow map: the synthetic generator is mapped as IGNORE with a reason", "erp/marketing/sample_placements.py" in fd.IGNORE and "SYNTHETIC" in fd.IGNORE["erp/marketing/sample_placements.py"])
    add("flow map: the rollup step is not a new job (no new node, no new trigger): it rides the existing rollup run", len([n for n in fd.NODES if "placement" in n["id"]]) == 0)
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    arch = (ROOT / "docs" / "marketing-data-architecture.md").read_text(encoding="utf-8")
    add("docs: README has a Placements section with the command to try it", "placement" in readme.lower() and "sample_placements" in readme and "11_marketing_placements.sql" in readme)
    add("docs: the architecture document describes the placement rollup and its measurement", "interaction_placement_rollup" in arch and "placement_performance" in arch)
    sql = (ROOT / MIG_11).read_text(encoding="utf-8")
    add("migration 11: additive and idempotent - only CREATE ... IF NOT EXISTS, no ALTER, DROP or DELETE of anything existing", "IF NOT EXISTS" in sql and not re.search(r"^\s*(ALTER|DROP|DELETE|TRUNCATE|UPDATE)\b", sql, re.I | re.M))
    add("migration 11: it names no company, brand or person", not re.search(r"@|http", sql))
    pyfiles = [ROOT / "erp/marketing/connectors/placement_performance.py", ROOT / "erp/marketing/rollup_placements.py", ROOT / "erp/marketing/sample_placements.py", ROOT / "desktop/placements_data.py", ROOT / "desktop/placements_queries.py", ROOT / "tests/placement_scenarios.py"]
    add("hygiene: the new modules carry a docstring and never call logging.basicConfig at import time (L-010)", all(f.read_text(encoding="utf-8").lstrip().startswith('"""') for f in pyfiles) and not any("basicConfig" in f.read_text(encoding="utf-8").split("def main")[0] for f in pyfiles[:5]))
    add("hygiene: no new dependency (only the standard library and packages already used)", not re.search(r"^\s*(import|from)\s+(pandas|numpy|requests|scipy)\b", "\n".join(f.read_text(encoding="utf-8") for f in pyfiles[:5]), re.M))


# =========================================================================================== runner
def run(db_ok: bool = True) -> list[tuple[str, bool, object]]:
    rows: list[tuple[str, bool, object]] = []

    def add(name: str, ok, detail=None) -> None:
        rows.append((name, bool(ok), detail if not ok else ""))

    loggers = [logging.getLogger(n) for n in ("erp_desk.channels", "erp_desk.leads", "erp_desk.insights", "erp_desk.placements", "erp.marketing", "erp.marketing.autorun",
                                              "erp.marketing.pipeline", "erp.marketing.rollup", "erp.marketing.rollup_placements", "erp.marketing.ingest")]
    old = [lg.disabled for lg in loggers]
    for lg in loggers:
        lg.disabled = True
    logging.disable(logging.CRITICAL)
    try:
        with tempfile.TemporaryDirectory(prefix="placement_scn_") as t:
            tmp = Path(t)
            fns = [(_check_connector, False), (_check_detection, False), (_check_rules, False), (_check_pure_feed, False), (_check_flow_and_docs, False)]
            if db_ok:
                fns += [(_check_database, True), (_check_sample_and_autorun, True)]
            else:
                add("database scenarios skipped (no database)", True)
            for fn, needs_tmp in fns:
                try:
                    fn(add, tmp) if needs_tmp else fn(add)
                except Exception as exc:  # noqa: BLE001
                    import traceback
                    add(f"{fn.__name__} crashed", False, f"{type(exc).__name__}: {exc} | {' / '.join(traceback.format_exc().splitlines()[-4:-1])[:300]}")
    finally:
        logging.disable(logging.NOTSET)
        for lg, o in zip(loggers, old):
            lg.disabled = o
        with contextlib.suppress(Exception):
            from desktop import leads_data as ld
            ld._logged.clear()
    return rows


def main() -> int:
    try:
        from sqlalchemy import text

        from erp.db import engine
        with engine.connect() as c:
            c.execute(text("SELECT 1"))
        db_ok = True
    except Exception:  # noqa: BLE001
        db_ok = False
    t0 = time.monotonic()
    rows = run(db_ok)
    for name, ok, detail in rows:
        print(f"{'ok  ' if ok else 'FAIL'} {name}" + (f"   <- {detail}" if not ok else ""))
    bad = [r for r in rows if not r[1]]
    print(f"\n{len(rows) - len(bad)}/{len(rows)} scenarios passed in {time.monotonic() - t0:.1f}s")
    return 1 if bad else 0


if __name__ == "__main__":
    for _stream in (sys.stdout, sys.stderr):
        try:
            _stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass
    sys.exit(main())
